"""The watermark (a deck's watermark:, the livery's look) on every slide of a deck.

Its text is `engine-watermark` in docProps/custom.xml, which slide-layout.lua writes from
engine-common.lua's one reader (the deck's watermark:, else the livery's default): "none", no
watermark. This step draws and never decides. A text box "Watermark", the slide's width wide and
its size's line tall, centered on the slide and turned (OOXML turns clockwise: rot is 360 less the
livery's counter-clockwise angle), one bold run in its color with the livery's opacity as an
a:alpha; the color is the livery's on a content slide, a composed slide's text color (its
composition's column role: a cover, divider or closing by its layout, a feature by its photo) and
on-dark on a dark slide, as HTML and the PDF choose. It goes in before the slide's first shape that
holds text (a content slide's title): behind the body and the title, over the master's furniture,
a cover's photo, field, waves and marks, a feature's photo and bar. type_step, after it, scales its
size with the deck's."""

from pathlib import Path

from lxml import etree
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Pt

import livery as livery_mod

from . import docprops
from .common import BAND, LIVERY, PHOTO, is_dark

WM = LIVERY["watermark"]   # color hex, opacity, angle (degrees counter-clockwise), size in points


def ink(slide) -> str:
    """The watermark's six-digit hex on this slide (the module's docstring)."""
    if is_dark(slide):
        return livery_mod.hex(LIVERY, "on-dark")
    comp = livery_mod.comp_of(slide.slide_layout.name, BAND)   # a composed slide's, by its layout
    if comp is None and any(sh.name == PHOTO for sh in slide.shapes):   # slide_layout.draw_feature()'s
        comp = "feature"
    return livery_mod.hex(LIVERY, LIVERY["compositions"][comp]["column"]["role"] if comp else WM["color"])


def run(prs, path: Path) -> None:
    text = docprops.prop(prs, "engine-watermark")
    if not text or text == "none":
        return
    W, H = prs.slide_width, prs.slide_height
    h = int(Pt(WM["size"]) * 1.2)
    for slide in prs.slides:
        box = slide.shapes.add_textbox(0, (H - h) // 2, W, h)
        box.name = "Watermark"
        box.rotation = (-WM["angle"]) % 360
        tf = box.text_frame
        tf.word_wrap = False
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = text
        r.font.size = Pt(WM["size"])
        r.font.bold = True
        r.font.color.rgb = RGBColor.from_string(ink(slide))
        clr = r._r.find(qn("a:rPr")).find(".//" + qn("a:srgbClr"))   # its alpha, in thousandths of a percent
        etree.SubElement(clr, qn("a:alpha")).set("val", str(round(WM["opacity"] * 100000)))
        tree = slide.shapes._spTree   # behind the first shape that holds text (the module's docstring)
        first = next((sh._element for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.strip()
                      and sh._element is not box._element), None)
        if first is not None:
            tree.remove(box._element)
            tree.insert(list(tree).index(first), box._element)
    print(f"watermark: {text!r} on {len(prs.slides)} slide(s) → {path.name}")
