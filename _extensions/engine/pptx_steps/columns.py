"""A two-column slide's columns at their widths and the livery's gap, as HTML and PDF lay them out.

pandoc's pptx writer puts a `.columns` div's two columns into the Two Content layout's placeholders,
each the template's half width, and drops their `width=`. grid-pptx.lua leaves the shares on a
notes line, `engine-columns: 0.55 0.45` (`wide` added on a `{.wide}` slide); this reads and
deletes the line and frames the columns: the content box is the slide less the margin a side
(geometry.wide-margin on a wide slide), each column its share of the box less the gaps
(boxes.columns.gap), x_1 at the margin and each next one a gap after the last; a livery's
boxes.columns.rule draws a rule centered in each gap, the row's height. Text placeholders
take the frame (all four fields: a placeholder inherits its frame, and setting one leaves the
others at 0). A picture or table pandoc placed in a column moves with it: a picture as wide as the
old column scales to the new one (its aspect kept, no taller than the column), a narrower one only
shrinks if it no longer fits, centered either way; a table takes the column's width, its grid
columns scaled to it (style_tables then re-splits them inside it). A columns slide pandoc laid out
on another layout (Comparison: text above a picture or table in a column) keeps the template's
halves, with a warning naming the slide. Idempotent: the line is gone
after one pass. After slide_layout (which leaves a columns slide's frames to this), before
style_tables and type_step.
"""

import sys

from pptx.util import Emu

import livery as livery_mod

from .common import BOXES, GEOM, LIVERY, column_frames, fit_table, is_picture, is_table, rect, take_marker
from .measure import text_height

RULE = livery_mod.blend(LIVERY, "rule")   # HTML's --engine-rule, beamer's BrandRule

KEY = "engine-columns:"


def frames(shares, width, wide=False):
    """(left, width) in EMU of each column, for shares of the content box less its gaps
    (common.column_frames, grid.py's too)."""
    m = GEOM["wide-margin" if wide else "margin"]
    return column_frames(round(m * width), round((1 - 2 * m) * width), shares, width)


def rules(slide, frames, top, height, slide_width, slide_height):
    """boxes.columns.rule: a rule centered in each gap between frames ((left, width) EMU), its width
    a fraction of the slide's height, from top for height (the columns' row: its tallest column, as
    HTML's and beamer's rules run); none when the livery's rule is 0."""
    w = round(BOXES["columns"]["rule"] * slide_height)
    if w <= 0:
        return
    g = round(BOXES["columns"]["gap"] * slide_width)
    for x, cw in frames[:-1]:
        rect(slide, x + cw + (g - w) // 2, top, w, height, RULE, "Column rule")


def _old(width):
    """The template's two half frames (livery_template.py's Two Content: geometry.half), (left, width) in EMU."""
    m, g, half = GEOM["margin"], BOXES["columns"]["gap"], GEOM["half"]
    return [(round(m * width), round(half * width)), (round((m + half + g) * width), round(half * width))]


def _column_of(shape, old):
    """The old column a shape sits in, by its center."""
    c = shape.left + shape.width // 2
    return min(range(len(old)), key=lambda i: abs(old[i][0] + old[i][1] // 2 - c))


def frame_slide(slide, words, width, height):
    try:
        shares = [float(w) for w in words if w != "wide"]
    except ValueError:
        return False
    if len(shares) != 2:
        return False
    new, old = frames(shares, width, "wide" in words), _old(width)
    area = {}   # each column's top and height, from its placeholder (the layout's when inherited)
    row = []    # each column's top and the foot of its content, for the rule
    for sh in slide.placeholders:
        idx = sh.placeholder_format.idx
        if idx in (1, 2) and sh.has_text_frame:
            x, w = new[idx - 1]
            sh.left, sh.top, sh.width, sh.height = x, sh.top, w, sh.height
            if sh.text_frame.text.strip():
                row.append((sh.top, sh.top + min(text_height(sh, slide), sh.height)))
    for ph in slide.slide_layout.placeholders:
        if ph.placeholder_format.idx in (1, 2):
            area[ph.placeholder_format.idx - 1] = (ph.top, ph.height)
    for sh in slide.shapes:
        if not (is_table(sh) or is_picture(sh)):
            continue
        i = _column_of(sh, old)
        x, w = new[i]
        if is_table(sh):
            fit_table(sh, x, w)
            row.append((sh.top, sh.top + sh.height))
            continue
        top, height = area.get(i, (sh.top, sh.height))
        k = w / sh.width if abs(sh.width - old[i][1]) <= old[i][1] * 0.02 else min(1.0, w / sh.width)
        k = min(k, height / sh.height)
        nw, nh = round(sh.width * k), round(sh.height * k)
        cy = sh.top + sh.height // 2
        sh.left, sh.top, sh.width, sh.height = Emu(x + (w - nw) // 2), Emu(cy - nh // 2), Emu(nw), Emu(nh)
        row.append((sh.top, sh.top + sh.height))
    if row:
        top = min(t for t, _ in row)
        rules(slide, new, top, max(b for _, b in row) - top, width, height)
    return True


def run(prs, name: str) -> None:
    done = 0
    for slide in prs.slides:
        words = take_marker(slide, KEY)
        if words is None:
            continue
        if slide.slide_layout.name == "Two Content":
            done += frame_slide(slide, words, prs.slide_width, prs.slide_height)
        else:   # pandoc chose another layout (Comparison): its frames are not the two placeholders'
            print(f"columns: WARNING: slide {prs.slides.index(slide) + 1}: columns laid out on "
                  f"{slide.slide_layout.name!r}, not Two Content; their widths are not applied", file=sys.stderr)
    if done:
        print(f"columns: framed {done} two-column slide(s) → {name}")
