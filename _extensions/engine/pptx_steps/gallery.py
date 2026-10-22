"""An image gallery, drawn from grid-pptx.lua's markers (grid.py's lay_out() hands it the segment).

grid-pptx.lua writes a top-level `::: {.gallery}` that engine-common.lua's gallery_items() reads as
`⟦engine:gallery cols=3 rows=2 fit=contain|cover[ height=0.6][ full]⟧`, per image `⟦engine:item
src=<path>⟧` (the path percent-encoded where it holds white space, `=`, `⟦` or `⟧`, as pandoc's URL form
of a src writes a space; decoded here once, to the file's own name),
its caption a paragraph, `⟦engine:end⟧`, and a last `⟦engine:end⟧`: text, so pandoc keeps Title and
Content and nothing splits. This reads each file from disk, relative to the deck's directory
(pandoc never embedded it; an executed cell's plot or a baked diagram is refused by the reader,
since it is not on disk after a render), an .svg rasterized to a PNG (no svgBlip twin: that is
svg_fallback's zip surgery, which this Presentation stage cannot do), and draws the cells row-major
and equal, boxes.gallery.gap apart both ways (of the slide's width): per cell a picture in its box
less boxes.gallery.pad, fit contain (the whole image, centered) or cover (cropped symmetrically, an
a:srcRect), and a caption under it in the caption size, muted, centered, the gallery's one caption
height (its tallest, measured) reserved in every cell when any image has one, so a row's pictures
and captions line up; a cell of boxes.gallery.fill behind when the livery sets one. The box is the
body's room under what is above it, with height= that share of the body box's height (at most the
room), and on a {.full} slide (slide_layout.py has hidden the furniture) the slide less
geometry.full's pads, footline and caption line, the slide caption (the text after the gallery) in
that line. One group, "Gallery": "Gallery k picture", "Gallery k caption", "Gallery k cell", and
"Gallery caption" on a full slide. A missing file warns and leaves its cell empty.
"""

import sys
from urllib.parse import unquote

from lxml import etree
from pptx.oxml.ns import qn
from pptx.util import Emu

import livery as livery_mod

from .common import BOXES, GEOM, LIVERY, SLACK, body_box, fit_picture, group, pads, ppr_place, rect, style_runs
from .measure import text_height
from .svg_fallback import DPI, rasterize_svg   # an .svg item's PNG at the deck's fig-dpi

CAPTION_PT = livery_mod.type_pt(LIVERY, "caption")
MUTED = livery_mod.hex(LIVERY, "muted")


def cells(x, y, w, h, cols, rows, gap):
    """The cells' frames (x, y, w, h EMU), row-major: the box less the gaps, shared equally both ways."""
    cw, ch = round((w - gap * (cols - 1)) / cols), round((h - gap * (rows - 1)) / rows)
    return [(x + c * (cw + gap), y + r * (ch + gap), cw, ch) for r in range(rows) for c in range(cols)]


def _caption(p):
    """A caption paragraph: centered, no space before it, the caption size, muted."""
    ppr = p.find(qn("a:pPr"))
    if ppr is None:
        ppr = etree.Element(qn("a:pPr")); p.insert(0, ppr)
    ppr.set("algn", "ctr")
    for old in ppr.findall(qn("a:spcBef")):
        ppr.remove(old)
    bef = etree.Element(qn("a:spcBef"))
    etree.SubElement(bef, qn("a:spcPts")).set("val", "0")
    ppr_place(ppr, bef)
    style_runs(p, CAPTION_PT, color=MUTED)


def _title(slide):
    t = slide.shapes.title
    return t.text_frame.text if t is not None else ""


def draw_gallery(slide, box, attrs, items, caption, slide_width, slide_height, deck_dir):
    """The gallery in box (x, y, w, h EMU: the room under what is above it); attrs the marker's, items
    [("item", attrs, [caption paragraphs])], caption the slide caption's paragraphs on a full slide
    (else None), deck_dir where the files are. The height it takes, a vertical pad under it (none on
    a full slide)."""
    g = BOXES["gallery"]
    cols, rows, fit = int(attrs.get("cols", 1)), int(attrs.get("rows", 1)), attrs.get("fit", "contain")
    gap = round(g["gap"] * slide_width)
    vpad, hpad = pads("gallery")
    full = bool(attrs.get("full"))
    shapes = []
    x, y, w, h = box
    if full:
        f = GEOM["full"]
        pad, floor = round(f["pad"] * slide_width), round(f["footline"] * slide_height)
        cap = round(f["caption"] * slide_height) if caption else 0
        x, y, w, h = pad, pad, slide_width - 2 * pad, slide_height - pad - floor - cap
        if caption:
            for p in caption:
                _caption(p)
            cb = body_box(slide, pad, slide_height - floor - cap, slide_width - 2 * pad, cap, caption, "Gallery caption")
            cb.text_frame._txBody.find(qn("a:bodyPr")).set("anchor", "ctr")
            shapes.append(cb)
    elif attrs.get("height"):
        h = min(h, round(float(attrs["height"]) * GEOM["body"][3] * slide_height))
    frames = cells(x, y, w, h, cols, rows, gap)
    caps = []
    for k, ((_, _, paras), (cx, cy, cw, ch)) in enumerate(zip(items, frames), 1):
        cb = None
        if paras:
            for p in paras:
                _caption(p)
            cb = body_box(slide, cx + hpad, cy, max(cw - 2 * hpad, 1), ch, paras, f"Gallery {k} caption")
            tf = cb.text_frame
            tf.margin_top, tf.margin_bottom, tf.margin_left, tf.margin_right = vpad, 0, 0, 0   # vpad: under the picture
        caps.append(cb)
    cap_h = max((int(text_height(cb, slide, gap=0) * SLACK) for cb in caps if cb is not None), default=0)
    for k, ((_, iattrs, _), (cx, cy, cw, ch), cb) in enumerate(zip(items, frames, caps), 1):
        px, py, pw, ph = cx + hpad, cy + vpad, max(cw - 2 * hpad, 1), max(ch - 2 * vpad - cap_h, 1)
        if g["fill"]:
            shapes.append(rect(slide, cx, cy, cw, ch, livery_mod.hex(LIVERY, g["fill"]), f"Gallery {k} cell"))
        src = deck_dir / unquote(iattrs.get("src", ""))
        if src.is_file():
            data = src.read_bytes()
            if src.suffix.lower() == ".svg":
                data = rasterize_svg(data, DPI)
            shapes.append(fit_picture(slide, data, px, py, pw, ph, fit, name=f"Gallery {k} picture"))
        else:
            print(f"gallery: WARNING: slide {_title(slide)!r}: {src} not found; its cell left empty", file=sys.stderr)
        if cb is not None:
            cb.left, cb.top, cb.width, cb.height = Emu(px), Emu(py + ph), Emu(pw), Emu(cap_h)
            shapes.append(cb)
    if shapes:
        group(slide, shapes, "Gallery")
    return h if full else h + vpad
