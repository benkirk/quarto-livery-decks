#!/usr/bin/env python3
"""Caption linked pictures in a pandoc-rendered pptx.

A linked image in the deck, [![](shot.png)](https://...), becomes a picture
with a click hyperlink, which PowerPoint shows only as a hand cursor in slide
show. This step adds a small muted caption with the URL at the bottom center
of each such slide (the HTML and PDF decks get the same line from
the engine's linked-images.lua). The caption box is itself the link, so its text
keeps the muted color instead of PowerPoint's hyperlink blue, and the picture
gets the URL as its hover tooltip.

Runs after pandoc, so it cannot split a slide. Idempotent: a slide that
already has a caption is skipped."""

from pathlib import Path

from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Pt

import livery as livery_mod

from .common import LIVERY, is_picture, palette

NAME       = "link-caption"
SIZE       = Pt(livery_mod.type_pt(LIVERY, "caption"))   # the caption role
BAND_WIDTH = 0.6                             # of the slide; clear of number and date
BAND_TOP   = 0.056                           # of the height, above the slide's bottom edge
BAND_TALL  = 0.04


def _picture_link(shape):
    if not is_picture(shape):
        return None
    try:
        return shape.click_action.hyperlink.address
    except (AttributeError, KeyError):
        return None


def _set_tooltip(shape, url):
    for hlink in shape._element.xpath(".//a:hlinkClick"):
        hlink.set("tooltip", url)


def run(prs, path: Path) -> None:
    """The open deck, changed in place; path for the message."""
    width = int(prs.slide_width * BAND_WIDTH)
    left = (prs.slide_width - width) // 2
    top = round(prs.slide_height * (1 - BAND_TOP))
    added = 0
    for slide in prs.slides:
        if any(shape.name == NAME for shape in slide.shapes):
            continue
        url = None
        for shape in slide.shapes:
            link = _picture_link(shape)
            if link:
                _set_tooltip(shape, link)
                url = url or link
        if not url:
            continue
        box = slide.shapes.add_textbox(left, top, width, round(prs.slide_height * BAND_TALL))
        box.name = NAME
        box.click_action.hyperlink.address = url
        para = box.text_frame.paragraphs[0]
        para.alignment = PP_ALIGN.CENTER
        run = para.add_run()
        run.text = url.split("://", 1)[-1]
        run.font.size = SIZE
        run.font.color.rgb = RGBColor.from_string(palette(slide)["muted"])   # a dark slide's dark-muted
        added += 1
    print(f"link_captions: captioned {added} linked picture slide(s) → {path.name}")
