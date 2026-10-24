#!/usr/bin/env python3
"""Restyle on-slide footnotes in a pandoc-rendered pptx.

pandoc's pptx writer ignores inline size and color, so a footnote authored in the .qmd (a
paragraph starting with `†`, `‡`, `§` or `¶`; `*† text*` works too; a `^[...]` note, which
inline-footnotes.lua writes as one) renders like any
paragraph. This finds each such paragraph, or a caption line that holds one, and restyles it as
HTML and the PDF do, from boxes.footnote: left-aligned in the livery's muted color at the
footnote size, the marker bold in the accent role and hanging (marL/indent by `hang` em of that
size), `gap` em above the first footnote of a text frame and `step` em above each following one;
a lifted note's mark in the text (a superscript marker) takes the accent too. HTML and PDF also
move footnotes to the slide's foot under a short rule, which pptx cannot do without splitting the
slide (the paragraph's position is not known before PowerPoint opens the file), so it draws no
rule and the footnotes stay in place.

Idempotent; a deck with no footnotes is untouched."""

import copy
from pathlib import Path

from lxml import etree
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Pt
from pptx.dml.color import RGBColor

import livery as livery_mod

from .common import BOXES, LIVERY, em, palette, ppr_place

MARKERS     = ("†", "‡", "§", "¶")     # the printer's order (engine-common.lua MARKERS)
FOOT_PT     = livery_mod.type_pt(LIVERY, "footnote")   # the themes' footnote size (livery.json type)
FOOT_SIZE   = Pt(FOOT_PT)
FOOT_COLOR  = RGBColor.from_string(livery_mod.color(LIVERY, "muted"))
MARK_COLOR  = RGBColor.from_string(livery_mod.hex(LIVERY, "accent"))   # the accent role, as HTML (--engine-accent) and beamer (BrandAccent)
FOOTNOTE    = BOXES["footnote"]        # hang, gap and step in em of the footnote type (boxes.footnote)


def _footnote_start(para):
    """Index of the run a footnote starts at: the first, or the first after a line
    break (single-body.lua folds a footnote into a table's caption, which pandoc
    joins with line breaks). None when the paragraph holds no footnote."""
    runs = para.runs
    if not runs:
        return None
    if runs[0].text.lstrip().startswith(MARKERS):
        return 0
    after_break, index = False, -1
    for child in para._p:
        if child.tag == qn("a:br"):
            after_break = True
        elif child.tag == qn("a:r"):
            index += 1
            if after_break and runs[index].text.lstrip().startswith(MARKERS):
                return index
            after_break = False
    return None


def _split_marker(para, start):
    """The run at start, cut so the marker is a run of its own (idempotent)."""
    first = para.runs[start]
    text = first.text.lstrip()
    marker = text[0]
    if text.rstrip() != marker:
        mark = copy.deepcopy(first._r)
        first._r.addprevious(mark)
        first.text = text[1:].lstrip()
        para.runs[start].text = marker + " "
    return para.runs[start]


def style(para, following: bool = False, color: str | None = None) -> bool:
    """The paragraph's footnote styled, if it holds one (slide_layout.py calls it again on a note
    line it splits out of a caption: idempotent); `following`, one right after another footnote,
    takes the step above it instead of the gap; color its text's (six digits: the slide's muted,
    common.palette()), FOOT_COLOR by default."""
    start = _footnote_start(para)
    if start is None:
        return False
    if start == 0:
        # a caption is centered; a footnote reads from the left, as in HTML and PDF, its marker
        # hanging in the margin: marL/indent by the hang, no bullet
        para.space_before = Pt((FOOTNOTE["step"] if following else FOOTNOTE["gap"]) * FOOT_PT)
        para.alignment = PP_ALIGN.LEFT
        hang = em(FOOTNOTE["hang"], FOOT_PT)
        ppr = para._p.get_or_add_pPr()
        ppr.set("marL", str(hang))
        ppr.set("indent", str(-hang))
        for old in ppr.findall(qn("a:buChar")) + ppr.findall(qn("a:buAutoNum")):
            ppr.remove(old)
        if ppr.find(qn("a:buNone")) is None:
            ppr_place(ppr, etree.Element(qn("a:buNone")))
    mark = _split_marker(para, start)
    for run in para.runs[start:]:  # every run — the code span too
        run.font.size = FOOT_SIZE
        run.font.color.rgb = MARK_COLOR if run._r is mark._r else (RGBColor.from_string(color) if color else FOOT_COLOR)
    mark.font.bold = True   # as HTML's .engine-dagger and beamer's footnote mark
    return True


def noteref(para) -> int:
    """A lifted ^[...] note's mark in the text (inline-footnotes.lua: a superscript holding one
    marker) colored as the footnote's marker is, since HTML (.engine-noteref) and PDF (footnote
    mark) set it in the accent and pandoc drops inline color; the count colored."""
    marks = 0
    for run in para.runs:
        rpr = run._r.find(qn("a:rPr"))
        if rpr is not None and int(rpr.get("baseline") or 0) > 0 and run.text.strip() in MARKERS:
            run.font.color.rgb = MARK_COLOR
            marks += 1
    return marks


def run(prs, path: Path) -> None:
    """The open deck, changed in place; path for the message."""
    styled, marks = 0, 0
    for slide in prs.slides:
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            previous = False   # a footnote right after another takes the step, the first the gap
            muted = palette(slide)["muted"]   # a dark slide's dark-muted
            for para in shape.text_frame.paragraphs:
                previous = style(para, following=previous, color=muted)
                styled += previous
                marks += noteref(para)
    print(f"style_footnotes: restyled {styled} footnote paragraph(s), {marks} mark(s) in the text → {path.name}")
