#!/usr/bin/env python3
"""Apply the theme's per-slide layout controls to a pandoc-rendered pptx.

pandoc drops slide classes, so the engine's slide-layout.lua leaves one speaker-notes line,
"engine-layout: vcenter scale=1.4", on each slide that asks for a layout; this removes the line
and applies what pptx can express. vcenter: body text anchored mid-placeholder, a table moved to
the middle of its area (its rows measured; PowerPoint sizes rows itself). scale=S: every body and
table run at S times its inherited size. hcenter, fill: nothing (a text box cannot shrink-wrap).
caution: the body on the livery's caution-fill with a warning-colored bar at its left (footnotes
stay inside; pptx cannot move them). full: the one picture fills the slide, no title or master
art, pandoc's caption box centered under it in the muted color, the slide number kept; a gallery's
full slide (grid-pptx.lua's marker in its body) only loses its furniture here, grid.py fills it.
note (single-body.lua's word for a `::: note` folded into the caption box): the box's lines each a
paragraph, the note's in the caption size and muted color, centered, a footnote's left as
style_footnotes styled it, and the box moved up to just under the picture or table.
references (slide-layout.lua's word for the slide holding ::: {#refs}, which citeproc fills): every
body paragraph in the footnote size with a hanging indent and a little space above, as HTML and
PDF set the list.
wide: the body placeholders and a table take the slide less geometry.wide-margin a side (a columns
slide's are columns.py's); a picture keeps its size, as HTML's does.
closing: the slide moves onto the template's Closing layout (Closing (light) under themeoptions
title=light, which slide-layout.lua records as `engine-title: light` and which also moves the
title slide onto Title Slide (light); a band title ignores it), its placeholders dropping pandoc's
copied frames to take the layout's. The title slide with a `titlegraphic:` (docProps/custom.xml;
else the livery's art.titlegraphic) is drawn from its composition (livery.json compositions):
the photo at the back, cropped to cover its box, the field and waves over it, the marks in their
boxes; the band takes the photo into its box over the layout's default. The cover's byline and fine
print (the deck's fineprint:, else the livery's default) are text boxes in theirs. feature: by the
livery's feature.style, the photo under a dark bar with the title and the body on a panel, or the
cover with the title and body in its column. Footnote paragraphs
(style_footnotes.py) keep their size. Every slide also gets its layout's slide-number placeholder, which pandoc never writes
(so decks had no numbers). Idempotent: the marker is gone after one pass, a slide that has a
number keeps it."""

import copy
import sys
from pathlib import Path

from lxml import etree
from pptx.dml.color import RGBColor
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.text.text import _Run
from pptx.util import Emu, Pt

import livery as livery_mod
import pptx_waves   # the waves, shared with the engine's template builder

from . import docprops, measure, style_footnotes, style_tables
from .common import (BOXES, GEOM, K, LIVERY, SLACK, contain_box, em, fit_picture, fit_table, inset, is_body, is_picture, is_table,
                     marker_spans, on_dark, pads, palette, para_text, ppr_place, rect, rpr, slide_title, split_br, take_marker)
from .measure import LINE, inherited, other_size, para_lines, para_size, pieces, text_height

MARKER    = "engine-layout:"
FULL_PAD   = LIVERY["geometry"]["full"]["pad"]       # full: margin round the figure (of the width),
FULL_FLOOR = LIVERY["geometry"]["full"]["footline"]  # the footline (slide number) and the caption
FULL_CAP   = LIVERY["geometry"]["full"]["caption"]   # line (of the height), as the themes draw them
CAP_SIZE   = Pt(livery_mod.type_pt(LIVERY, "caption"))   # the themes' caption size, the muted color
CAP_COLOR  = RGBColor.from_string(livery_mod.color(LIVERY, "muted"))
NOTE_GAP   = Pt(6 * K)    # note: between the picture or table and the note under it
REF_PT     = livery_mod.type_pt(LIVERY, "footnote")   # references: the list in the footnote size,
REF_HANG   = BOXES["references"]["hang"]   # hung and spaced by boxes.references (em of it), as HTML and PDF
REF_BEFORE = BOXES["references"]["gap"]
CAUTION_FILL = RGBColor.from_string(livery_mod.blend(LIVERY, "caution-fill"))  # warning at 18% on white
# the warning role, not the accent5 slot: the ncar template's accent5 (FDD509) is not brand yellow
CAUTION_BAR  = livery_mod.color(LIVERY, "warning")
SLIDE_H      = Pt(LIVERY["type"]["canvas"]["h"])   # the canvas height: bars are fractions of it
CAUTION_W    = Emu(round(BOXES["caution"]["bar"] * SLIDE_H))
WAVES        = LIVERY.get("art", {}).get("waves")
BAND         = (LIVERY.get("title") or {}).get("style") == "band"
COMP         = LIVERY["compositions"]   # the title, closing, section and feature slides as boxes


def _scale_paragraphs(paragraphs, scale, size_at):
    for p in paragraphs:
        if p.text.lstrip().startswith(style_footnotes.MARKERS):
            continue
        base = size_at(p.level)
        for r in p._p.iter(qn("a:rPr"), qn("a:endParaRPr")):
            sz = int(r.get("sz")) if r.get("sz") else base
            r.set("sz", str(round(livery_mod.half_pt(sz * scale / 100) * 100)))   # to the half point


def _table_height(table, other, biggest, pad=2 * style_tables.CELL_MAR):
    """A table's height: each row as tall as its cell of most lines, measured in the livery's base
    face in its column's width (less the cell's margins), at the table's biggest size, plus pad a
    row: the cell's top and bottom margins as style_tables sets them (it runs after this step, so
    the cells still carry PowerPoint's defaults here)."""
    widths = [col.width for col in table.columns]
    total = 0
    for row in table.rows:
        most = 1
        for cell, w in zip(row.cells, widths):
            tf = cell.text_frame
            inner = w - cell.margin_left - cell.margin_right
            n = 0
            for p in tf.paragraphs:
                sz = next((int(r.get("sz")) for r in p._p.iter(qn("a:rPr")) if r.get("sz")), other)
                n += para_lines(p, measure.BODY, Pt(sz / 100), inner)
            most = max(most, n)
        total += int(most * Pt(biggest / 100) * LINE + pad)
    return total


def table_height(shape, slide):
    """A table's measured height on its slide (_table_height at the slide's size): what a note
    under it or a slide centering it reads, since the frame's own height is pandoc's nominal."""
    other = other_size(slide)
    return _table_height(shape.table, other, _biggest(shape.table, other))


def _biggest(table, other):
    """The size (hundredths of a point) of a table's biggest run, at least other."""
    return max([other] + [int(r.get("sz")) for r in table._tbl.iter(qn("a:rPr")) if r.get("sz")])


def caution_box(slide, shape, vcenter):
    """The body text box on the soft yellow field, a brand-yellow bar at its left, cut down to
    about its text (pandoc gives it the whole content area); 10% slack keeps autofit idle."""
    vpad, hpad = pads("caution")
    inset(shape.text_frame, vpad, hpad)
    left, top, width, height = shape.left, shape.top, shape.width, shape.height
    est = int(text_height(shape, slide) * SLACK)
    if est < height:
        if vcenter:
            top += (height - est) // 2
        height = est
    # all four: a placeholder inherits its frame, and setting one leaves the rest at 0
    shape.left, shape.top, shape.width, shape.height = left, top, width, height
    shape.fill.solid()
    shape.fill.fore_color.rgb = CAUTION_FILL
    rect(slide, shape.left, shape.top, CAUTION_W, shape.height, CAUTION_BAR)


def lay_out(slide, words, width):
    if "wide" in words and slide.slide_layout.name != "Two Content":   # columns.py frames a columns slide
        m = GEOM["wide-margin"]
        for shape in [sh for sh in slide.shapes if is_body(sh)]:   # all four: a placeholder inherits its frame
            shape.left, shape.top, shape.width, shape.height = round(m * width), shape.top, round((1 - 2 * m) * width), shape.height
        for shape in [sh for sh in slide.shapes if is_table(sh)]:   # a graphicFrame, not a body
            fit_table(shape, round(m * width), round((1 - 2 * m) * width))
    scale = 1.0
    for w in words:
        if w.startswith("scale="):
            try:
                scale = float(w[len("scale="):])
            except ValueError:
                pass
    vcenter = "vcenter" in words
    for shape in slide.shapes:
        if is_body(shape):
            if scale != 1.0:
                for r in shape.text_frame._txBody.iter(qn("a:r")):   # runs without an rPr inherit too: give each one, so sz has a home
                    rpr(r)
                _scale_paragraphs(shape.text_frame.paragraphs, scale,
                                  lambda lvl, s=shape: inherited(s, slide, lvl))
            if vcenter:
                shape.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        elif is_table(shape):
            other = other_size(slide)
            rows = list(shape.table.rows)
            biggest = other
            for row in rows:
                for cell in row.cells:
                    tf = cell.text_frame
                    if scale != 1.0:
                        for r in tf._txBody.iter(qn("a:r")):
                            rpr(r)
                        _scale_paragraphs(tf.paragraphs, scale, lambda lvl: other)
                    for r in tf._txBody.iter(qn("a:rPr")):
                        if r.get("sz"):
                            biggest = max(biggest, int(r.get("sz")))
            if vcenter and rows:
                est = _table_height(shape.table, other, biggest)
                if est < shape.height:
                    shape.top = shape.top + (shape.height - est) // 2
                    shape.height = est
    # after scale=: the box is cut to the text measured at the size it is set in
    if "caution" in words:
        for shape in [sh for sh in slide.shapes if is_body(sh)]:
            caution_box(slide, shape, vcenter)


def _note_parts(slide):
    """A note slide's caption box and the one picture or table it is under, or None."""
    texts = [sh for sh in slide.shapes if is_body(sh)]
    figs = [sh for sh in slide.shapes if is_picture(sh) or is_table(sh)]
    return (texts[0], figs[0]) if len(texts) == 1 and len(figs) == 1 else None


def note_style(slide):
    """note (a `::: note` single-body.lua folded into the caption box): each line a paragraph, the
    note's in the caption size and muted color, centered, as full_slide draws its caption; a
    footnote line keeps style_footnotes' look. Before lay_out, whose scale= then sizes the note
    with the body; False (slide untouched) when the slide is not one picture or table and its
    caption."""
    parts = _note_parts(slide)
    if parts is None:
        return False
    tf = parts[0].text_frame
    split_br(tf._txBody)
    spans = marker_spans(tf._txBody.findall(qn("a:p")))   # a lead folded in after the caption: grid.py's
    for i, p in enumerate(tf.paragraphs):
        if i in spans or style_footnotes.style(p, color=palette(slide)["muted"]):
            continue
        p.alignment = PP_ALIGN.CENTER
        for r in p.runs:   # muted: the dark slide's dark-muted (palette())
            r.font.size, r.font.color.rgb = CAP_SIZE, RGBColor.from_string(palette(slide)["muted"])
    return True


def references_style(slide):
    """references: each body paragraph (a citeproc entry) in the footnote size, hanging by REF_HANG
    em, REF_BEFORE em above it. Before lay_out, whose scale= then sizes it with the rest."""
    hang = em(REF_HANG, REF_PT)
    for shape in [sh for sh in slide.shapes if is_body(sh)]:
        for p in shape.text_frame.paragraphs:
            ppr = p._p.get_or_add_pPr()
            ppr.set("marL", str(hang))
            ppr.set("indent", str(-hang))
            for old in ppr.findall(qn("a:buChar")) + ppr.findall(qn("a:buAutoNum")):
                ppr.remove(old)
            if ppr.find(qn("a:buNone")) is None:
                ppr_place(ppr, etree.Element(qn("a:buNone")))
            p.space_before = Pt(REF_BEFORE * REF_PT)
            for r in p.runs:
                r.font.size = Pt(REF_PT)


def note_place(slide):
    """The note's box moved from the content area's foot (where pandoc puts a caption) to just
    under the picture or table, cut to its text, measured; a picture too tall to leave the note
    room on the slide shrinks (a table cannot, and the note then stays at the foot)."""
    caption, fig = _note_parts(slide)
    floor = caption.top + caption.height
    est = int(text_height(caption, slide) * SLACK)   # the room the note takes, as frame_note sets it
    if is_table(fig):
        bottom = fig.top + table_height(fig, slide)
    else:
        bottom = fig.top + fig.height
        room = floor - est - NOTE_GAP
        if bottom > room and room > fig.top:
            k, middle = (room - fig.top) / fig.height, fig.left + fig.width // 2
            fig.width, fig.height = int(fig.width * k), room - fig.top
            fig.left = middle - fig.width // 2
            bottom = room
    frame_note(caption, slide, caption.left, caption.width, bottom + NOTE_GAP, floor)


def frame_note(note, slide, left, width, top, floor):
    """The note's box at left for width, cut to its text, measured (the slack keeps autofit idle);
    its top at `top`, or higher when the text would then pass `floor`. Returns its top."""
    note.text_frame.vertical_anchor = MSO_ANCHOR.TOP
    est = int(text_height(note, slide) * SLACK)
    y = min(top, floor - est)
    # all four: a placeholder inherits its frame, and setting one leaves the rest at 0
    note.left, note.top, note.width, note.height = left, y, width, est
    return y


def hide_furniture(slide):
    """A full slide's furniture gone: the master's logo, tab, rule and waves hidden (showMasterSp),
    the title deleted (it stays in the source); the slide number kept."""
    slide._element.set("showMasterSp", "0")
    title = slide.shapes.title
    if title is not None:
        title._element.getparent().remove(title._element)


def has_gallery(slide):
    """A body placeholder holding grid-pptx.lua's gallery marker (grid.py draws it, later)."""
    return any(para_text(p).startswith("⟦engine:gallery") for sh in slide.shapes if is_body(sh)
               for p in sh.text_frame._txBody.findall(qn("a:p")))


def full_slide(slide, width, height):
    """One picture over the whole slide; False (slide untouched) for any other body."""
    pics = [sh for sh in slide.shapes if is_picture(sh)]
    texts = [sh for sh in slide.shapes if is_body(sh)]
    if len(pics) != 1 or len(texts) > 1:
        return False
    pic, caption = pics[0], (texts[0] if texts else None)
    hide_furniture(slide)
    pad, floor = int(FULL_PAD * width), int(FULL_FLOOR * height)
    if caption is not None:
        cap = int(FULL_CAP * height)
        caption.left, caption.width = pad, width - 2 * pad
        caption.top, caption.height = height - floor - cap, cap
        tf = caption.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        for p in tf.paragraphs:
            p.alignment = PP_ALIGN.CENTER
            for r in p.runs:
                r.font.size, r.font.color.rgb = CAP_SIZE, CAP_COLOR
        floor += cap
    box_w, box_h = width - 2 * pad, height - pad - floor
    pic.left, pic.top, pic.width, pic.height = contain_box(*pic.image.size, pad, pad, box_w, box_h)
    return True


def _title_graphic(prs, deck_dir):
    """The deck's `titlegraphic:` path from docProps/custom.xml, None without one, or "none"
    when the deck asks for the plain cover (`titlegraphic: none`)."""
    value = docprops.prop(prs, "titlegraphic")
    return "none" if value == "none" else deck_dir / value if value else None


def use_layout(prs, slide, name):
    """Point the slide at the named layout; its placeholders lose the frames pandoc copied from
    the old one and take the new layout's. False when the template has no such layout."""
    layout = next((l for l in prs.slide_layouts if l.name == name), None)
    if layout is None:
        return False
    for rel in slide.part.rels.values():
        if rel.reltype == RT.SLIDE_LAYOUT:
            rel._target = layout.part
            for cached in ("target_part", "target_ref"):   # python-pptx lazyproperties: drop the old layout's
                rel.__dict__.pop(cached, None)
    for sh in slide.placeholders:
        sppr = sh._element.find(qn("p:spPr"))   # a table's graphicFrame has none; its frame stays
        xfrm = sppr.find(qn("a:xfrm")) if sppr is not None else None
        if xfrm is not None:
            sppr.remove(xfrm)
    return True


def _background_photo(slide):
    """The bytes of pandoc's `background-image` (a slide <p:bg> blip), removed; or None."""
    csld = slide._element.find(qn("p:cSld"))
    bg = csld.find(qn("p:bg"))
    blip = bg.find(".//" + qn("a:blip")) if bg is not None else None
    if blip is None:
        return None
    blob = slide.part.related_part(blip.get(qn("r:embed"))).blob
    csld.remove(bg)
    return blob


def _frame(shape, x, y, w, h, width, height):
    """All four of a placeholder's frame, in slide fractions (setting one leaves the rest at 0)."""
    shape.left, shape.top = int(x * width), int(y * height)
    shape.width, shape.height = int(w * width), int(h * height)


def _back_picture(slide, photo, x, y, w, h):
    """The photo cropped to cover the box, at the back of the slide (above the layout's shapes)."""
    tree = slide.shapes._spTree
    pic = fit_picture(slide, photo, x, y, w, h, "cover")
    tree.remove(pic._element)
    tree.insert(2, pic._element)   # after nvGrpSpPr and grpSpPr: the back of the slide
    return pic


def band_photo(slide, width, height, photo):
    """The band title's photo in its box (title-band.photo); nothing else sits there."""
    x, y, w, h = COMP["title-band"]["photo"]["box"]
    _back_picture(slide, photo, int(x * width), int(y * height), int(w * width), int(h * height)).name = "Title photo"


def band_stack(slide, width, height):
    """The band's title stack (date, title, subtitle, rule and author) centered on the column's
    middle, as HTML and PDF center theirs. pandoc puts the subtitle and the author in one
    paragraph of the subtitle placeholder, a blank line between them: the rule (a layout shape,
    moved onto the slide; this deck's Title Slide layout serves the title slide alone) goes in
    that blank line, under a subtitle alone, or over an author alone. From the
    title's and subtitle's text heights, measured in the livery's faces."""
    ph = {sh.placeholder_format.type: sh for sh in slide.placeholders}
    title, date, sub = (ph.get(t) for t in (PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.SUBTITLE))
    if title is None:
        return
    x, y, w, h = COMP["title-band"]["column"]["box"]
    X, W = int(x * width), int(w * width)
    dh = int(0.04 * height) if date is not None else 0
    gaps = [int(g * height) for g in (0.01 if date is not None else 0, 0.015, 0.02)]   # date|title, title|sub, rule
    th = text_height(title, slide, measure.TITLE, 700)   # the master's titles are bold
    text = sub.text_frame.text.strip() if sub is not None else ""
    sh = text_height(sub, slide) if text else 0
    author = slide.part.package.core_properties.author or ""
    head, both, tail = text.partition("\v\v")
    if both:
        rule = "inside"                      # subtitle, blank line, author
    elif text and text == author.strip():
        rule = "over"                        # an author alone
    else:
        rule = "under"                       # a subtitle alone, or nothing
    extra = 2 * gaps[2] if rule != "inside" else 0
    total = dh + gaps[0] + th + gaps[1] + sh + extra
    top = max(int(y * height), int((y + h / 2) * height - total / 2))
    if date is not None:
        date.left, date.top, date.width, date.height = X, top, W, dh
    title.left, title.top, title.width, title.height = X, top + dh + gaps[0], W, th
    sub_top = title.top + th + gaps[1] + (extra if rule == "over" else 0)
    if rule == "over":
        rule_y = sub_top - gaps[2]
    elif rule == "under":
        rule_y = sub_top + sh + gaps[2]
    else:
        tf = sub.text_frame
        p = tf.paragraphs[0]
        parts = pieces(p, measure.BODY, para_size(p, sub, slide))
        cut, last = len(parts), None   # the subtitle ends at the blank line (two breaks)
        for i, (text, _f) in enumerate(parts):
            if text == "\v":
                if last is not None:
                    cut = last
                    break
                last = i
            elif text:
                last = None
        n = measure.runs_lines(parts[:cut], (W - tf.margin_left - tf.margin_right) / measure.PX)   # the subtitle's lines
        size = para_size(p, sub, slide)
        rule_y = sub_top + tf.margin_top + int((n + 0.5) * size * LINE)
    if sub is not None:
        sub.left, sub.top, sub.width = X, sub_top, W
        sub.height = max(sh, int((y + h) * height) - sub.top)
    lay = slide.slide_layout.shapes._spTree
    for el in [e for e in lay if e.tag == qn("p:cxnSp") and e.find(".//" + qn("p:cNvPr")).get("name") == "Title rule"]:
        lay.remove(el)
        el.find(".//" + qn("a:off")).set("y", str(rule_y))
        slide.shapes._spTree.append(el)

def _rgb(role):
    """A role's color for python-pptx's font and fill setters."""
    return RGBColor.from_string(livery_mod.hex(LIVERY, role))


def _box_picture(slide, path, box, width, height, name):
    x, y, w, h = box
    pic = slide.shapes.add_picture(str(path), int(x * width), int(y * height), int(w * width), int(h * height))
    pic.name = name
    return pic


def draw_cover(slide, width, height, photo, comp):
    """A composition with a photo (the title's, or the feature cover): the photo at the back,
    cropped to cover its box, the field and waves over it (pptx_waves.comp_shapes), the mark and
    sign-off in their boxes. Returns the next free shape id."""
    tree = slide.shapes._spTree
    x, y, w, h = comp["photo"]["box"]
    _back_picture(slide, photo, int(x * width), int(y * height), int(w * width), int(h * height)).name = "Photo"
    sid = max(int(e.get("id")) for e in tree.iter(qn("p:cNvPr"))) + 1
    slot = pptx_waves.ROLE_SLOT[comp["field"]["role"]] if comp["field"] else "tx2"   # field: null draws none
    shapes = pptx_waves.comp_shapes(WAVES and WAVES["lines"], comp, sid, width, height, field_slot=slot,
                                    line_slot=pptx_waves.ROLE_SLOT[comp["waves"]["role"]] if comp["waves"] else "accent1")
    for i, xml in enumerate(shapes):
        tree.insert(3 + i, etree.fromstring(xml))
        sid += 1
    for key, name in (("mark", "Logo"), ("signoff", "Sign-off")):
        if comp[key] and comp[key].get("img"):   # its file, the composition's own (livery.json)
            _box_picture(slide, Path(LIVERY["dir"]) / comp[key]["img"],
                         comp[key]["box"], width, height, f"{name} ({comp[key]['role']})")
    return sid


def title_cover(slide, width, height, photo):
    """The title slide over its photo: its composition (title-photo, or -light on Title Slide
    (light)); the layout's tab and rule (hidden by the photo) are copied onto the slide."""
    light = slide.slide_layout.name == "Title Slide (light)"
    draw_cover(slide, width, height, photo, COMP["title-photo-light" if light else "title-photo"])
    tree = slide.shapes._spTree
    for sp in slide.slide_layout.shapes:
        if sp.is_placeholder or is_picture(sp) or sp.name.startswith("Wave "):
            continue
        tree.append(copy.deepcopy(sp._element))


def cover_byline(prs, slide, comp, width, height):
    """A cover's byline (compositions: byline) in the pptx. pandoc puts the subtitle and the authors
    in one paragraph of the subtitle placeholder, a blank line between them, the date in the date
    placeholder, and drops the institute (slide-layout.lua records the authors and institutes).
    With a box, the byline's elements become one text box in it (its order, the type roles' sizes,
    its role, alignment and anchor), and the placeholders keep only the column's elements; without
    one, the byline's runs are recolored where pandoc put them."""
    by = comp.get("byline")
    if not by:
        return
    ph = {sh.placeholder_format.type: sh for sh in slide.placeholders}
    date, sub = ph.get(PP_PLACEHOLDER.DATE), ph.get(PP_PLACEHOLDER.SUBTITLE)
    subtitle = docprops.prop(prs, "subtitle") or ""
    color = RGBColor.from_string(livery_mod.color(LIVERY, by["role"]))
    if not by["box"]:
        _recolor_byline(date, sub, by["elements"], bool(subtitle), color)
        return
    authors = docprops.prop(prs, "engine-author") or prs.core_properties.author or ""
    lines = {"date": [date.text_frame.text.strip()] if date is not None and date.text_frame.text.strip() else [],
             "subtitle": [subtitle] if subtitle else [],
             "author": [a for a in authors.split("\n") if a.strip()],
             "institute": [i for i in (docprops.prop(prs, "engine-institute") or "").split("\n") if i.strip()]}
    x, y, w, h = by["box"]
    box = slide.shapes.add_textbox(int(x * width), int(y * height), int(w * width), int(h * height))
    box.name = "Byline"
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE if by["anchor"] == "middle" else MSO_ANCHOR.TOP
    first = True
    for e in by["elements"]:   # as HTML sets them: the date in bold capitals, subtitle and author bold
        for text in lines[e]:
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            p.alignment = PP_ALIGN.CENTER if by["align"] == "center" else PP_ALIGN.LEFT
            if e == "date" and not first:
                p.space_before = int(0.02 * height)
            r = p.add_run()
            r.text = text
            r.font.size = Pt(livery_mod.type_pt(LIVERY, e))
            r.font.bold = e in ("date", "subtitle", "author")
            r.font.color.rgb = color
            if e == "date":
                r.font._rPr.set("cap", "all")
            first = False
    keep = comp["column"]["elements"]
    if date is not None and "date" not in keep:
        date._element.getparent().remove(date._element)
    if sub is not None:
        parts = ([subtitle] if subtitle and "subtitle" in keep else []) + (["\v".join(lines["author"])] if "author" in keep and lines["author"] else [])
        if parts:
            sub.text_frame.text = "\v\v".join(parts)
        else:
            sub._element.getparent().remove(sub._element)


def cover_fineprint(prs, slide, comp, width, height):
    """A cover's fine print (compositions: fineprint), as HTML and beamer draw it: the deck's
    fineprint: (slide-layout.lua's engine-fineprint, a line per paragraph; `none`, none), else the
    livery's default paragraphs, as one text box in its box, hanging from its top or standing on its
    bottom, left-aligned in its role's color, the fineprint size at its leading. PowerPoint cannot
    flow it, so a text taller than its box is warned."""
    fp = comp.get("fineprint")
    deck = docprops.prop(prs, "engine-fineprint")
    if not fp or not fp.get("box") or deck == "none":
        return
    paras = [p for p in (deck.split("\n") if deck else fp.get("paragraphs") or []) if p.strip()]
    if not paras:
        return
    x, y, w, h = fp["box"]
    box = slide.shapes.add_textbox(int(x * width), int(y * height), int(w * width), int(h * height))
    box.name = "Fineprint"
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.BOTTOM if fp["anchor"] == "bottom" else MSO_ANCHOR.TOP
    pt, lead = livery_mod.type_pt(LIVERY, "fineprint"), LIVERY["type"]["roles"]["fineprint"]["leading"]
    color = RGBColor.from_string(livery_mod.color(LIVERY, fp["role"]))
    for i, text in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.line_spacing = Pt(pt * lead)   # exact points, which type_step scales with the size
        r = p.add_run()
        r.text = text
        r.font.size = Pt(pt)
        r.font.color.rgb = color
    need = sum(para_lines(p, measure.BODY, Pt(pt), box.width) for p in tf.paragraphs) * Pt(pt * lead)
    if need > box.height:
        print(f"slide_layout: WARNING: the cover's fine print needs {need / box.height:.2f} times its box; "
              "it runs past it", file=sys.stderr)


def _recolor_byline(date, sub, els, has_subtitle, color):
    """A box-less byline: its elements' runs in its color, where pandoc put them (the subtitle's
    runs before the blank line, the authors' after it, or all of them without a subtitle)."""
    if date is not None and "date" in els:
        for p in date.text_frame.paragraphs:
            for r in p.runs:
                r.font.color.rgb = color
    if sub is None:
        return
    for p in sub.text_frame.paragraphs:
        part, brs = "subtitle" if has_subtitle else "author", 0
        for el in p._p:
            if el.tag == qn("a:br"):
                brs += 1
                if brs == 2:
                    part = "author"
            elif el.tag == qn("a:r"):
                brs = 0
                if part in els:
                    _Run(el, p).font.color.rgb = color


def closing_top(slide, width, height, comp):
    """A closing column anchored at its top (column.anchor: top): the headline hangs from the
    column's top at its measured height, the body right under it, as HTML and PDF flow them (the
    layout can only split the column at a fixed place). A middle-anchored column is left alone."""
    col = comp["column"]
    if BAND or col["anchor"] != "top":
        return
    title = slide.shapes.title
    body = next((s for s in slide.placeholders if s.placeholder_format.idx == 1), None)
    if title is None:
        return
    x, y, w, h = col["box"]
    X, Y, W, Hh = int(x * width), int(y * height), int(w * width), int(h * height)
    title.left, title.top, title.width, title.height = X, Y, W, Hh   # all four: a placeholder inherits its frame
    size = para_size(title.text_frame.paragraphs[0], title, slide) if title.text_frame.paragraphs else Pt(36)
    title.height = min(Hh, text_height(title, slide, measure.TITLE, 700))
    if body is not None:
        top = title.top + title.height + int(0.4 * size)   # the h2's margin-bottom, 0.4em
        body.left, body.top, body.width, body.height = X, top, W, max(Pt(12), Y + Hh - top)


def _left(shape):
    """A feature title takes its composition column's alignment (column.align), whatever
    frametitle.align gives the content titles (the layout's title style); a no-op when they agree."""
    want = COMP["feature"]["column"].get("align", "left")
    if LIVERY.get("frametitle", {}).get("align", "left") != want:
        for p in shape.text_frame.paragraphs:
            p.alignment = PP_ALIGN.CENTER if want == "center" else PP_ALIGN.LEFT


def _recolor(shape, color, pt=None):
    """Every run of the shape in a color (RGBColor), and at pt points when given."""
    for p in shape.text_frame.paragraphs:
        for r in p.runs:
            r.font.color.rgb = color
            if pt:
                r.font.size = Pt(pt)


def _rect(slide, box, width, height, hexv, alpha=None, name=None):
    """A filled rectangle at a box (slide fractions); alpha (0-1) written as <a:alpha>, which
    python-pptx lacks."""
    x, y, w, h = box
    r = rect(slide, int(x * width), int(y * height), int(w * width), int(h * height), hexv, name)
    if alpha is not None:
        clr = r.fill._xPr.find(qn("a:solidFill"))[0]
        etree.SubElement(clr, qn("a:alpha")).set("val", str(round(alpha * 100000)))
    return r


def draw_feature(slide, width, height, photo):
    """A {.feature} slide by feature.style: bar (the photo full, the title in the column on the
    dark bar, the body on the panel cut to its text) or cover (draw_cover, the title, rule and
    body in the column). No master shapes either way."""
    c = COMP["feature"]
    slide._element.set("showMasterSp", "0")
    title, body = slide.shapes.title, [sh for sh in slide.shapes if is_body(sh)]
    x, y, w, h = c["column"]["box"]
    text = _rgb(c["column"]["role"])
    if c["bar"]:
        _back_picture(slide, photo, 0, 0, width, height).name = "Photo"
        bar = _rect(slide, c["bar"]["box"], width, height, livery_mod.hex(LIVERY, c["bar"]["role"]), c["bar"]["opacity"], "Feature bar")
        tree = slide.shapes._spTree
        tree.remove(bar._element); tree.insert(3, bar._element)   # over the photo, under the text
        for sh in slide.placeholders:   # the number on the photo: the on-dark role, as HTML's and the PDF's
            if sh.placeholder_format.type == PP_PLACEHOLDER.SLIDE_NUMBER:
                for para in sh.text_frame.paragraphs:
                    para.font.color.rgb = _rgb("on-dark")
        if title is not None:
            _frame(title, x, y, w, h, width, height)
            title.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE if c["column"]["anchor"] == "middle" else MSO_ANCHOR.TOP
            _recolor(title, text)
            _left(title)
        px, py, pw, ph = c["panel"]["box"]
        for shape in body:
            _frame(shape, px, py, pw, ph, width, height)
            vpad, hpad = pads("panel")
            inset(shape.text_frame, vpad, hpad)
            est = min(int(text_height(shape, slide) * SLACK), int(ph * height))
            panel = _rect(slide, [px, py, pw, est / height], width, height, livery_mod.hex(LIVERY, c["panel"]["role"]), c["panel"]["opacity"], "Feature panel")
            tree = slide.shapes._spTree
            tree.remove(panel._element); tree.insert(list(tree).index(shape._element), panel._element)
            shape.height = est
            _recolor(shape, _rgb("on-dark"))   # the panel's text: the on-dark role, as HTML's and the PDF's
    else:
        draw_cover(slide, width, height, photo, c)
        if c["column"]["tab"]:
            _rect(slide, [0, y, GEOM["tab-width"], 0.16], width, height, livery_mod.hex(LIVERY, "accent"), name="Title tab")
        if c["column"]["rule"]:
            _rect(slide, [x, y + 0.185, c["column"]["rule"][0], c["column"]["rule"][1]], width, height,
                  livery_mod.hex(LIVERY, "secondary"), name="Title rule")
        if title is not None:
            _frame(title, x, y, w, 0.16, width, height)
            title.text_frame.vertical_anchor = MSO_ANCHOR.TOP
            _recolor(title, text, livery_mod.type_pt(LIVERY, "frametitle"))
            _left(title)
        for shape in body:
            _frame(shape, x, y + 0.21, w, h - 0.21, width, height)
            _recolor(shape, text, livery_mod.type_pt(LIVERY, "body"))


def draw_dark(slide, width, height, words):
    """A {.brand-dark} slide on a light livery: the slide on the dark role (its background a solid
    fill, common.is_dark()), the master's furniture hidden (showMasterSp: no content art, no waves)
    and redrawn from livery.json furniture: the tab in the accent ("Dark tab"), the hairline in the
    dark-rule blend ("Dark rule"), the dark role's mark and sign-off in their corners ("Dark logo",
    "Dark signoff"), at the back; the title and the slide number on-dark, and every body
    placeholder's text that takes the master's color (on_dark()), but a grid slide's marker spans
    (grid.py draws those lines, a card's or a lead's on its own fill) and a caution slide's body
    (on its own field). The steps after this read the swapped colors from common.palette()."""
    slide._element.set("showMasterSp", "0")
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = _rgb("dark")
    fu = LIVERY["furniture"]
    made = []
    if fu["tab"]:
        made.append(_rect(slide, fu["tab"], width, height, livery_mod.hex(LIVERY, "accent"), name="Dark tab"))
    if fu["rule"]:   # the line along y, h its width
        x, y, w, h = fu["rule"]
        made.append(_rect(slide, [x, y - h / 2, w, h], width, height, livery_mod.hex(LIVERY, "dark-rule"), name="Dark rule"))
    for corner in sorted(c for c, mk in fu["marks"].items() if mk):
        mk = fu["marks"][corner]
        path = livery_mod.mark(LIVERY, "dark", mk["mark"])
        if path:
            made.append(_box_picture(slide, Path(LIVERY["dir"]) / path, mk["box"], width, height, f"Dark {mk['mark']}"))
    tree = slide.shapes._spTree
    for sh in reversed(made):   # at the back, after the tree's own properties
        tree.remove(sh._element)
        tree.insert(2, sh._element)
    on = livery_mod.hex(LIVERY, "on-dark")
    for sh in slide.placeholders:
        kind = sh.placeholder_format.type
        if kind == PP_PLACEHOLDER.SLIDE_NUMBER:
            for para in sh.text_frame.paragraphs:
                para.font.color.rgb = _rgb("on-dark")
        elif kind in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE):
            _recolor(sh, _rgb("on-dark"))
        elif is_body(sh) and "caution" not in words:
            paras = sh.text_frame._txBody.findall(qn("a:p"))
            on_dark(paras, on, marker_spans(paras))


def add_slide_number(slide):
    """The layout's slide-number placeholder, copied onto the slide (pandoc leaves it out)."""
    for sh in slide.placeholders:
        if sh.placeholder_format.type == PP_PLACEHOLDER.SLIDE_NUMBER:
            return False
    for sh in slide.slide_layout.placeholders:
        if sh.placeholder_format.type == PP_PLACEHOLDER.SLIDE_NUMBER:
            el = copy.deepcopy(sh._element)
            for e in el.iter(qn("p:cNvPr")):
                e.set("id", str(max(int(c.get("id")) for c in slide.shapes._spTree.iter(qn("p:cNvPr"))) + 1))
            slide.shapes._spTree.append(el)
            return True
    return False


def run(prs, path: Path) -> None:
    """The open deck, changed in place; path the pptx's (a titlegraphic: is relative to its directory)."""
    w, h = prs.slide_width, prs.slide_height
    done = 0
    light = docprops.prop(prs, "engine-title") == "light"
    if light and BAND:
        print("slide_layout: WARNING: themeoptions title=light is ignored under the livery's band title", file=sys.stderr)
        light = False
    if light and len(prs.slides) and prs.slides[0].slide_layout.name == "Title Slide":
        if not use_layout(prs, prs.slides[0], "Title Slide (light)"):
            print("slide_layout: WARNING: title=light, but the template has no Title Slide (light)", file=sys.stderr)
    graphic = _title_graphic(prs, path.parent)
    default = LIVERY.get("art", {}).get("titlegraphic")
    if graphic is None and default and not BAND:   # the band's layout carries the default itself
        graphic = Path(LIVERY["dir"]) / default
    if graphic == "none":   # the plain cover, not the livery's default photo
        graphic = None
    if graphic is not None and len(prs.slides) and prs.slides[0].slide_layout.name.startswith("Title Slide"):
        if graphic.exists():
            (band_photo if BAND else title_cover)(prs.slides[0], w, h, graphic.read_bytes())
            done += 1
        else:
            print(f"slide_layout: WARNING: titlegraphic {graphic} not found", file=sys.stderr)
    markers = [take_marker(slide, MARKER) for slide in prs.slides]
    for slide, words in zip(prs.slides, markers):   # before numbering: the number is the new layout's
        if words and "closing" in words:
            if use_layout(prs, slide, "Closing (light)" if light else "Closing"):
                words.remove("closing")
                closing_top(slide, w, h, COMP["closing-light" if light else "closing"])
            else:
                print("slide_layout: WARNING: {.closing} slide, but the template has no Closing layout", file=sys.stderr)
    numbered = sum(add_slide_number(slide) for slide in prs.slides)
    for slide, words in zip(prs.slides, markers):
        if words is None:
            continue
        if not words or words == ["closing"]:   # a closing slide with nothing else to do
            done += 1
            continue
        if "feature" in words:
            photo = _background_photo(slide)
            if photo is None:
                print("slide_layout: WARNING: {.feature} slide without a background-image; left as it is",
                      file=sys.stderr)
            else:
                draw_feature(slide, w, h, photo)
        elif "full" in words:   # its caption, a note's lines too, is the full caption
            if has_gallery(slide):   # grid.py fills the slide with the gallery and its caption
                hide_furniture(slide)
            elif not full_slide(slide, w, h):
                print("slide_layout: WARNING: {.full} slide %r needs one picture and at most a "
                      "caption; left as it is" % slide_title(slide), file=sys.stderr)
        else:
            if "dark" in words and LIVERY["mode"] != "dark":   # before the note's style, which reads palette()
                draw_dark(slide, w, h, words)
            note = "note" in words and note_style(slide)
            if "references" in words:
                references_style(slide)
            lay_out(slide, words, w)
            if note:   # after scale=: the box is cut to the note at the size it is set in
                note_place(slide)
        done += 1
    if BAND and len(prs.slides) and prs.slides[0].slide_layout.name == "Title Slide":   # after any re-framing
        band_stack(prs.slides[0], w, h)
    if len(prs.slides) and prs.slides[0].slide_layout.name.startswith("Title Slide"):   # the cover's byline
        cover = "title-band" if BAND else ("title-photo" if graphic is not None else "title") + ("-light" if light else "")
        cover_byline(prs, prs.slides[0], COMP[cover], w, h)
        cover_fineprint(prs, prs.slides[0], COMP[cover], w, h)
    print(f"slide_layout: laid out {done} slide(s), numbered {numbered} → {path.name}")
