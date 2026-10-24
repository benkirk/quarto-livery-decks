"""A grid slide (three or more columns, a lead box, cards, strips, quotes, steps, a timeline, an agenda, a
gallery) redrawn from grid-pptx.lua's markers.

grid-pptx.lua flattens a grid slide's columns into one text flow, which pandoc writes into the
Title and Content body placeholder, with marker paragraphs (a run of `⟦engine:...⟧`):
`⟦engine:columns widths=a,b,c⟧`, then per column `⟦engine:column⟧` ... `⟦engine:end⟧`, and a last
`⟦engine:end⟧` closing the columns; `⟦engine:lead⟧` ... `⟦engine:end⟧` a lead box;
`⟦engine:rail kind=steps|timeline role=r⟧`, per item `⟦engine:item⟧` (`⟦engine:item done⟧`, a done
milestone) ... `⟦engine:end⟧`, and a last `⟦engine:end⟧`, steps or a timeline (draw_rail());
`⟦engine:quote role=r⟧` ... `⟦engine:end⟧` a quote (draw_quote(): its lines in the quote size, an
`⟦engine:attribution⟧` line at the caption size, muted, beside a bar; on its own, in a card inside
its pads, or in a plain column between its text boxes); `⟦engine:agenda current=n⟧`, per numbered
section `⟦engine:item n=k state=..⟧` ... `⟦engine:end⟧`, and a last `⟦engine:end⟧`, the agenda
(draw_agenda(): a number box beside a text box per row); `⟦engine:gallery cols=c rows=r fit=f ..⟧`, per
image `⟦engine:item src=..⟧`, its caption, `⟦engine:end⟧`, and a last `⟦engine:end⟧`, an image gallery
(gallery.py's draw_gallery(), which reads the files from the deck's directory); in a column or a card,
`⟦engine:picture src=..[ data=..][ size=..]⟧`, its caption, `⟦engine:end⟧`, a picture (figures.py's
draw_picture(): the column's width, its aspect kept, its caption under it; "Column n picture",
"Card n picture", a card's inside its pads and its group; pictures give way, never below PIC_FLOOR
em, when the text beside them needs the room); in a plain column, `⟦engine:table cols=n[ header]⟧`,
per row `⟦engine:row⟧` and its cells led by `⟦engine:cell⟧`, its caption, `⟦engine:end⟧`, a native
table (figures.py's draw_table(), "Column n table", which style_tables styles after). This
parses the placeholder's paragraphs into segments, text and columns, and lays them top down from the
placeholder's frame (its top below the title, its foot geometry.body-bottom; its left and width,
so a {.wide} slide's widening by slide_layout carries over): text before the columns stays in the
placeholder, cut to its measured height; the columns are free text boxes (Column 1, Column 2, ...; a card, a lead, a strip or a quote is one group of
its shapes, "Card n", "Lead", "Strip", "Quote", so a click in PowerPoint selects it whole)
at columns.py's frames (each its share of the box less boxes.columns.gap, a livery's
boxes.columns.rule in each gap), their list style the master's body style, so pandoc's paragraphs,
moved into them with their own properties, keep the template's bullets and sizes; a lead is its
text box on a fill (boxes.lead.fill) with a bar at its left (bar, bar-color), padded (pad), cut to
its text, a pad's gap under it; text after the columns is a box of its own. Marker paragraphs are
deleted; a placeholder left with no text goes. After columns, before style_tables and type_step
(which scales the boxes' sizes with the rest). A grid slide's body-layout classes (.caution,
.vcenter) do not apply; scale= has run before.
"""

import copy
import re
import sys
from pathlib import Path

from lxml import etree
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Pt

import livery as livery_mod

from .columns import rules
from .figures import draw_picture, draw_table, place_picture, shrink_picture
from .gallery import draw_gallery
from .common import (BODY_PT, BOXES, GEOM, LIVERY, MARK, SLACK, body_box, body_frame, column_frames, em, fit_table, group, inset,
                     is_body, is_picture, is_table, marker_spans, pads, para_text, rect, rpr, shape, slide_title, space_before,
                     split_br, style_runs)
from .slide_layout import NOTE_GAP, frame_note, note_style, table_height, use_layout
from .measure import text_height

PREFIX = re.compile(r"^⟦engine:(title|sub|label|big|attribution)⟧")   # a card's lines, a quote's attribution (grid-pptx.lua)
CARD_GAP = 0.25   # a card's paragraphs sit this share of the body's spacing apart (beamer's card \parskip)
BODY_GAP = livery_mod.SPC_BEFORE[0] * BODY_PT * CARD_GAP   # pt: CARD_GAP of the template's space before a body paragraph
PIC_FLOOR = 2   # em of the body type: the least height a picture keeps when the lead under it needs its room


def _attrs(rest):
    """`widths=a,b role=x pick` -> {"widths": "a,b", "role": "x", "pick": True}."""
    out = {}
    for w in (rest or "").split():
        k, eq, v = w.partition("=")
        out[k] = v if eq else True
    return out


def parse(paragraphs):
    """The flow as segments: ("text", [p]), ("lead", attrs, [p]), ("strip" or "quote", attrs, [p]),
    ("columns", shares, [(attrs, [p])]; a strip or quote in a column one of its items),
    (kind, attrs, [("item", attrs, [p])]) for a rail, an agenda or a gallery; None when it holds no
    marker of a box (not a grid slide). Markers nest: a column inside its columns, an item inside its
    rail, agenda or gallery, each closed by its `end`."""
    segs, stack = [], []   # stack: the open containers, innermost last

    def text(p):
        if segs and segs[-1][0] == "text":
            segs[-1][1].append(p)
        else:
            segs.append(("text", [p]))
    for p in paragraphs:
        m = MARK.match(para_text(p).strip())
        if m is None:
            top = stack[-1] if stack else None
            if top is None or top[0] in ("columns", "rail", "agenda", "gallery"):   # a stray paragraph between columns or items: text
                text(p)
            else:
                top[-1].append(p)
            continue
        kind, attrs = m.group(1), _attrs(m.group(2))
        if kind == "columns":
            seg = ("columns", [float(v) for v in attrs.get("widths", "").split(",") if v], [])
            segs.append(seg); stack.append(seg)
        elif kind == "column" and stack and stack[-1][0] == "columns":
            col = ("column", attrs, [])
            stack[-1][2].append(col); stack.append(col)
        elif kind == "rail" and not stack:
            seg = ("rail", attrs, [])
            segs.append(seg); stack.append(seg)
        elif kind in ("agenda", "gallery") and not stack:
            seg = (kind, attrs, [])
            segs.append(seg); stack.append(seg)
        elif kind == "item" and stack and stack[-1][0] in ("rail", "agenda", "gallery"):
            item = ("item", attrs, [])
            stack[-1][2].append(item); stack.append(item)
        elif kind == "lead" and not stack:
            seg = ("lead", attrs, [])
            segs.append(seg); stack.append(seg)
        elif kind in ("strip", "quote", "picture", "table"):   # in a column, an item of its flow; else a segment of its own
            seg = (kind, attrs, [])
            if stack and stack[-1][0] == "column":
                stack[-1][2].append(seg)
            elif not stack and kind not in ("picture", "table"):   # a column's alone (grid-pptx.lua's cell)
                segs.append(seg)
            stack.append(seg)
        elif kind in ("row", "cell") and stack and stack[-1][0] == "table":   # a row's opener, an empty cell: draw_table's
            stack[-1][2].append(p)
        elif kind == "end" and stack:
            stack.pop()
    return segs if any(s[0] != "text" for s in segs) else None


def draw_lead(slide, x, y, w, room, paras, slide_height):
    """The lead box: its text on boxes.lead.fill, a bar of bar-color at its left, the text padded
    by boxes.lead.pad, the box cut to its measured text; the height it takes, a pad's gap under it
    included (HTML's margin-bottom)."""
    lead = BOXES["lead"]
    bar = round(lead["bar"] * slide_height)
    vpad, hpad = pads("lead")
    tb = body_box(slide, x + bar, y, w - bar, room, paras, "Lead text")
    inset(tb.text_frame, vpad, hpad)
    h = min(int(text_height(tb, slide) * SLACK), room)
    tb.height = h
    fill = rect(slide, x, y, w, h, _hex(lead["fill"]), "Lead fill", behind=tb)
    bar_shape = rect(slide, x, y, bar, h, _hex(lead["bar-color"]), "Lead bar", behind=tb)
    group(slide, [fill, bar_shape, tb], "Lead")
    return h + vpad


def _hex(value):
    """A role's, or a box color's (#RRGGBB), six digits."""
    return livery_mod.hex(LIVERY, value)


def _prefix(p):
    """A paragraph's prefix marker (title, sub, label, big), taken off its first run; None if none."""
    t = p.find(".//" + qn("a:t"))
    m = PREFIX.match(t.text or "") if t is not None else None
    if m is None:
        return None
    t.text = t.text[m.end():]
    if not t.text:   # the marker's own run
        r = t.getparent()
        if r.getnext() is not None:
            r.getparent().remove(r)
    return m.group(1)


def _close(p):
    """A card paragraph's space before: CARD_GAP of the master body's (a:spcBef), as beamer's card
    lines sit close (HTML uses reveal's block margin); type_step scales it with the sizes."""
    space_before(p, BODY_GAP)


def _card_height(tb, slide):
    return int(text_height(tb, slide, gap=Pt(BODY_GAP)))


def _card_text(paras, solid):
    """A card's paragraphs, their prefix markers taken off and styled: (head, body) lists. The band's
    title in the subhead size, bold, and its sub in the body size, both primary-text; a label in the
    eyebrow size, caps, muted; a big number in the display size, bold, primary-contrast; on a solid
    card the body in primary-text."""
    head, body = [], []
    on_role = _hex("primary-text")
    for p in paras:
        if isinstance(p, tuple):   # a strip: its own lines styled as a card's body; a quote, a picture: their drawers'
            body.append(p if p[0] in ("quote", "picture") else (p[0], p[1], _card_text(p[2], solid)[1]))   # a table: never (tabled_row())
            continue
        kind = _prefix(p)
        _close(p)
        if kind == "title":
            style_runs(p, livery_mod.type_pt(LIVERY, "subhead"), True, on_role); head.append(p)
        elif kind == "sub":
            style_runs(p, livery_mod.type_pt(LIVERY, "body"), False, on_role); head.append(p)
        elif kind == "label":
            style_runs(p, livery_mod.type_pt(LIVERY, "eyebrow"), True, on_role if solid else _hex("muted"), caps=True); body.append(p)
        elif kind == "big":
            style_runs(p, livery_mod.type_pt(LIVERY, "display"), True, on_role if solid else _hex("primary-contrast")); body.append(p)
        else:
            if solid:
                style_runs(p, color=on_role)
            body.append(p)
    return head, body


def _flat(items):
    """A column's flow without its strips' boxes (a plain column draws no strip): paragraphs."""
    out = []
    for it in items:
        out.extend(it[2] if isinstance(it, tuple) else [it])
    return out


def draw_strip(slide, x, y, w, room, paras, role, solid, name):
    """A strip: its text box padded by boxes.strip.pad over a rectangle the strip.tint of its role
    over the background (on a solid card the role itself, its text primary-text), cut to its text.
    (text box, rectangle) at y, its measured height."""
    vpad, hpad = pads("strip")
    tb = body_box(slide, x, y, w, room, paras, f"{name} text")
    inset(tb.text_frame, vpad, hpad)
    if solid:
        for p in paras:
            style_runs(p, color=_hex("primary-text"))
    tb.height = min(_card_height(tb, slide), room)
    fill = _hex(role) if solid else _hex(livery_mod.mix(_hex(role), _hex("background"), BOXES["strip"]["tint"]))
    return tb, rect(slide, x, y, w, tb.height, fill, name, behind=tb)


# a:rPr's children before a:latin, in the schema's order (a fill and a line among them)
_BEFORE_LATIN = {qn(f"a:{t}") for t in ("ln", "noFill", "solidFill", "gradFill", "blipFill", "pattFill", "grpFill",
                                         "effectLst", "effectDag", "highlight", "uLnTx", "uLn", "uFillTx", "uFill")}


def _latin(el, face):
    """A run's a:rPr set in a typeface: its a:latin, made in the schema's place when it has none."""
    lat = el.find(qn("a:latin"))
    if lat is None:
        lat = etree.Element(qn("a:latin"))
        before = [c for c in el if c.tag in _BEFORE_LATIN]
        before[-1].addnext(lat) if before else el.insert(0, lat)
    lat.set("typeface", face)


QUOTE_FACE = LIVERY["pptx"].get("quote-typeface") or BOXES["quote"]["family"]   # None: the body's


def draw_quote(slide, x, y, w, room, paras, color, slide_height, name="Quote"):
    """A quote: its lines in the quote size, in boxes.quote.family (PowerPoint's pptx.quote-typeface,
    else the family's name; none: the body's) and boxes.quote.style, the attribution (led by
    ⟦engine:attribution⟧) at the caption size in muted after "— ", in the body's face, upright; in a
    text box padded by boxes.quote.pad beside a bar of boxes.quote.bar of the color at its left, no
    fill, cut to its text (measured in the family's files). (text box, bar) at y."""
    bar = round(BOXES["quote"]["bar"] * slide_height)
    vpad, hpad = pads("quote")
    italic = BOXES["quote"]["style"] == "italic"
    for p in paras:
        if _prefix(p) == "attribution":
            t = p.find(".//" + qn("a:t"))
            if t is not None:
                t.text = "— " + (t.text or "")
            caption_pt = livery_mod.type_pt(LIVERY, "caption")
            style_runs(p, caption_pt, color=_hex("muted"))
            space_before(p, BOXES["quote"]["attribution-gap"] * caption_pt)   # under the lines, as HTML and the PDF
        else:
            style_runs(p, livery_mod.type_pt(LIVERY, "quote"))
            for r in p.iter(qn("a:r")):
                el = rpr(r)
                if italic:
                    el.set("i", "1")
                if QUOTE_FACE:
                    _latin(el, QUOTE_FACE)
    tb = body_box(slide, x + bar, y, max(w - bar, 0), room, paras, f"{name} text")
    inset(tb.text_frame, vpad, hpad)
    tb.height = min(int(text_height(tb, slide) * SLACK), room)
    return tb, rect(slide, x, y, bar, tb.height, color, f"{name} bar", behind=tb)


def _quote_color(attrs, slide):
    """A quote marker's role= as a hex when it is a role or a blend of the livery, else (none, or a
    name that is neither, warned of by _pick_role) the box's bar-color."""
    return _hex(_pick_role(attrs, BOXES["quote"]["bar-color"], slide))


AGENDA_COLORS = {"done": _hex("muted"), "current": _hex(BOXES["agenda"]["current-color"]), "next": None}   # None: the text's


def draw_agenda(slide, x, y, w, room, items):
    """The agenda's rows, top down from y, boxes.agenda.gap above each: a number box number-width wide (the
    eyebrow size, bold, boxes.agenda.number-color) beside a text box, the title in the subhead size
    (done muted, current bold in boxes.agenda.current-color, next the text's) and the divider's
    subtitle under it (the caption size, muted); each cut to its text, measured. When the rows run
    past the room, the subtitles go (the room is what a slide has: no autofit on native boxes).
    One group, "Agenda", of "Agenda k number" and "Agenda k text". The height it takes."""
    gap, col = em(BOXES["agenda"]["gap"]), em(BOXES["agenda"]["number-width"])
    rows = []
    for k, (_, attrs, paras) in enumerate(items, 1):
        title, sub = [], []
        for p in paras:
            (sub if _prefix(p) == "sub" else title).append(p)
        state = attrs.get("state", "next")
        for p in title:
            style_runs(p, livery_mod.type_pt(LIVERY, "subhead"), state == "current", AGENDA_COLORS.get(state))
            space_before(p, 0)
        for p in sub:
            style_runs(p, livery_mod.type_pt(LIVERY, "caption"), color=_hex("muted"))
            _close(p)
        num = copy.deepcopy(title[0]) if title else etree.Element(qn("a:p"))
        runs = num.findall(qn("a:r"))
        for r in runs[1:]:
            num.remove(r)
        if runs:
            runs[0].find(qn("a:t")).text = attrs.get("n", str(k))
        style_runs(num, livery_mod.type_pt(LIVERY, "eyebrow"), True, _hex(BOXES["agenda"]["number-color"]))
        rows.append((k, num, title, sub))

    def lay(with_subs):
        shapes, cy = [], y
        for k, num, title, sub in rows:
            cy += gap
            tb = body_box(slide, x + col, cy, max(w - col, 0), room, title + (sub if with_subs else []), f"Agenda {k} text")
            tb.text_frame.margin_top = tb.text_frame.margin_bottom = 0   # the rows apart by the gap alone, as HTML's
            th = int(text_height(tb, slide))
            tb.height = int(th * SLACK)   # the box with its slack (no fill: an overlap does not show), the row its text
            nb = body_box(slide, x, cy, col, tb.height, [num], f"Agenda {k} number")
            # the number on the title's first line: down by the difference of the two sizes' ascents
            nb.text_frame.margin_top = em(0.8, livery_mod.type_pt(LIVERY, "subhead") - livery_mod.type_pt(LIVERY, "eyebrow"))
            nb.text_frame.margin_bottom = 0
            shapes += [nb, tb]
            cy += th
        return shapes, cy - y
    shapes, h = lay(True)
    if h > room and any(sub for *_, sub in rows):   # past the room: drop the subtitles and lay out again
        for sh in shapes:
            sh._element.getparent().remove(sh._element)
        print(f"grid: WARNING: slide {slide_title(slide)!r}: the agenda runs past the slide; its subtitles dropped",
              file=sys.stderr)
        shapes, h = lay(False)
    if shapes:
        group(slide, shapes, "Agenda")
    return min(h, room)


def _yield_pictures(pics, over, slide):
    """The drawn pictures [(picture, caption box)] shorter by over (EMU) in all, each by its share of
    their heights and to no less than PIC_FLOOR em: the room text beside them needs. The EMU taken."""
    total = sum(pic.height for pic, _ in pics)
    if over <= 0 or not total:
        return 0
    taken = 0
    for pic, cb in pics:
        by = min(round(over * pic.height / total), max(pic.height - em(PIC_FLOOR), 0))
        shrink_picture(pic, cb, by)
        taken += by
    if taken < over:
        print(f"grid: WARNING: slide {slide_title(slide)!r}: a column's pictures and text need more room than it has",
              file=sys.stderr)
    return taken


def _plain_column(slide, x, y, w, room, items, name, slide_height, deck_dir):
    """A plain column's flow (a strip's lines as text: a plain column draws no strip) as its parts,
    [(kind, shape, other)]: one text box, or, when it holds a quote, a picture or a table, its text runs,
    its quotes (draw_quote(), at the column's width: (quote, text box, bar)), its pictures
    (draw_picture(): (picture, picture, caption box or None)) and its tables (draw_table(): (table,
    frame, caption box or None)) in order, each at y, the pictures shortened when the column passes
    the room; _place() stacks them."""
    if not any(isinstance(it, tuple) and it[0] in ("quote", "picture", "table") for it in items):
        return [("text", body_box(slide, x, y, w, room, _flat(items), name), None)]
    parts, run = [], []

    def text():
        nonlocal run
        if run:
            parts.append(("text", body_box(slide, x, y, w, room, run, name if not parts else f"{name} {len(parts) + 1}"), None))
        run = []
    for it in items:
        if isinstance(it, tuple) and it[0] == "quote":
            text()
            tb, bar = draw_quote(slide, x, y, w, room, it[2], _quote_color(it[1], slide), slide_height,
                                 f"{name} quote {sum(1 for p in parts if p[0] == 'quote') + 1}")
            parts.append(("quote", tb, bar))
        elif isinstance(it, tuple) and it[0] == "picture":
            text()
            k = sum(1 for p in parts if p[0] == "picture") + 1
            drawn = draw_picture(slide, x, y, w, room, it[1], it[2], deck_dir, f"{name} picture" + (f" {k}" if k > 1 else ""))
            if drawn is None:   # not drawn (warned of): its caption stays text
                run.extend(it[2])
                continue
            parts.append(("picture", drawn[0], drawn[1]))
        elif isinstance(it, tuple) and it[0] == "table":
            text()
            k = sum(1 for p in parts if p[0] == "table") + 1
            drawn = draw_table(slide, x, y, w, room, it[1], it[2], f"{name} table" + (f" {k}" if k > 1 else ""))
            if drawn is not None:
                parts.append(("table", drawn[0], drawn[1]))
        else:
            run.extend(it[2] if isinstance(it, tuple) else [it])
    text()
    _yield_pictures([(p, cb) for kind, p, cb in parts if kind == "picture"], int(_column_height(parts, slide) * SLACK) - room, slide)
    return parts


def _part_height(kind, sh, other, slide):
    """A plain column part's measured height: a text run's text, a quote's box, a picture with its caption."""
    if kind == "quote":
        return sh.height
    if kind in ("picture", "table"):
        return sh.height + (other.height if other is not None else 0)
    return int(text_height(sh, slide))


def _column_height(parts, slide):
    """A plain column's measured height: its parts' heights added."""
    return sum(_part_height(kind, sh, other, slide) for kind, sh, other in parts)


def _place(slide, parts, y, h):
    """A plain column's parts stacked from y, each text run its measured height and the last running
    to y + h (one text box: at y, h tall, as before quotes); a quote grouped once in place ("<name>"),
    a picture or a table with its caption under it."""
    cy = y
    for i, (kind, tb, bar) in enumerate(parts):
        if kind == "quote":
            tb.top = bar.top = cy
            group(slide, [bar, tb], tb.name[:-len(" text")])
            cy += tb.height
            continue
        if kind in ("picture", "table"):
            place_picture(tb, bar, cy)
            cy += _part_height(kind, tb, bar, slide)
            continue
        th = max(y + h - cy, 0) if i == len(parts) - 1 else int(text_height(tb, slide) * SLACK)
        tb.top, tb.height = cy, th
        cy += th


def _pick_role(attrs, default, slide):
    """A marker's role= when it is a role or a blend of the livery, else the default, warned of
    (engine-common.lua's box_role, for HTML and PDF)."""
    r = attrs.get("role", default)
    if r == default or r in LIVERY["roles"] or r in LIVERY["blends"]:
        return r
    print(f"grid: WARNING: slide {slide_title(slide)!r}: role={r!r} is not a role or a blend of the livery; {default}",
          file=sys.stderr)
    return default


ROLES = BOXES["card"]["roles"]   # a card without role=: by its place in the row


def draw_cards(slide, columns, frames, y, room, n_col, slide_height, deck_dir):
    """A row of cards (and plain columns beside them): each card a head band in its role over a
    body on boxes.card.fill (all of it the role's on a .solid card), rounded by card.radius, a
    .pick outlined in outline-color; its body a stack of text and strips (each at the card's full
    width), quotes and pictures (inside its pads, a picture a body pad below what is above it, shortened
    when the card would pass the room); each head its own text's height (as HTML's) and the cards one height, the row's
    tallest, capped at the room left. Returns the row's height and the column boxes' count."""
    card = BOXES["card"]
    radius = round(card["radius"] * slide_height)
    hv, hh = pads("card", "head-pad")
    bv, bh = pads("card", "pad")
    cards = []
    for k, ((_, attrs, paras), (x, w)) in enumerate(zip(columns, frames), 1):
        n_col += 1
        if not attrs.get("card"):
            cards.append(dict(plain=_plain_column(slide, x, y, w, room, paras, f"Column {n_col}", slide_height, deck_dir), attrs=attrs))
            continue
        solid = attrs.get("solid")
        head, items = _card_text(paras, solid)
        hb = body_box(slide, x, y, w, room, head, f"Card {n_col} title") if head else None
        if hb is not None:
            inset(hb.text_frame, hv, hh)
        parts, run, n_strip, n_quote, n_pic = [], [], 0, 0, 0   # the body: text boxes, strips, quotes and pictures, in order
        def text_part():
            nonlocal run
            if run or not parts:
                tb = body_box(slide, x, y, w, room, run, f"Card {n_col} body" + (f" {len(parts) + 1}" if parts else ""))
                inset(tb.text_frame, bv, bh)
                parts.append(("text", tb, None, _card_height(tb, slide)))
            run = []
        for it in items:
            if isinstance(it, tuple) and it[0] == "quote":   # inside the body's pads, as HTML's and the PDF's
                text_part() if run else None
                n_quote += 1
                tb, r = draw_quote(slide, x + bh, y, w - 2 * bh, room, it[2], _quote_color(it[1], slide), slide_height,
                                   f"Card {n_col} quote {n_quote}")
                parts.append(("quote", tb, r, tb.height))
            elif isinstance(it, tuple) and it[0] == "picture":   # inside the body's pads, a pad below what is above
                drawn = draw_picture(slide, x + bh, y, w - 2 * bh, max(room - bv, 0), it[1], it[2], deck_dir,
                                     f"Card {n_col} picture" + (f" {n_pic + 1}" if n_pic else ""))
                if drawn is None:   # not drawn (warned of): its caption stays text
                    run.extend(it[2])
                    continue
                text_part() if run else None
                n_pic += 1
                parts.append(("picture", drawn[0], drawn[1], bv + drawn[2]))
            elif isinstance(it, tuple):
                text_part() if run else None
                n_strip += 1
                tb, r = draw_strip(slide, x, y, w, room, it[2], _pick_role(it[1], "accent", slide), solid, f"Card {n_col} strip {n_strip}")
                parts.append(("strip", tb, r, tb.height + 2 * em(BOXES["strip"]["gap"])))   # its room: the gap above and below
            else:
                run.append(it)
        if run or not parts or parts[-1][0] in ("strip", "quote", "picture"):
            text_part()
        first = hb if hb is not None else parts[0][1]
        head_h = _card_height(hb, slide) if hb is not None else 0
        body_h = sum(h for _, _, _, h in parts)   # measured, no slack: the last part runs to the foot
        pics = [(sh, cb) for kind, sh, cb, _ in parts if kind == "picture"]
        if pics and _yield_pictures(pics, head_h + body_h - int(room / SLACK), slide):   # the card past the room
            parts = [(kind, sh, r, bv + sh.height + (r.height if r is not None else 0)) if kind == "picture" else (kind, sh, r, h)
                     for kind, sh, r, h in parts]
            body_h = sum(h for _, _, _, h in parts)
        cards.append(dict(attrs=attrs, k=k, x=x, w=w, n=n_col, hb=hb, parts=parts, first=first, head_h=head_h, h=head_h + body_h))
    total = min(int(max(c["h"] for c in cards if "plain" not in c) * SLACK), room)   # the slack once, at the foot
    for c in cards:
        if "plain" in c:
            _place(slide, c["plain"], y, total)
            continue
        attrs, x, w, n, hb, head_h = c["attrs"], c["x"], c["w"], c["n"], c["hb"], c["head_h"]
        role, solid = _hex(_pick_role(attrs, ROLES[(c["k"] - 1) % len(ROLES)], slide)), attrs.get("solid")
        mine = [shape(slide, x, y, w, total, radius, role if solid else _hex(card["fill"]), f"Card {n} fill", behind=c["first"])]
        if hb is not None and not solid:   # rounded at the top only, square where the body starts
            mine.append(shape(slide, x, y, w, head_h, radius, role, f"Card {n} head", behind=c["first"], top_only=True))
        if attrs.get("pick"):
            mine.append(shape(slide, x, y, w, total, radius, None, f"Card {n} outline",
                               line=(_hex(card["outline-color"]), round(card["outline"] * slide_height))))
        if hb is not None:
            hb.top, hb.height = y, head_h
        cy = y + head_h
        gap = em(BOXES["strip"]["gap"])   # a strip sits boxes.strip.gap inside its room, above and below
        for i, (kind, tb, r, h) in enumerate(c["parts"]):
            if kind == "picture":
                place_picture(tb, r, cy + bv)
                cy += h
                continue
            if i == len(c["parts"]) - 1 and kind == "text":   # the last text runs to the card's foot
                h = max(y + total - cy, 0)
            inset_y = gap if kind == "strip" else 0
            tb.top, tb.height = cy + inset_y, h - 2 * inset_y
            if r is not None:
                r.top = cy + inset_y
            cy += h
        mine += [hb] + [part for _, tb, r, _ in c["parts"] for part in (tb, r)]
        group(slide, mine, f"Card {n}")   # one object to move or copy in PowerPoint
    return total, n_col


def _oval(slide, x, y, d, fill, name, line=None):
    """A circle of diameter d (EMU) at x, y: fill a six-digit hex, line (hex, width EMU) an outline
    inside its edge (the circle shrinks by half the width, as HTML's border-box), no shadow."""
    edge = line[1] // 2 if line else 0
    return shape(slide, x + edge, y + edge, d - 2 * edge, d - 2 * edge, 0, fill, name, line=line, preset=MSO_SHAPE.OVAL)


def _badge(slide, x, y, d, fill, k, name):
    """A step's badge: a circle in the role, its number centered in it, the subhead size, bold,
    knocked out in the background color (the shape's own text frame, no insets, no wrap)."""
    o = _oval(slide, x, y, d, fill, name)
    tf = o.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0   # left first: the attributes' order in the XML
    tf.word_wrap = False
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p._p.get_or_add_pPr().set("algn", "ctr")
    p.add_run().text = str(k)
    style_runs(p._p, livery_mod.type_pt(LIVERY, "subhead"), True, _hex("background"))
    return o


def draw_rail(slide, x, y, w, room, kind, items, attrs, slide_width, slide_height):
    """Steps or a timeline: the items in equal columns (column_frames(), the columns' gap), a rail
    of boxes.rail.line in line-color through the marks, from the first mark's middle to the row's
    right edge; a step's badge (badge of the height, the role's, its number in the background color,
    counted from the marker's start=, the list's first number) over a
    text box of its title (the subhead size, bold) and text, or a milestone's date (subhead, bold,
    to the box's foot) over its dot (dot of the height, the role's when done, else the background
    in a ring of the role) over its label. The text boxes padded by boxes.rail.pad, a mark and a
    date at the text's left; a cell of boxes.rail.fill behind each text box when the livery sets
    one; the text boxes one height, the tallest measured, capped at the room. One group, "Steps"
    or "Timeline"; the height it takes, a pad's gap under it included."""
    rail = BOXES["rail"]
    steps = kind == "steps"
    noun = "Step" if steps else "Milestone"
    role = _hex(_pick_role(attrs, BOXES["rail"]["role"], slide))
    start = int(attrs.get("start", 1))
    mark = round(rail["badge" if steps else "dot"] * slide_height)
    vpad, hpad = pads("rail")
    frames = column_frames(x, w, [1 / len(items)] * len(items), slide_width)
    subhead = livery_mod.type_pt(LIVERY, "subhead")
    shapes, texts, dates = [], [], []
    for k, ((_, iattrs, paras), (cx, cw)) in enumerate(zip(items, frames), 1):
        title, body = [], []
        for p in paras:
            kind_ = _prefix(p)
            _close(p)
            if kind_ == "title":
                style_runs(p, subhead, True); title.append(p)
            else:
                body.append(p)
        if not steps:   # the date, a box of its own above the dot
            db = body_box(slide, cx, y, cw, room, title, f"{noun} {k} title")
            db.text_frame.margin_top = 0; db.text_frame.margin_bottom = vpad
            db.text_frame.margin_left = db.text_frame.margin_right = hpad
            db.text_frame.vertical_anchor = MSO_ANCHOR.BOTTOM
            dates.append(db)
            title = []
        tb = body_box(slide, cx, y, cw, room, title + body, f"{noun} {k} text")
        inset(tb.text_frame, vpad, hpad)
        texts.append((k, cx, cw, iattrs, tb))
    date_h = max((_card_height(db, slide) for db in dates), default=0)
    for db in dates:
        db.top, db.height = y, date_h
    my = y + date_h   # the marks' row
    ty = my + mark    # the text boxes' top
    text_h = min(int(max(_card_height(tb, slide) for *_, tb in texts) * SLACK), max(y + room - ty, 0))
    line = round(rail["line"] * slide_height)
    x0 = frames[0][0] + hpad + mark // 2
    shapes.append(rect(slide, x0, my + (mark - line) // 2, x + w - x0, line, _hex(rail["line-color"]), "Rail"))
    for k, cx, cw, iattrs, tb in texts:
        tb.top, tb.height = ty, text_h
        if rail["fill"]:
            shapes.append(rect(slide, cx, ty, cw, text_h, _hex(rail["fill"]), f"{noun} {k} cell", behind=tb))
        if steps:
            shapes.append(_badge(slide, cx + hpad, my, mark, role, k + start - 1, f"{noun} {k} badge"))
        elif iattrs.get("done"):
            shapes.append(_oval(slide, cx + hpad, my, mark, role, f"{noun} {k} dot"))
        else:
            shapes.append(_oval(slide, cx + hpad, my, mark, _hex("background"), f"{noun} {k} dot",
                                line=(role, max(round(BOXES["rail"]["ring"] * mark), 1))))
        shapes.append(tb)
    shapes += dates
    group(slide, shapes, "Steps" if steps else "Timeline")
    return ty + text_h - y + vpad


def lay_out(slide, ph, segs, slide_width, slide_height, deck_dir):
    left, top, width = ph.left, ph.top, ph.width
    bottom = round(GEOM["body-bottom"] * slide_height)
    y, n_col, keep, taken = top, 0, False, set()
    for i, seg in enumerate(segs):
        later = segs[i + 1:]
        if i in taken:   # a full slide's caption, drawn with its gallery
            continue
        if seg[0] == "gallery":   # on a {.full} slide the text after it is the slide's caption
            caption = None
            if seg[1].get("full") and later and later[0][0] == "text":
                caption = later[0][1]; taken.add(i + 1)
            y += draw_gallery(slide, (left, y, width, max(bottom - y, 0)), seg[1], seg[2], caption,
                              slide_width, slide_height, deck_dir)
            continue
        if seg[0] == "text":
            if i == 0:   # before the columns: the placeholder keeps it, cut to its text
                body = ph.text_frame._txBody
                for p in body.findall(qn("a:p")):
                    if p not in seg[1]:
                        body.remove(p)
                h = min(int(text_height(ph, slide) * SLACK), bottom - y)
                ph.left, ph.top, ph.width, ph.height = left, y, width, h
                keep = True
            else:
                tb = body_box(slide, left, y, width, max(bottom - y, 0), seg[1], "Body text")
                h = min(int(text_height(tb, slide) * SLACK), max(bottom - y, 0))
                tb.height = h
            y += h
            continue
        if seg[0] == "lead":
            y += draw_lead(slide, left, y, width, max(bottom - y, 0), seg[2], slide_height)
            continue
        if seg[0] == "agenda":
            y += draw_agenda(slide, left, y, width, max(bottom - y, 0), seg[2])
            continue
        if seg[0] == "rail":
            y += draw_rail(slide, left, y, width, max(bottom - y, 0), seg[1].get("kind", "steps"), seg[2], seg[1],
                           slide_width, slide_height)
            continue
        if seg[0] == "quote":   # on its own: as wide as the body, a vertical pad under it
            tb, bar = draw_quote(slide, left, y, width, max(bottom - y, 0), seg[2], _quote_color(seg[1], slide), slide_height)
            group(slide, [bar, tb], "Quote")
            y += tb.height + pads("quote")[0]
            continue
        if seg[0] == "strip":   # on its own: as wide as the body, boxes.strip.gap above and below it
            gap = em(BOXES["strip"]["gap"])
            y += gap
            tb, r = draw_strip(slide, left, y, width, max(bottom - y, 0), seg[2], _pick_role(seg[1], "accent", slide), False, "Strip fill")
            tb.name = "Strip text"
            group(slide, [r, tb], "Strip")
            y += tb.height + gap
            continue
        _, shares, columns = seg
        frames = column_frames(left, width, shares[:len(columns)] or [1 / len(columns)] * len(columns), slide_width)
        room = max(bottom - y, 0)
        if any(c[1].get("card") for c in columns):
            h, n_col = draw_cards(slide, columns, frames, y, room, n_col, slide_height, deck_dir)
            rules(slide, frames, y, h, slide_width, slide_height)
            y += h
            continue
        cols = []
        for (_, attrs, paras), (x, w) in zip(columns, frames):
            n_col += 1
            cols.append(_plain_column(slide, x, y, w, room, paras, f"Column {n_col}", slide_height, deck_dir))
        h = room
        tallest = min(max(_column_height(c, slide) for c in cols), room)
        if any(s[0] == "text" for s in later):   # text below: the row as tall as its tallest column
            h = min(int(tallest * SLACK), room)
        for c in cols:
            _place(slide, c, y, h)
        rules(slide, frames, y, tallest, slide_width, slide_height)
        y += h
    if not keep:
        ph._element.getparent().remove(ph._element)
    return y


def run(prs, path: Path) -> None:
    """The open deck, changed in place; path the pptx's (its name for the messages; its directory the deck's,
    where a gallery's files are)."""
    done = 0
    for slide in prs.slides:
        for ph in [sh for sh in slide.shapes if is_body(sh)]:
            txBody = ph.text_frame._txBody
            if not any("⟦engine:" in para_text(p) for p in txBody.findall(qn("a:p"))):
                continue
            # a lead after a figure: single-body.lua folds it into the figure's caption, which pandoc
            # writes as one centered paragraph, its lines parted by a:br and the markers inline
            split_br(txBody)
            paragraphs = txBody.findall(qn("a:p"))
            segs = parse(paragraphs)
            if segs is None:
                continue
            for i in marker_spans(paragraphs):   # the fold's centering, off the lines grid draws
                ppr = paragraphs[i].find(qn("a:pPr"))
                if ppr is not None:
                    ppr.attrib.pop("algn", None)
            for p in paragraphs:   # the marker paragraphs go
                if MARK.match(para_text(p).strip()):
                    p.getparent().remove(p)
            above = [sh for sh in slide.shapes if is_picture(sh) and sh.top + sh.height <= ph.top]
            # all four of a placeholder's sides at once: it inherits its frame, and setting one alone
            # leaves the rest at 0 (pandoc's caption box carries a frame, a Content with Caption's text box none)
            if above:   # pandoc's caption box, under the picture: the flow starts just below it
                ph.left, ph.top, ph.width, ph.height = ph.left, max(sh.top + sh.height for sh in above) + em(0.5), ph.width, ph.height
            beside = [sh for sh in slide.shapes if is_table(sh) and sh.left >= ph.left + ph.width - em(1)]
            if beside:   # a lead before a table: pandoc's Content with Caption set them side by side
                use_layout(prs, slide, "Title and Content")
                ph.left, ph.top, ph.width, ph.height = body_frame(prs, ph)
            left, width = ph.left, ph.width
            before = {sh._element for sh in slide.shapes}
            y = lay_out(slide, ph, segs, prs.slide_width, prs.slide_height, path.parent)
            bottom = round(GEOM["body-bottom"] * prs.slide_height)
            if above and y > bottom:   # the lead needs more room than the caption had: the picture yields
                # no further than PIC_FLOOR em of its height: past that the lead runs over the foot, warned,
                # rather than the picture going negative (python-pptx refuses it and the build fails)
                room = min(pic.height for pic in above) - em(PIC_FLOOR)
                over = max(0, min(y - bottom, room))
                if y - bottom > room:
                    print(f"grid: WARNING: slide {prs.slides.index(slide) + 1}: the lead under the picture needs more room "
                          f"than the picture has; it runs past the body's foot", file=sys.stderr)
                for pic in above:
                    k, middle = (pic.height - over) / pic.height, pic.left + pic.width // 2
                    pic.width, pic.height = int(pic.width * k), pic.height - over
                    pic.left = middle - pic.width // 2
                for sh in slide.shapes:
                    if sh._element not in before or sh._element is ph._element:
                        sh.top -= over
            if beside:   # the note's look, which slide_layout skipped beside two bodies (one is left now)
                note_style(slide)
            notes = [sh for sh in slide.shapes if is_body(sh) and sh._element is not ph._element]
            for t in beside:   # the table under the flow, as wide as the body, its note under it
                fit_table(t, left, width)
                t.top, t.height = y + em(0.5), table_height(t, slide)
                for note in notes:   # under the table, raised when the text would pass the body's foot
                    frame_note(note, slide, left, width, t.top + t.height + NOTE_GAP, bottom)
            done += 1
        left = [t.text for t in slide._element.iter(qn("a:t")) if t.text and "⟦" in t.text]
        if left:
            print(f"grid: WARNING: slide {prs.slides.index(slide) + 1}: markers left: {left[:3]}", file=sys.stderr)
    if done:
        print(f"grid: drew {done} grid slide(s) → {path.name}")
