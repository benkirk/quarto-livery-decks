"""What the pptx steps share: the livery and the numbers every step reads from it, the shape and run
helpers each step wrote for itself (text boxes in the body's style, groups, card shapes, a picture
fitted to a box), the idioms the drawers share (an em in EMU, a frame's insets, a paragraph's space
before, a slide's title for a warning), and the shape tests.

The livery is read once here (livery.load(): $LIVERY_DIR, exported by Make.common, else a `_livery`
link in the current directory) and every step imports it, so one deck's steps agree on it. The
constants are the canvas scale (K: absolute pads were drawn for the 540 pt canvas), the boxes and
geometry blocks of livery.json, the body type in points, and SLACK, the share a measured text
height is grown by so PowerPoint's autofit stays idle. measure.py holds the text measuring
(text_height); this module holds no measuring.
"""

import copy
import io
import re

from lxml import etree
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.oxml.ns import qn
from pptx.util import Emu

import livery as livery_mod

LIVERY = livery_mod.load()
K = livery_mod.canvas_scale(LIVERY)                # the canvas height over 540 pt: scales absolute pads
BOXES = LIVERY["boxes"]                            # the boxes' numbers, every format's (livery.json boxes)
GEOM = LIVERY["geometry"]
BODY_PT = livery_mod.type_pt(LIVERY, "body")       # pads are em of the body type
TRACKING = LIVERY["type"]["tracking"]              # small caps' letter spacing, em of the run's size (type.tracking)
SLACK = 1.1                                        # measured height x this: autofit stays idle
NOT_BODY = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.DATE,
            PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.SLIDE_NUMBER}
# a:pPr's children, in schema order: a bullet's elements go before the rest
PPR_ORDER = ["lnSpc", "spcBef", "spcAft", "buClrTx", "buClr", "buSzTx", "buSzPct", "buSzPts", "buFontTx", "buFont",
             "buNone", "buAutoNum", "buChar", "buBlip", "tabLst", "defRPr", "extLst"]


def is_body(shape) -> bool:
    """A body text placeholder (not the title, date, footer or number)."""
    return (shape.is_placeholder and shape.has_text_frame
            and shape.placeholder_format.type not in NOT_BODY)


def is_table(shape) -> bool:
    """A table's graphic frame."""
    return bool(getattr(shape, "has_table", False))


def is_picture(shape) -> bool:
    return shape.shape_type == MSO_SHAPE_TYPE.PICTURE


def rpr(run):
    """A run's a:rPr, made when it has none (a run without one inherits; sz and color need a home)."""
    el = run.find(qn("a:rPr"))
    if el is None:
        el = etree.Element(qn("a:rPr"))
        run.insert(0, el)
    return el


def set_fill(rpr_el, hexv: str) -> None:
    """The run's color: one a:solidFill of the six-digit hex, after a:ln when there is one."""
    for old in rpr_el.findall(qn("a:solidFill")):
        rpr_el.remove(old)
    fill = etree.Element(qn("a:solidFill"))
    etree.SubElement(fill, qn("a:srgbClr")).set("val", hexv)
    ln = rpr_el.find(qn("a:ln"))
    ln.addnext(fill) if ln is not None else rpr_el.insert(0, fill)


def ppr_place(ppr, el) -> None:
    """el into a:pPr at its schema position (PPR_ORDER)."""
    rank = PPR_ORDER.index(etree.QName(el).localname)
    for child in ppr:
        name = etree.QName(child).localname
        if name in PPR_ORDER and PPR_ORDER.index(name) > rank:
            child.addprevious(el)
            return
    ppr.append(el)


def rect(slide, x, y, w, h, hexv: str, name=None, behind=None):
    """A filled rectangle (EMU, a six-digit hex), no line or shadow; behind: a shape it goes under."""
    r = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(x), Emu(y), Emu(w), Emu(h))
    r.fill.solid(); r.fill.fore_color.rgb = RGBColor.from_string(hexv)
    r.line.fill.background(); r.shadow.inherit = False
    if name:
        r.name = name
    if behind is not None:
        behind._element.addprevious(r._element)
    return r


def take_marker(slide, key):
    """The words of the slide's notes lines that start with key, those paragraphs removed (for
    engine-layout:, slide-layout.lua writes one, single-body.lua another for a note; columns.py
    reads grid-pptx.lua's engine-columns:). None when the slide has no such line."""
    if not slide.has_notes_slide:
        return None
    tf = slide.notes_slide.notes_text_frame
    if tf is None:
        return None
    words = None
    for p in tf.paragraphs:
        if p.text.strip().startswith(key):
            words = (words or []) + p.text.strip()[len(key):].split()
            body = p._p.getparent()
            body.remove(p._p)
            if body.find(qn("a:p")) is None:
                etree.SubElement(body, qn("a:p"))
    return words


# a light livery's dark content slide ({.brand-dark}): the colors a step draws with there, the
# light ones swapped as livery.css swaps the roles (livery_mod.DARK_SWAP; background to the field)
LIGHT_INK = {k: livery_mod.hex(LIVERY, k) for k in (*livery_mod.DARK_SWAP, "background")}
DARK_INK = {**LIGHT_INK, **({k: livery_mod.hex(LIVERY, v) for k, v in {**livery_mod.DARK_SWAP, "background": "dark"}.items()}
                            if LIVERY["mode"] != "dark" else {})}


def is_dark(slide) -> bool:
    """A light livery's dark content slide: its background a solid fill of the dark role
    (slide_layout.draw_dark()), or, for a step before that one, its engine-layout notes line naming
    dark (slide-layout.lua's). Never on a dark livery, where the class is a no-op."""
    if LIVERY["mode"] == "dark":
        return False
    bg = slide._element.find(qn("p:cSld") + "/" + qn("p:bg"))
    clr = bg.find(".//" + qn("a:srgbClr")) if bg is not None else None
    if clr is not None and clr.get("val", "").upper() == DARK_INK["background"]:
        return True
    if slide.has_notes_slide:
        for line in slide.notes_slide.notes_text_frame.text.splitlines():
            if line.strip().startswith("engine-layout:") and "dark" in line.split(":", 1)[1].split():
                return True
    return False


def palette(slide) -> dict:
    """The six-digit hexes a step draws with on this slide, by the light role or blend they stand
    for (foreground, muted, primary-contrast, primary, table-rule, table-band, background): the
    livery's, or on a dark slide (is_dark()) the dark ones, on-dark, dark-muted, secondary, the
    dark table blends, the dark role."""
    return DARK_INK if is_dark(slide) else LIGHT_INK


def on_dark(paragraphs, hexv, skip=()):
    """Every run of the paragraphs that takes its color from the master (no a:solidFill of its own)
    in hexv; a run a step colored keeps its color. skip: the indexes left alone."""
    for i, p in enumerate(paragraphs):
        if i in skip:
            continue
        for r in p.iter(qn("a:r")):
            el = rpr(r)
            if el.find(qn("a:solidFill")) is None:
                set_fill(el, hexv)


# a light ink -> its dark twin (palette()'s swap), the foreground last: it wins where a livery's
# foreground is also one of the others
DARK_MAP = {LIGHT_INK[k]: DARK_INK[k] for k in ("muted", "primary-contrast", "primary", "foreground")}


def darken(el):
    """A drawn shape's (an element's) ink for a dark slide: every run that takes the master's color
    on-dark, and a run, fill or line in one of the light inks of palette() in its dark twin
    (muted to dark-muted, primary and primary-contrast to secondary, foreground to on-dark); any
    other color kept."""
    for r in el.iter(qn("a:r")):
        if rpr(r).find(qn("a:solidFill")) is None:
            set_fill(rpr(r), DARK_INK["foreground"])
    for clr in el.iter(qn("a:srgbClr")):
        v = (clr.get("val") or "").upper()
        if v in DARK_MAP and clr.getparent().tag == qn("a:solidFill"):
            clr.set("val", DARK_MAP[v])


def column_frames(left, width, shares, slide_width):
    """(left, width) in EMU of each column in a box (left, width EMU): its share of the box less
    the gaps (boxes.columns.gap of the slide's width); columns.py's and grid.py's one formula."""
    g = round(BOXES["columns"]["gap"] * slide_width)
    room, x, out = width - g * (len(shares) - 1), left, []
    for s in shares:
        w = round(room * s)
        out.append((x, w))
        x += w + g
    return out


def em(v, pt=BODY_PT):
    """v em of a type size (pt, the body's by default) in EMU, rounded (Pt() truncates 91439.99)."""
    return Emu(round(v * pt * 12700))


def pads(box, key="pad"):
    """A box's pads, (vertical, horizontal) in EMU: its em times the body type."""
    v, h = BOXES[box][key]
    return em(v), em(h)


def inset(tf, v, h):
    """A text frame's four insets: v (EMU) top and bottom, h left and right."""
    tf.margin_top = tf.margin_bottom = v; tf.margin_left = tf.margin_right = h


def slide_title(slide):
    """The slide's title text, for a warning; "" when it has no title."""
    t = slide.shapes.title
    return t.text_frame.text if t is not None else ""


def space_before(p, pts):
    """A paragraph's space before (a:spcBef), in points."""
    ppr = p.find(qn("a:pPr"))
    if ppr is None:
        ppr = etree.Element(qn("a:pPr")); p.insert(0, ppr)
    for old in ppr.findall(qn("a:spcBef")):
        ppr.remove(old)
    bef = etree.Element(qn("a:spcBef"))
    etree.SubElement(bef, qn("a:spcPts")).set("val", str(round(pts * 100)))
    ppr_place(ppr, bef)


def fit_table(shape, left, width):
    """A table's frame at left for width (EMU), its grid columns scaled to it in their shares
    (style_tables re-splits them inside the frame later); its top and height kept."""
    cols = shape.table._tbl.findall(qn("a:tblGrid") + "/" + qn("a:gridCol"))
    total = sum(int(c.get("w")) for c in cols)
    for c in cols:   # integers: a float makes PowerPoint repair the deck
        c.set("w", str(round(int(c.get("w")) * width / total)))
    shape.left, shape.top, shape.width, shape.height = Emu(left), shape.top, Emu(width), shape.height


def para_text(p):
    """A paragraph's (a:p) text, every run's."""
    return "".join(t.text or "" for t in p.iter(qn("a:t")))


MARK = re.compile(r"^⟦engine:(\w+)(?:\s+(.*))?⟧$")   # a grid-pptx.lua marker paragraph, on its own


def split_br(txBody):
    """Each paragraph cut at its line breaks (a:br) into paragraphs of their own, each with a copy
    of the paragraph's properties: pandoc writes a caption single-body.lua folded (a ::: note, a
    lead's markers and lines) as one paragraph, its lines joined by a:br."""
    for p in list(txBody.iterchildren(qn("a:p"))):
        ppr, cur = p.find(qn("a:pPr")), p
        for el in list(p):
            if el.tag == qn("a:br"):
                new = etree.Element(qn("a:p"))
                if ppr is not None:
                    new.append(copy.deepcopy(ppr))
                cur.addnext(new)
                p.remove(el)
                cur = new
            elif cur is not p and el.tag != qn("a:pPr"):
                cur.append(el)   # moves it, the end-of-paragraph properties too


def marker_spans(paragraphs):
    """The indexes of the paragraphs inside a marker span (an opener `⟦engine:x⟧` to its
    `⟦engine:end⟧`, both included, nested spans too): what grid.py draws, which another step
    styling the box leaves alone."""
    out, depth = set(), 0
    for i, p in enumerate(paragraphs):
        m = MARK.match(para_text(p).strip())
        if m:
            depth = max(depth - 1, 0) if m.group(1) == "end" else depth + 1
            out.add(i)
        elif depth:
            out.add(i)
    return out


def body_frame(prs, ph):
    """The Title and Content layout's body frame (left, top, width, height), the placeholder's own
    when the template has no such layout: all four, since a placeholder inherits its frame and
    setting one side alone leaves the rest at 0."""
    layout = next((l for l in prs.slide_layouts if l.name == "Title and Content"), None)
    body = next((sh for sh in layout.placeholders if sh.placeholder_format.idx == 1), None) if layout else None
    return (body.left, body.top, body.width, body.height) if body is not None else (ph.left, ph.top, ph.width, ph.height)


def style_runs(p, size=None, bold=None, color=None, caps=False):
    """Every run of the paragraph at a size (pt), bold, a six-digit color, small caps tracked by
    type.tracking (em of the run's size; a:spc is in hundredths of a point)."""
    for r in p.iter(qn("a:r")):
        el = rpr(r)
        if size is not None:
            el.set("sz", str(round(livery_mod.half_pt(size) * 100)))
        if bold is not None:
            el.set("b", "1" if bold else "0")
        if caps:
            sz = int(el.get("sz") or round(BODY_PT * 100))
            el.set("cap", "all"); el.set("spc", str(round(sz * TRACKING)))
        if color is not None:
            set_fill(el, color)


CAPTION_PT = livery_mod.type_pt(LIVERY, "caption")


def caption_para(p, color=None):
    """A caption paragraph (a gallery's, a picture's in a cell): centered, no space before it, the
    caption size, muted (or color, a six-digit hex)."""
    ppr = p.find(qn("a:pPr"))
    if ppr is None:
        ppr = etree.Element(qn("a:pPr")); p.insert(0, ppr)
    ppr.set("algn", "ctr")
    space_before(p, 0)
    style_runs(p, CAPTION_PT, color=color or livery_mod.hex(LIVERY, "muted"))


def body_box(slide, x, y, w, h, paragraphs, name):
    """A text box that renders as the body placeholder does: the master's body list style (bullets,
    sizes, spacing per level) in its lstStyle and its insets, word wrap, shrink on overflow; the
    paragraphs moved into it with their own properties."""
    tb = slide.shapes.add_textbox(Emu(x), Emu(y), Emu(w), Emu(h))
    tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    body = tf._txBody
    bp = body.find(qn("a:bodyPr"))
    for fit in list(bp):
        if fit.tag in (qn("a:spAutoFit"), qn("a:noAutofit"), qn("a:normAutofit")):
            bp.remove(fit)
    etree.SubElement(bp, qn("a:normAutofit"))
    master = slide.slide_layout.slide_master._element
    for sp in master.iter(qn("p:sp")):   # the body placeholder's insets (a text box's default is 0.1 in)
        ph = sp.find(".//" + qn("p:ph"))
        mbp = sp.find(".//" + qn("a:bodyPr"))
        if ph is not None and ph.get("type") == "body" and mbp is not None:
            for k in ("lIns", "tIns", "rIns", "bIns"):
                if mbp.get(k) is not None:
                    bp.set(k, mbp.get(k))
            break
    lst = body.find(qn("a:lstStyle"))
    style = master.find(".//" + qn("p:bodyStyle"))
    for lvl in (style if style is not None else []):
        if lvl.tag.endswith("pPr"):
            lst.append(copy.deepcopy(lvl))
    for p in body.findall(qn("a:p")):
        body.remove(p)
    for p in paragraphs:
        body.append(p)   # moves it
    if not paragraphs:
        etree.SubElement(body, qn("a:p"))
    return tb


def group(slide, shapes, name):
    """The shapes as one PowerPoint group, named: a click selects the whole box (a card, a lead),
    a second click a part. At the first shape's place in the stacking (python-pptx appends a
    group on top); its child coordinates are the slide's (recalculate_extents)."""
    tree = slide.shapes._spTree
    shapes = sorted((sh for sh in shapes if sh is not None), key=lambda sh: list(tree).index(sh._element))
    at = list(tree).index(shapes[0]._element)
    g = slide.shapes.add_group_shape(shapes)
    tree.remove(g._element); tree.insert(at, g._element)
    g.name = name
    return g


def shape(slide, x, y, w, h, radius, fill, name, line=None, behind=None, top_only=False, preset=None):
    """A card's rectangle, rounded when radius (EMU) > 0 (top_only: its top corners, a head band's),
    or another preset (MSO_SHAPE.OVAL, a rail's marks); fill a six-digit hex or None; line (hex,
    width EMU) an outline; behind: a shape it goes under. Its own shape, not rect(): the preset
    geometry differs."""
    if preset is None:
        preset = (MSO_SHAPE.ROUND_2_SAME_RECTANGLE if top_only else MSO_SHAPE.ROUNDED_RECTANGLE) if radius > 0 else MSO_SHAPE.RECTANGLE
    r = slide.shapes.add_shape(preset, Emu(x), Emu(y), Emu(w), Emu(h))
    if radius > 0:
        r.adjustments[0] = min(0.5, round(radius / min(w, h), 5))
        if top_only:
            r.adjustments[1] = 0.0
    if fill:
        r.fill.solid(); r.fill.fore_color.rgb = RGBColor.from_string(fill)
    else:
        r.fill.background()
    if line:
        r.line.color.rgb = RGBColor.from_string(line[0]); r.line.width = Emu(line[1])
    else:
        r.line.fill.background()
    r.shadow.inherit = False
    r.name = name
    if behind is not None:
        behind._element.addprevious(r._element)
    return r


def contain_box(iw, ih, x, y, w, h):
    """The largest box of an iw x ih image's aspect inside the box (x, y, w, h EMU), centered in it:
    (x, y, w, h) in integer EMU."""
    k = min(w / iw, h / ih)
    cw, ch = int(iw * k), int(ih * k)
    return x + (w - cw) // 2, y + (h - ch) // 2, cw, ch


def cover_crop(iw, ih, w, h):
    """The symmetric crop, (left and right, top and bottom) shares of an iw x ih image, that makes it
    cover a w x h box with its aspect kept: one of the two is 0."""
    if iw * h > ih * w:
        return (1 - (w / h) / (iw / ih)) / 2, 0.0
    return 0.0, (1 - (iw / ih) / (w / h)) / 2


def fit_picture(slide, src, x, y, w, h, fit, name=None):
    """A picture of src (a path or the file's bytes) in the box (x, y, w, h EMU): fit "contain", the
    whole image at its aspect, centered (contain_box()); "cover", the box filled, the image cropped
    symmetrically (cover_crop(), the picture's a:srcRect). The picture."""
    pic = slide.shapes.add_picture(io.BytesIO(src) if isinstance(src, bytes) else str(src), Emu(x), Emu(y), Emu(w), Emu(h))
    iw, ih = pic.image.size
    if fit == "cover":
        lr, tb = cover_crop(iw, ih, w, h)
        if lr:
            pic.crop_left = pic.crop_right = lr
        else:
            pic.crop_top = pic.crop_bottom = tb
    else:
        pic.left, pic.top, pic.width, pic.height = (Emu(v) for v in contain_box(iw, ih, x, y, w, h))
    if name:
        pic.name = name
    return pic

