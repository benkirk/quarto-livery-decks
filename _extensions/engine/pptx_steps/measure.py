"""Text measured in the livery's own fonts (PIL): the pptx steps' one measurer. It counts the lines
paragraphs wrap to in a width, greedily, as PowerPoint wraps them, and from them a shape's text
height (text_height(): the lines of each paragraph in the face and size the pptx shows it, one LINE
per line, plus the paragraph spacing), which the steps that cut a box to its text read.

Sizes and widths are the revealjs theme's slide pixels (1600 x 900 over the livery's canvas,
livery.json type.canvas): font() takes points of type, lines() and runs_lines() a width in
pixels. A face is the one the pptx sets, measured in the livery's bundled files
(livery.font_file): its typeface is the base family's stand-in (cu's Helvetica Neue measures as
TeX Gyre Heros), titles take the headings family only when the livery names a pptx
heading-typeface (else the template sets them in the typeface too), a quote's lines its
boxes.quote.family (under its pptx.quote-typeface, or its own name), and a face the livery does
not bundle, pandoc's Courier for code, advances 0.6 em a character, as Courier does."""

import functools
import re

from PIL import ImageFont
from pptx.oxml.ns import qn
from pptx.util import Pt

import livery as livery_mod

from .common import BODY_PT, K, LIVERY

CANVAS = LIVERY["type"]["canvas"]
PX = CANVAS["w"] * 12700 / 1600             # EMU per slide pixel
PT_PX = CANVAS["h"] / 900                   # points of type per slide pixel
TYPE = LIVERY["typography"]
PPTX = LIVERY["pptx"]
BODY = TYPE["base"]                                            # the pptx's body text
TITLE = TYPE["headings"] if PPTX.get("heading-typeface") else BODY   # its titles (+mj-lt)
BUNDLED = {f["family"] for f in TYPE.get("fonts", []) if f.get("files")}
PPTX_FACES = {PPTX.get("typeface"): BODY, PPTX.get("heading-typeface"): TYPE.get("headings"),
              PPTX.get("quote-typeface"): LIVERY["boxes"]["quote"]["family"]}   # a quote's (grid.py), in its family's files
LINE = 1.2                                  # line height / font size, for the text estimates


class _Mono:
    """A face the livery does not bundle (pandoc's Courier): 0.6 em a character."""
    def __init__(self, px):
        self.px = px

    def getlength(self, text):
        return 0.6 * self.px * len(text)


@functools.lru_cache
def font(family, weight, pt):
    """The family's face at weight (400 when it has no such file) and pt points."""
    try:
        path = livery_mod.font_file(LIVERY, family, weight)
    except FileNotFoundError:
        path = livery_mod.font_file(LIVERY, family, 400)
    # BASIC, always: Pillow shapes with Raqm (HarfBuzz) whenever it can load a FriBiDi library at
    # run time, so the same text measured differently on a machine that had one (conda-forge
    # librsvg brings it) and moved a wrap
    return ImageFont.truetype(str(path), round(pt / PT_PX), layout_engine=ImageFont.Layout.BASIC)


def face(name, family, weight, pt):
    """The face a run named `name` (None: the placeholder's own, `family`) is shown in."""
    if name is None:
        return font(family, weight, pt)
    if name in BUNDLED:
        return font(name, weight, pt)
    if PPTX_FACES.get(name):
        return font(PPTX_FACES[name], weight, pt)
    return _Mono(pt / PT_PX)


def lines(paragraphs, font, width):
    """How many lines `paragraphs` (strings in one face) wrap to in `width` px."""
    n = 0
    for text in paragraphs:
        line = ""
        n += 1
        for word in text.split():
            trial = f"{line} {word}" if line else word
            if line and font.getlength(trial) > width:
                n, line = n + 1, word
            else:
                line = trial
    return n


def runs_lines(pieces, width):
    """How many lines one paragraph of (text, face) pieces wraps to in `width` px; a "\v" (a
    line break) starts a new line."""
    rows = [[]]
    for text, f in pieces:
        for i, part in enumerate(text.split("\v")):
            if i:
                rows.append([])
            if part:
                rows[-1].append((part, f))
    return sum(_wrap(row, width) for row in rows)


def _wrap(pieces, width):
    """Greedy, by words: a word's width the sum of its pieces' (one may cross runs: a code span
    and its comma), the space after it its run's."""
    words, cur, have = [], 0.0, False   # [width, space after]
    for text, f in pieces:
        for tok in re.split(r"(\s+)", text):
            if not tok:
                continue
            if tok.isspace():
                if have:
                    words.append([cur, f.getlength(" ")])
                    cur, have = 0.0, False
                elif words:
                    words[-1][1] = f.getlength(" ")
            else:
                cur, have = cur + f.getlength(tok), True
    if have:
        words.append([cur, 0.0])
    n, x, gap = 1, None, 0.0
    for w, space in words:
        if x is None:
            x = w
        elif x + gap + w > width:
            n, x = n + 1, w
        else:
            x += gap + w
        gap = space
    return n


# --- a shape's paragraphs, measured in the faces the pptx shows them in

def lvl_size(txstyle, level):
    """sz (hundredths of a point) for a paragraph level in a list-style element."""
    if txstyle is None:
        return None
    lvl = txstyle.find(qn(f"a:lvl{level + 1}pPr"))
    rpr = lvl.find(qn("a:defRPr")) if lvl is not None else None
    return int(rpr.get("sz")) if rpr is not None and rpr.get("sz") else None


def inherited(shape, slide, level):
    """A body placeholder's size at a level: its layout's lstStyle, then the master."""
    try:
        idx = shape.placeholder_format.idx
        lph = slide.slide_layout.placeholders.get(idx=idx)
    except (KeyError, ValueError, AttributeError):
        lph = None
    if lph is not None:
        sz = lvl_size(lph._element.find(".//" + qn("a:lstStyle")), level)
        if sz:
            return sz
    master = slide.slide_layout.slide_master._element
    return lvl_size(master.find(".//" + qn("p:bodyStyle")), level) or 1800


def other_size(slide):
    """The master's otherStyle size (hundredths of a point): a table's text, a text box's."""
    master = slide.slide_layout.slide_master._element
    return lvl_size(master.find(".//" + qn("p:otherStyle")), 0) or 1800


def para_size(p, shape, slide):
    """A paragraph's size in EMU: its first sized run's, else what the placeholder inherits."""
    sized = [r.font.size for r in p.runs if r.font.size]
    return sized[0] if sized else Pt(inherited(shape, slide, p.level) / 100)


def pieces(p, family, size, weight=None):
    """A paragraph as (text, face) pieces for runs_lines(): each run in the face the pptx shows it
    in, at its own size (else the paragraph's, EMU) and weight (else `weight`, else regular); a
    line break (pandoc's subtitle and author share a paragraph) as "\v"."""
    out = []
    for el in p._p:
        if el.tag == qn("a:br"):
            out.append(("\v", None))
        elif el.tag in (qn("a:r"), qn("a:fld")):
            rpr, t = el.find(qn("a:rPr")), el.find(qn("a:t"))
            latin = rpr.find(qn("a:latin")) if rpr is not None else None
            name = latin.get("typeface") if latin is not None else None
            name = None if name and name.startswith("+") else name   # +mn-lt: the placeholder's own
            pt = int(rpr.get("sz")) / 100 if rpr is not None and rpr.get("sz") else size / 12700
            b = rpr.get("b") if rpr is not None else None
            w = 700 if b in ("1", "true") else 400 if b in ("0", "false") else (weight or 400)
            out.append((t.text or "" if t is not None else "", face(name, family, w, pt)))
    return out


def para_lines(p, family, size, width, weight=None):
    """The lines a paragraph wraps to in width (EMU), measured in the faces the pptx shows."""
    return runs_lines(pieces(p, family, size, weight), width / PX)


def text_height(shape, slide, family=None, weight=None, gap=None):
    """The text's height (EMU): each paragraph's lines measured in the livery's face (family, the
    body's by default), one LINE per line, plus the template's space before the paragraph
    (livery.SPC_BEFORE of the body type, by its level), or gap (EMU) when given."""
    tf = shape.text_frame
    width = shape.width - tf.margin_left - tf.margin_right - Pt(18 * K)   # less the bullet indent (the template's marL)
    total = tf.margin_top + tf.margin_bottom
    for p in tf.paragraphs:
        size = para_size(p, shape, slide)
        before = Pt(BODY_PT * livery_mod.SPC_BEFORE[0 if p.level == 0 else 1]) if gap is None else gap
        total += int(para_lines(p, family or BODY, size, width, weight) * size * LINE) + before
    return total
