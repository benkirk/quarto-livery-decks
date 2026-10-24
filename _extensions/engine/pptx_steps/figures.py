"""A picture in a grid column or a card, and a table in a grid column, drawn from grid-pptx.lua's
markers (grid.py hands it the item).

grid-pptx.lua writes a picture in a flattened column or a card (engine-common.lua's cell_figure()) as
`⟦engine:picture src=<path>[ data=<base64>][ size=<width>]⟧`, its caption a paragraph, `⟦engine:end⟧`:
text, so pandoc embeds no picture and nothing splits. The bytes are data='s when it is there (an
executed cell's plot or a baked diagram, gone from disk by now), else the file at src (percent-encoded
as a gallery item's), relative to the deck's directory; an .svg rasterized to a PNG (no svgBlip twin,
as the gallery's). The picture is the column's width (a sized one its width=: a percentage of the
column, or a length), its aspect kept, centered, and no taller than the room less its caption; the
caption under it in the caption look (the caption size, muted, centered), boxes.note.gap em of the
caption type below the picture, as a ::: note under a figure. "<name>" and "<name> caption". A
missing file warns and draws nothing (its caption stays text).

A simple table in a column (engine-common.lua's simple_table()) is `⟦engine:table cols=n[ header]⟧`,
per row `⟦engine:row⟧` and a paragraph per cell led by `⟦engine:cell⟧`, the caption a paragraph,
`⟦engine:end⟧`: drawn as a native table at the column's width, its columns equal (style_tables, which
runs after grid, re-splits them by their longest cells as it does pandoc's, and styles the table
as the HTML deck's), each cell's paragraph moved in with its runs, the master's table size
inherited as pandoc's are, the first row the header when the marker says so; its caption under it
as a picture's. "<name>" and "<name> caption". A table is never in a card: a row with a card and a
table is drawn as plain columns (tabled_row()). A table is not shortened: one taller than its room
warns.
"""

import base64
import re
import sys
from urllib.parse import unquote

from lxml import etree
from pptx.oxml.ns import qn
from pptx.parts.image import Image
from pptx.util import Emu

from .common import BOXES, CAPTION_PT, SLACK, body_box, caption_para, em, fit_picture, para_text, slide_title
from .measure import text_height
from .slide_layout import table_height
from .svg_fallback import DPI, rasterize_svg   # an .svg's PNG at the deck's fig-dpi

EMU_PER = {"in": 914400, "cm": 360000, "mm": 36000, "pt": 12700, "px": 9525, "": 9525}   # pandoc's bare number: pixels
SIZE = re.compile(r"^([\d.]+)\s*(%|in|cm|mm|pt|px|)$")
ROW, CELL = "⟦engine:row⟧", "⟦engine:cell⟧"


def width_of(size, w):
    """A sized picture's width in EMU (a percentage of w, or a length), at most w; None when size is
    not one of those."""
    m = SIZE.match((size or "").strip())
    if m is None:
        return None
    v, unit = float(m.group(1)), m.group(2)
    return min(w, round(w * v / 100 if unit == "%" else v * EMU_PER[unit]))


def picture_bytes(attrs, deck_dir, slide):
    """A picture marker's image bytes (a PNG for an .svg), or None (warned of)."""
    src = unquote(attrs.get("src", ""))
    if attrs.get("data"):
        data = base64.b64decode(attrs["data"])
    else:
        path = deck_dir / src
        if not path.is_file():
            print(f"figures: WARNING: slide {slide_title(slide)!r}: {path} not found; the picture is not drawn", file=sys.stderr)
            return None
        data = path.read_bytes()
    return rasterize_svg(data, DPI) if src.lower().endswith(".svg") else data


def draw_picture(slide, x, y, w, room, attrs, caption, deck_dir, name):
    """The picture of a marker's attrs at (x, y) in a column w wide, the caption paragraphs under it,
    the two no taller than room: (picture, caption box or None, height) or None when there is no
    picture to draw."""
    data = picture_bytes(attrs, deck_dir, slide)
    if data is None:
        return None
    iw, ih = Image.from_blob(data).size
    cb, cap_h = _caption_box(slide, x, y, w, room, caption, name)
    pw = width_of(attrs.get("size"), w) or w
    ph = round(pw * ih / iw)
    if ph > room - cap_h:   # too tall: the room's height, narrower at its aspect
        ph = max(room - cap_h, 1)
        pw = max(round(ph * iw / ih), 1)
    pic = fit_picture(slide, data, x + (w - pw) // 2, y, pw, ph, "contain", name=name)
    if cb is not None:
        cb._element.addprevious(pic._element)   # the picture first in the stacking: a card's fill goes behind it
        cb.left, cb.top, cb.width, cb.height = Emu(x), Emu(y + ph), Emu(w), Emu(cap_h)
    return pic, cb, ph + cap_h


def shrink_picture(pic, cb, by):
    """A drawn picture by EMU shorter (its aspect kept, centered where it was; its caption raised
    under it): what a column or card gives up when its text needs the room."""
    if by <= 0:
        return
    k = max(pic.height - by, 1) / pic.height
    middle = pic.left + pic.width // 2
    pic.width, pic.height = Emu(max(round(pic.width * k), 1)), Emu(max(pic.height - by, 1))
    pic.left = Emu(middle - pic.width // 2)
    if cb is not None:
        cb.top = Emu(pic.top + pic.height)


def place_picture(pic, cb, top):
    """A picture or a table and its caption moved to top."""
    pic.top = Emu(top)
    if cb is not None:
        cb.top = Emu(top + pic.height)


def _caption_box(slide, x, y, w, room, caption, name):
    """A caption's box under a picture or table: (box, its height), (None, 0) without a caption."""
    if not caption:
        return None, 0
    for p in caption:
        caption_para(p)
    cb = body_box(slide, x, y, w, room, caption, f"{name} caption")
    tf = cb.text_frame
    tf.margin_top, tf.margin_bottom, tf.margin_left, tf.margin_right = em(BOXES["note"]["gap"], CAPTION_PT), 0, 0, 0
    return cb, int(text_height(cb, slide, gap=0) * SLACK)


def _cell(p):
    """A cell's paragraph as pandoc writes one in a table: its ⟦engine:cell⟧ prefix off, no indent, no bullet."""
    t = p.find(".//" + qn("a:t"))
    if t is not None and (t.text or "").startswith(CELL):
        t.text = t.text[len(CELL):]
        r = t.getparent()
        if not t.text and r.getparent() is p and len(p.findall(qn("a:r"))) > 1:
            p.remove(r)
    for old in p.findall(qn("a:pPr")):
        p.remove(old)
    ppr = etree.Element(qn("a:pPr"), lvl="0", indent="0", marL="0")
    etree.SubElement(ppr, qn("a:buNone"))
    p.insert(0, ppr)
    return p


def draw_table(slide, x, y, w, room, attrs, paras, name):
    """The table of a marker's attrs and paragraphs (rows, their cells, the caption) at (x, y), w wide:
    (table frame, caption box or None, height) or None when it has no row."""
    rows, caption = [], []
    for p in paras:
        text = para_text(p).strip()
        if text == ROW and not caption:
            rows.append([])
        elif text.startswith(CELL) and rows and not caption:
            rows[-1].append(_cell(p))
        else:
            caption.append(p)
    cols = int(attrs.get("cols", max((len(r) for r in rows), default=0)))
    if not rows or not cols:
        return None
    frame = slide.shapes.add_table(len(rows), cols, Emu(x), Emu(y), Emu(w), Emu(max(room, 1)))
    frame.name = name
    tbl = frame.table
    tbl.first_row = bool(attrs.get("header"))
    tbl.horz_banding = False
    for c, col in enumerate(tbl.columns):   # equal, in integers summing to w
        col.width = Emu(w * (c + 1) // cols - w * c // cols)
    for r, row in enumerate(rows):
        tbl.rows[r].height = Emu(0)   # pandoc's: PowerPoint grows a row to its text
        for c in range(cols):
            body = tbl.cell(r, c).text_frame._txBody
            for old in body.findall(qn("a:p")):
                body.remove(old)
            body.append(row[c] if c < len(row) else _cell(etree.Element(qn("a:p"))))
    th = table_height(frame, slide)
    cb, cap_h = _caption_box(slide, x, y + th, w, max(room - th, 0), caption, name)
    frame.height = Emu(th)
    if cb is not None:
        cb.left, cb.top, cb.width, cb.height = Emu(x), Emu(y + th), Emu(w), Emu(cap_h)
    if th + cap_h > room:
        print(f"figures: WARNING: slide {slide_title(slide)!r}: {name} is taller than its column's room", file=sys.stderr)
    return frame, cb, th + cap_h
