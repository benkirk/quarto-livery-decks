"""The marking strips (livery.yml marking:) on every slide of a deck.

Each strip's text is `engine-marking-<name>` in docProps/custom.xml, which slide-layout.lua writes
from engine-common.lua's one reader of the deck's `marking:` (a string sets the first strip, `none`
turns every strip off, a map sets strips by name, a strip the deck leaves alone keeps the livery's
default): "none", no strip. This step draws and never decides. A text box in the strip's box, over
everything on the slide: no insets, the text centered on its height in the strip's size, its color
and weight, a solid fill behind it when the livery names one. The template has no placeholder for
them (pandoc fills none it does not know)."""

from pathlib import Path

from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Pt

from . import docprops
from .common import LIVERY

MARKINGS = LIVERY.get("marking") or []   # in order, colors hex and sizes in points: livery_build.markings()
ALIGN = {"center": PP_ALIGN.CENTER, "left": PP_ALIGN.LEFT, "right": PP_ALIGN.RIGHT}


def texts(prs):
    """[(strip, text)] for the strips that have a text (slide-layout.lua's engine-marking-<name>)."""
    out = []
    for s in MARKINGS:
        text = docprops.prop(prs, f"engine-marking-{s['name']}")
        if text and text != "none":
            out.append((s, text))
    return out


def run(prs, path: Path) -> None:
    strips = texts(prs)
    if not strips:
        return
    W, H = prs.slide_width, prs.slide_height
    for slide in prs.slides:
        for s, text in strips:
            x, y, w, h = s["box"]
            box = slide.shapes.add_textbox(int(x * W), int(y * H), int(w * W), int(h * H))
            box.name = f"Marking {s['name']}"
            if s.get("fill"):
                box.fill.solid()
                box.fill.fore_color.rgb = RGBColor.from_string(s["fill"][1:])
            tf = box.text_frame
            tf.word_wrap = True
            tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = tf.paragraphs[0]
            p.alignment = ALIGN[s.get("align", "center")]
            r = p.add_run()
            r.text = text
            r.font.size = Pt(s["size"])
            r.font.color.rgb = RGBColor.from_string(s["color"][1:])
            r.font.bold = bool(s.get("bold"))
    print(f"marking: {len(strips)} marking strip(s) on {len(prs.slides)} slide(s) → {path.name}")
