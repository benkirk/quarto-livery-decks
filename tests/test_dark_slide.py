"""Check the inverted content slide, `{.brand-dark}`: its reader and each format's drawing.

    python3 -m pytest tests/test_dark_slide.py

engine-common.lua's take_layout() reads `.brand-dark` as a layout class (`dark`, the class taken off
the heading) under `pandoc lua` (conftest's lua_common). livery.css on a light livery swaps the
roles a dark slide's text reads and restores them in a box with a fill of its own; a dark livery's
has no swap (the class is a no-op there).
"""

import json
import os
import subprocess
import sys

import pytest

from conftest import ENGINE, ROOT, flat, lua_common, pandoc_exe, slide_layout_shim

# take_layout() on each heading of the markdown from stdin: {dark, classes} or {nil}
HARNESS = r"""
local doc = pandoc.read(io.read("a"), "markdown")
for _, b in ipairs(doc.blocks) do
  if b.t == "Header" then
    local lay = C.take_layout(b, "t")
    io.write(pandoc.json.encode(lay and { dark = lay.dark, classes = b.classes } or { ["nil"] = true }) .. "\n")
  end
end
"""


def test_take_layout_dark(tmp_path):
    """`.brand-dark` is a layout class: dark, the class off the heading, another class kept; with
    .vcenter both; a heading without one nil."""
    out, warnings = lua_common(tmp_path, HARNESS, "## A {.brand-dark .mine}\n\n## B {.brand-dark .vcenter}\n\n## C\n")
    got = [json.loads(x) for x in out.splitlines()]
    assert got == [{"dark": True, "classes": ["mine"]}, {"dark": True, "classes": []}, {"nil": True}], got
    assert not warnings


def test_take_layout_dark_ignored(tmp_path):
    """A .full or .feature slide ignores .brand-dark in every format (the picture fills it): not
    dark, the class taken off, one warning each."""
    out, warnings = lua_common(tmp_path, HARNESS, "## F {.full .brand-dark}\n\n## P {.feature .brand-dark}\n")
    got = [json.loads(x) for x in out.splitlines()]
    assert got == [{"dark": False, "classes": ["full"]}, {"dark": False, "classes": ["feature"]}], got
    assert warnings == ['t: {.brand-dark} on the {.full} slide "F" is ignored',
                        't: {.brand-dark} on the {.feature} slide "P" is ignored'], warnings


def test_css_swaps_on_a_light_livery():
    """ncar: on .engine-dark the text's roles take their dark values (on-dark, dark-muted, secondary,
    the dark table blends), reveal's text colors and the hairline with them; a card, lead, strip,
    callout or caution body gets the livery's own back. nasa-dark (a dark livery): no swap."""
    j = json.loads((ROOT / "liveries" / "ncar" / "build" / "livery.json").read_text())
    css = (ROOT / "liveries" / "ncar" / "build" / "livery.css").read_text()
    c = {**j["roles"], **j["blends"]}
    on = css[css.index("/* the dark slide (.engine-dark)"):]
    on, back = on[:on.index("}")], on[on.index("}") + 1:]
    back = back[:back.index("}")]
    for role, dark in (("foreground", "on-dark"), ("muted", "dark-muted"), ("primary-contrast", "secondary"),
                       ("table-rule", "dark-rule"), ("table-band", "dark-band")):
        assert f"--engine-{role}: {c[dark]};" in on, (role, on)
        assert f"--engine-{role}: {c[role]};" in back, (role, back)
    assert f"--engine-title-color: {c['on-dark']};" in on, "the title on-dark"
    assert f"--r-main-color: {c['on-dark']};" in on and f"--engine-title-rule-color: {c['dark-rule']};" in on
    assert ":is(.engine-card, .engine-lead, .engine-strip, .callout)" in back, "the boxes' restore"
    assert ".engine-dark.engine-caution > .engine-body" in css
    dark = (ROOT / "liveries" / "nasa-dark" / "build" / "livery.css").read_text()
    assert "/* the dark slide (.engine-dark)" not in dark, "a dark livery swaps nothing"


# the pptx: a dark text slide (a list, inline code, a footnote), a dark grid slide (a card beside a
# plain column), a dark table slide, a light slide; through slide-layout.lua (conftest's shim) and
# grid-pptx.lua, then the driver, under the livery
DECK = ("## Dark text {.brand-dark}\n\n- One item\n- Two, with `code`\n\n† A footnote.\n\n"
        "## Dark grid {.brand-dark}\n\n:::: {.columns}\n::: {.column .card}\n### Card\nCard body.\n:::\n"
        "::: {.column}\nPlain column.\n:::\n::::\n\n"
        "## Dark table {.brand-dark}\n\n| a | b |\n|---|---|\n| 1 | 2 |\n| 3 | 4 |\n\n"
        "## Light\n\nText.\n")


@pytest.fixture(scope="module", params=["ncar", "nsi", "nasa-dark"])
def dark_deck(request, tmp_path_factory):
    name = request.param
    d = ROOT / "liveries" / name
    tmp = tmp_path_factory.mktemp(f"dark-{name}")
    out = tmp / "t.pptx"
    subprocess.run([pandoc_exe(), "--lua-filter", str(slide_layout_shim(tmp)), "--lua-filter", str(ENGINE / "grid-pptx.lua"),
                    "--reference-doc", str(d / "template.pptx"), "-o", str(out)],
                   input=f"---\nlivery: {d}\n---\n\n" + DECK, text=True, cwd=tmp, check=True, capture_output=True)
    env = {**os.environ, "LIVERY_DIR": str(d), "PYTHONPATH": str(ENGINE)}
    r = subprocess.run([sys.executable, str(ENGINE / "pptx_post.py"), str(out)], env=env, cwd=tmp, capture_output=True, text=True)
    assert r.returncode == 0, r.stderr
    from pptx import Presentation
    return name, json.loads((d / "build" / "livery.json").read_text()), Presentation(str(out))


def colors(shape):
    """The explicit run colors (six digits; None for a run that takes the master's) of a shape's text
    (a shape or a table cell)."""
    el = shape._element if hasattr(shape, "_element") else shape._tc
    return [(r.find(f"{A}rPr/{A}solidFill/{A}srgbClr").get("val") if r.find(f"{A}rPr/{A}solidFill/{A}srgbClr") is not None else None)
            for r in el.iter(f"{A}r")]


A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def bg(slide):
    clr = slide._element.find(f"{{http://schemas.openxmlformats.org/presentationml/2006/main}}cSld/"
                              f"{{http://schemas.openxmlformats.org/presentationml/2006/main}}bg//{A}srgbClr")
    return clr.get("val") if clr is not None else None


def test_dark_furniture(dark_deck):
    """A dark slide: its background the dark role, the master's shapes hidden, the furniture redrawn at
    livery.json's boxes (the tab in the accent, the rule in dark-rule; nsi has neither), the dark role's
    marks (nsi's white lockup and sign-off, ncar primary's), the title on-dark; the light slide as it
    was. Under nasa-dark (a dark livery) the class is a no-op: no background, the master shown."""
    name, j, prs = dark_deck
    h = lambda v: j["roles"].get(v, j["blends"].get(v, "#"))[1:].upper()
    W, H = prs.slide_width, prs.slide_height
    dark, light = list(prs.slides)[:3], prs.slides[3]
    if name == "nasa-dark":
        assert all(bg(s) is None and s._element.get("showMasterSp") != "0" for s in dark), "a dark livery's no-op"
        return
    for s in dark:
        assert bg(s) == h("dark") and s._element.get("showMasterSp") == "0", s.shapes.title.text_frame.text
        by = {sh.name: sh for sh in s.shapes}
        fu = j["furniture"]
        assert ("Dark tab" in by) == bool(fu["tab"]) and ("Dark rule" in by) == bool(fu["rule"]), "the furniture's presence"
        if fu["tab"]:
            t = by["Dark tab"]
            assert (t.left, t.top, t.width, t.height) == tuple(int(v * d) for v, d in zip(fu["tab"], (W, H, W, H)))
            assert str(t.fill.fore_color.rgb) == h("accent")
        if fu["rule"]:
            assert str(by["Dark rule"].fill.fore_color.rgb) == h("dark-rule")
        logo = by["Dark logo"]
        assert logo.image.blob == (ROOT / "liveries" / name / j["marks"]["dark"]["logo"]).read_bytes(), "the dark role's mark"
        if name == "nsi":
            assert j["marks"]["dark"]["logo"].endswith("nsi-lockup-white.png") and "Dark signoff" in by
        assert set(colors(s.shapes.title)) == {h("on-dark")}, "the title on-dark"
    assert bg(light) is None and light._element.get("showMasterSp") != "0", "the light slide"


def test_dark_text(dark_deck):
    """The text slide: its list on-dark, the footnote in dark-muted (the marker in the accent)."""
    name, j, prs = dark_deck
    if name == "nasa-dark":
        return
    h = lambda v: j["roles"].get(v, j["blends"].get(v, "#"))[1:].upper()
    body = next(sh for sh in prs.slides[0].placeholders if sh.placeholder_format.idx == 1)
    paras = body.text_frame.paragraphs
    items = [p for p in paras if p.text.startswith(("One", "Two"))]
    assert all(set(colors_of(p)) == {h("on-dark")} for p in items), [colors_of(p) for p in items]
    foot = next(p for p in paras if "A footnote" in p.text)
    assert h("dark-muted") in colors_of(foot) and h("accent") in colors_of(foot), colors_of(foot)


def colors_of(p):
    return [(r._r.find(f"{A}rPr/{A}solidFill/{A}srgbClr").get("val") if r._r.find(f"{A}rPr/{A}solidFill/{A}srgbClr") is not None else None)
            for r in p.runs]


def test_dark_grid(dark_deck):
    """The grid slide: the plain column's text on-dark; the card keeps its own colors (its body takes
    the master's, its title primary-text on its band)."""
    name, j, prs = dark_deck
    if name == "nasa-dark":
        return
    h = lambda v: j["roles"].get(v, j["blends"].get(v, "#"))[1:].upper()
    by = {sh.name: sh for sh in flat(prs.slides[1].shapes)}
    assert set(colors(by["Column 2"])) == {h("on-dark")}, colors(by["Column 2"])
    assert h("on-dark") not in colors(by["Card 1 body"]), colors(by["Card 1 body"])
    assert set(colors(by["Card 1 title"])) == {h("primary-text")}


def test_dark_table(dark_deck):
    """The table slide: the head rule in secondary, the row rules in dark-rule, a band in dark-band,
    every cell's text on-dark."""
    name, j, prs = dark_deck
    if name == "nasa-dark":
        return
    h = lambda v: j["roles"].get(v, j["blends"].get(v, "#"))[1:].upper()
    t = next(sh for sh in prs.slides[2].shapes if sh.has_table).table
    ln = lambda tc, side: tc._tc.find(f"{A}tcPr/{A}{side}/{A}solidFill/{A}srgbClr")
    assert ln(t.cell(0, 0), "lnB").get("val") == h("secondary"), "the head rule"
    assert ln(t.cell(1, 0), "lnB").get("val") == h("dark-rule"), "a row rule"
    band = t.cell(2, 0)._tc.find(f"{A}tcPr/{A}solidFill/{A}srgbClr")
    assert band is not None and band.get("val") == h("dark-band"), "the band"
    assert all(set(colors(t.cell(r, c))) == {h("on-dark")}
               for r in range(3) for c in range(2)), "the cells' text"
