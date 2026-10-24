"""Check the watermark (a deck's `watermark:`, a livery's `watermark:` look): its data from the build,
the deck's word on it (engine-common.lua), and the HTML filter's element.

    python3 -m pytest tests/test_watermark.py

The build resolves the look over WATERMARK (color to hex, size to the half point, the default size
scaled by the canvas) into every livery.json, livery.css and livery.sty, so a deck's word shows
under any livery. engine-common.lua's watermark_text() is the one reader: the deck's text, else the
livery's default, `none` for none. Needs pandoc ($QUARTO_PANDOC) for the reader.
"""

import json
import re

import pytest
from pptx import Presentation
from pptx.oxml.ns import qn

import livery_build as lb

from conftest import ROOT, lua_common, run_steps, slide_layout_shim


def test_defaults_in_every_livery(livery):
    j = json.loads((livery / "build" / "livery.json").read_text())
    w = j["watermark"]
    assert w["text"] is None and w["opacity"] == 0.08 and w["angle"] == 30, "the defaults: no text, faint, 30 degrees"
    assert w["color"] == j["roles"]["foreground"], "the foreground, as hex"
    assert w["size"] == lb.half(120 * j["type"]["canvas"]["h"] / 540), "120 pt on the default canvas, scaled to the livery's"
    css, sty = (livery / "build" / "livery.css").read_text(), (livery / "build" / "livery.sty").read_text()
    assert f"--engine-watermark-color: {w['color']};" in css and "--engine-watermark-angle: 30deg;" in css
    assert re.search(r"--engine-watermark-size: calc\(200px \* var\(--engine-type-scale, 1\)\);", css), \
        "0.22 of the slide's height in every livery"
    assert f"\\definecolor{{EngineWatermarkColor}}{{HTML}}{{{w['color'][1:]}}}" in sty
    assert "\\def\\engine@watermark@angle{30}" in sty and "\\def\\engine@watermark@default{}" in sty


def test_schema_defaults():
    s = json.loads((ROOT / "_extensions" / "engine" / "livery.schema.json").read_text())["properties"]["watermark"]["properties"]
    for k, v in lb.WATERMARK.items():
        if k != "text":
            assert s[k]["default"] == v, f"watermark.{k}'s default is not WATERMARK's {v}"


@pytest.mark.parametrize("block, check", [
    ("  text: DRAFT\n  color: primary\n", lambda w, j: w["text"] == "DRAFT" and w["color"] == j["roles"]["primary"]),
    ("  color: \"#FF0000\"\n  size: 12.3\n", lambda w, j: w["color"] == "#FF0000" and w["size"] == 12.5),
    ("  color: rule\n  opacity: 0.2\n  angle: -45\n", lambda w, j: w["color"] == j["blends"]["rule"] and (w["opacity"], w["angle"]) == (0.2, -45)),
])
def test_livery_look(tmp_path, tmp_livery, block, check):
    rc, err, j = tmp_livery(tmp_path, block, key="watermark")
    assert rc == 0, err
    assert check(j["watermark"], j), j["watermark"]
    if "text: DRAFT" in block:
        assert "\\def\\engine@watermark@default{DRAFT}" in (tmp_path / "liveries" / "ncar" / "build" / "livery.sty").read_text()


@pytest.mark.parametrize("block, why", [
    ("  angle: 120\n", "watermark.angle: 120 is greater than the maximum of 90"),
    ("  opacity: 2\n", "watermark.opacity: 2 is greater than the maximum of 1"),
    ("  font: Poppins\n", "watermark.font: unknown key"),
])
def test_refusals(tmp_path, tmp_livery, block, why):
    rc, err, _ = tmp_livery(tmp_path, block, key="watermark")
    assert rc != 0 and why in err, err


READER = """
local doc = pandoc.read(io.read("a"), "markdown")
local t = C.watermark_text(doc.meta, pandoc.json.decode(arg[2]))
print(t and pandoc.utils.stringify(t) or "-")
if t then print(#t, t[1].t) end
"""


@pytest.mark.parametrize("front, default, want", [
    ("watermark: DRAFT\n", None, ["DRAFT", "1", "Str"]),           # the deck's word
    ("watermark: \"*Review copy*\"\n", None, ["Review copy", "1", "Emph"]),   # inlines kept
    ("watermark: none\n", "SAMPLE", ["-"]),                       # none drops a default
    ("", "SAMPLE", ["SAMPLE", "1", "Str"]),                       # the livery's default
    ("", None, ["-"]),                                            # neither: none
])
def test_reader(tmp_path, front, default, want):
    md = f"---\n{front}---\n\nx\n" if front else "x\n"
    out, warnings = lua_common(tmp_path, READER, md, json.dumps({"watermark": {"text": default}}))
    lines = out.strip().splitlines()
    assert lines[0] == want[0], out
    assert [l.split() for l in lines[1:]] == ([want[1:]] if len(want) > 1 else []), "one inline, its kind"
    assert not warnings


# the pptx half: slide-layout.lua writes the text (engine-watermark), the step draws it
STEPS = ["slide_layout", "marking", "watermark", "type_step"]   # the driver's order around it


def pptx_marks(pandoc_deck, tmp, front, body="## a\n\nOne.\n\n## b {.brand-dark}\n\nTwo.\n", name="ncar"):
    """Each slide's Watermark shapes after the steps: (index in the tree, rot, size pt, color, alpha), and
    each slide's first text shape's index; under the livery name."""
    d = ROOT / "liveries" / name
    md = f"---\nlivery: {d}\n{front}---\n\n{body}"
    pandoc_deck(tmp, md, "t.pptx", "--lua-filter", str(slide_layout_shim(tmp)))
    run_steps(tmp / "t.pptx", STEPS, d)
    out = []
    for s in Presentation(str(tmp / "t.pptx")).slides:
        shapes = list(s.shapes)
        wm = [(i, sh) for i, sh in enumerate(shapes) if sh.name == "Watermark"]
        first = next((i for i, sh in enumerate(shapes) if sh.name != "Watermark" and sh.has_text_frame and sh.text_frame.text.strip()), None)
        rows = []
        for i, sh in wm:
            run = sh.text_frame.paragraphs[0].runs[0]
            clr = run._r.find(qn("a:rPr")).find(".//" + qn("a:srgbClr"))
            rows.append((i, sh._element.find(".//" + qn("a:xfrm")).get("rot"), run.font.size.pt, clr.get("val"),
                         clr.find(qn("a:alpha")).get("val"), sh.text_frame.text, run.font.bold))
        out.append((rows, first))
    return out


def test_pptx_every_slide(pandoc_deck, tmp_path):
    j = json.loads((ROOT / "liveries" / "ncar" / "build" / "livery.json").read_text())
    got = pptx_marks(pandoc_deck, tmp_path, "watermark: DRAFT\n")
    assert len(got) == 2 and all(len(rows) == 1 for rows, _ in got), got
    for (rows, first), ink in zip(got, (j["roles"]["foreground"], j["roles"]["on-dark"])):
        i, rot, pt, color, alpha, text, bold = rows[0]
        assert (rot, alpha, text, bold) == (str(330 * 60000), "8000", "DRAFT", True), rows
        assert pt == j["watermark"]["size"], "the livery's size, in points"
        assert color == ink.lstrip("#").upper(), "the livery's color; on-dark on the dark slide"
        assert first is not None and i < first, "behind the slide's first text (its title)"


def test_pptx_none_and_default(pandoc_deck, tmp_path):
    assert all(rows == [] for rows, _ in pptx_marks(pandoc_deck, tmp_path, "")), "no word, no default: nothing"
    (tmp_path / "n").mkdir()
    assert all(rows == [] for rows, _ in pptx_marks(pandoc_deck, tmp_path / "n", "watermark: none\n")), "none: nothing"


def test_pptx_scaled_by_the_type_step(pandoc_deck, tmp_path):
    import livery as livery_mod
    j = json.loads((ROOT / "liveries" / "ncar" / "build" / "livery.json").read_text())
    want = livery_mod.half_pt(j["watermark"]["size"] * 10 / j["type"]["roles"]["body"]["pt"])
    got = pptx_marks(pandoc_deck, tmp_path, "watermark: DRAFT\ntype-body: 10\n")
    assert all(rows[0][2] == want for rows, _ in got), (want, got)


def test_pptx_field_color(pandoc_deck, tmp_path):
    """A divider wears its composition's text color (white on ncar's blue), as HTML and the PDF."""
    j = json.loads((ROOT / "liveries" / "ncar" / "build" / "livery.json").read_text())
    got = pptx_marks(pandoc_deck, tmp_path, "watermark: DRAFT\n", body="# Part one\n\n## a\n\nOne.\n")
    want = j["roles"][j["compositions"]["section"]["column"]["role"]].lstrip("#").upper()
    assert got[0][0][0][3] == want and got[1][0][0][3] == j["roles"]["foreground"].lstrip("#").upper(), got
