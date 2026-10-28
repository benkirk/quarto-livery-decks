"""Check a list after a slide's one table or picture in the pptx (single-body.lua's fold, then
pptx_steps/grid.py), and a list in a table's cell (pptx_steps/style_tables.py).

    python3 -m pytest tests/test_list_after.py

pandoc's pptx writer moves anything after a table or picture onto an untitled continuation slide.
single-body.lua folds a bullet or numbered list there into the caption, as `⟦engine:list⟧` marker
paragraphs (a nested list a span of its own), with the paragraphs beside it; grid.py draws the list
as the body's list under the table or picture: one "Body list" text box carrying the master's body
style, each item a paragraph at its level, no marker left, the lines beside it left-aligned. A list
item holding more than a line (code, a second paragraph) still splits the slide. A list in a grid
table's cell, which pandoc writes as leveled paragraphs a cell gives no bullet, takes the body
list's bullet and indent; the cell's plain line keeps none. single-body.lua runs under plain pandoc
with a `quarto` shim, the steps in a subprocess under ncar's LIVERY_DIR. Needs pandoc
($QUARTO_PANDOC).
"""

import shutil
import subprocess

import pytest

from conftest import ENGINE, ROOT, flat, pandoc_exe, run_steps

NCAR = ROOT / "liveries" / "ncar"
IMAGE = ROOT / "decks" / "quickstart" / "images" / "example.jpg"
TABLE = "| a | b |\n|---|---|\n| 1 | 2 |\n| 3 | 4 |\n"
MD = ("## List under a table\n\n" + TABLE + "\n**The point.** A line before the list:\n\n- One\n- Two\n  - Two, under it\n\n"
      "† A footnote.\n\n"
      "## Numbers under a picture\n\n![](p.jpg){height=40%}\n\n3. Three\n4. Four\n\n"
      "## Code in an item\n\n" + TABLE + "\n- One\n\n  ```\n  code\n  ```\n\n"
      "## Bullets in a cell\n\n+-------+------------+\n| Who   | What       |\n+=======+============+\n"
      "| Staff | - One      |\n|       | - Two      |\n+-------+------------+\n")
SHIM = ('quarto = {{ doc = {{ is_format = function(f) return f == "pptx" end }},\n'
        '           utils = {{ resolve_path = function(p) return "{engine}/" .. p end }},\n'
        '           log = {{ warning = function(s) io.stderr:write("WARNING: " .. s .. "\\n") end }},\n'
        '           json = {{ decode = pandoc.json.decode }} }}\n'
        'dofile("{engine}/single-body.lua")\n')


@pytest.fixture(scope="module")
def deck(tmp_path_factory):
    tmp = tmp_path_factory.mktemp("list-after")
    shutil.copy(IMAGE, tmp / "p.jpg")
    shim = tmp / "single-body-shim.lua"
    shim.write_text(SHIM.format(engine=ENGINE))
    out = tmp / "t.pptx"
    subprocess.run([pandoc_exe(), "--lua-filter", str(shim), "--reference-doc", str(NCAR / "template.pptx"),
                    "-o", str(out)], input=MD, text=True, cwd=tmp, check=True, capture_output=True)
    run_steps(out, ["style_footnotes", "slide_layout", "columns", "grid", "style_tables"], NCAR)
    from pptx import Presentation
    return Presentation(str(out))


def body_list(slide):
    lists = [sh for sh in flat(slide.shapes) if sh.name == "Body list"]
    assert len(lists) == 1, [sh.name for sh in flat(slide.shapes)]
    return lists[0]


def levels(shape):
    return [(p.text, int(p._p.pPr.get("lvl", "0"))) for p in shape.text_frame.paragraphs]


def test_slides(deck):
    """The two folded slides stay one slide each; the item holding code still splits its slide."""
    titles = [s.shapes.title.text if s.shapes.title else "" for s in deck.slides]
    assert titles[:3] == ["List under a table", "Numbers under a picture", "Code in an item"], titles
    assert len(titles) == 5 and titles[3] == "", "the code item's continuation slide"


def test_list_under_a_table(deck):
    """The list in a box of the body's list style under the table, each item at its level, bulleted
    by the box's style (no buNone, no caption centering); the line before it left-aligned, above it."""
    from pptx.oxml.ns import qn
    slide = deck.slides[0]
    table = next(sh for sh in slide.shapes if sh.has_table)
    box = body_list(slide)
    rows = sum(r.height for r in table.table.rows)   # pandoc's frame is taller than its rows
    assert table.top + rows <= box.top < table.top + table.height, "the list under the table's rows"
    assert levels(box) == [("One", 0), ("Two", 0), ("Two, under it", 1)], levels(box)
    for p in box.text_frame.paragraphs:
        ppr = p._p.pPr
        assert ppr.find(qn("a:buNone")) is None and ppr.get("algn") is None, "the caption's look dropped"
    assert box._element.find(".//" + qn("a:lstStyle") + "/" + qn("a:lvl1pPr")) is not None, "the master's body style"
    lead = next(sh for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.startswith("The point."))
    assert lead.text_frame.paragraphs[0]._p.pPr.get("algn") == "l", "the line beside the list left-aligned"
    assert lead.top < box.top, "the line before the list above it"
    assert "⟦" not in slide._element.xml, "a marker left"


def test_numbers_under_a_picture(deck):
    """A numbered list under a picture: numbered from its start."""
    from pptx.oxml.ns import qn
    slide = deck.slides[1]
    pic = next(sh for sh in slide.shapes if sh.shape_type == 13)
    box = body_list(slide)
    assert box.top >= pic.top + pic.height, "the list under the picture"
    nums = [p._p.pPr.find(qn("a:buAutoNum")) for p in box.text_frame.paragraphs]
    assert all(n is not None and n.get("startAt") == "3" for n in nums), "numbered from 3"
    assert "⟦" not in slide._element.xml, "a marker left"


def test_bullets_in_a_cell(deck):
    """The list's items bulleted and indented as the body's first level; the plain cell no bullet."""
    from pptx.oxml.ns import qn
    table = next(sh for sh in deck.slides[4].shapes if sh.has_table).table
    staff, items = table.cell(1, 0), table.cell(1, 1)
    for p in items.text_frame.paragraphs:
        ppr = p._p.pPr
        assert ppr.find(qn("a:buChar")) is not None and int(ppr.get("marL", "0")) > 0, p.text
    assert [p.text for p in items.text_frame.paragraphs] == ["One", "Two"]
    assert staff.text_frame.paragraphs[0]._p.pPr.find(qn("a:buChar")) is None, "no bullet on a plain line"
