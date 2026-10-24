"""Check a columns slide's frames in the pptx (grid-pptx.lua, then pptx_steps/columns.py).

    python3 -m pytest tests/test_grid.py

pandoc writes a deck on the livery's template with grid-pptx.lua (slide-layout.lua's `wide` line
written by hand: that filter needs Quarto), then the driver's steps run in a subprocess under the
livery's LIVERY_DIR. A 55/45 slide's placeholders are the shares of the content box less the gap,
within 1 EMU; a picture pandoc put at the old half frame moves into its column and scales to its
width, its aspect kept; a {.wide} slide's box starts at geometry.wide-margin, its Title and Content
body and a table too (its grid columns scaled); every offset, extent and size is an integer EMU. Under ncar and nasa-jsc-eg (another
canvas, a dark livery). Needs pandoc ($QUARTO_PANDOC).
"""

import json
import os
import re
import shutil
import subprocess
import sys
import zipfile

import pytest

from conftest import ENGINE, ROOT, flat, pandoc_exe

IMAGE = ROOT / "decks" / "quickstart" / "images" / "example.jpg"
CARDS = ('## Cards\n\n:::: {.columns}\n'
         '::: {.column .card .pick role="primary"}\n### OPTION 1\n[Procure a system]{.sub}\n\n'
         '**Total**\n[LOCAL INVESTMENT]{.label}\n[$20M]{.big}\n\n::: {.strip role="accent"}\nTo the consortium\n:::\n\n'
         '- one\n- two\n- three\n:::\n'
         '::: {.column .card role="secondary"}\n### OPTION 2\nShort.\n:::\n'
         '::: {.column .card .solid role="accent"}\n### OPTION 3\nSolid.\n:::\n::::\n')
DEFAULTS = ('## Default roles\n\n:::: {.columns}\n' + ''.join(f'::: {{.column .card}}\n### {t}\nText.\n:::\n' for t in 'ABCD')
            + '::::\n')
WIDE = "\n::: notes\nengine-layout: wide\n:::\n"
RAILS = ('## Steps\n\n::: {.steps role="secondary"}\n1. **Propose** A one-page case.\n2. **Pilot** Six weeks.\n'
         '   A second line.\n3. **Scale** The allocation.\n:::\n\n'
         '## Timeline\n\n::: {.timeline}\n- [x] **Q1 2025** Accepted\n- **2026** Full allocation\n:::\n\n'
         '## Later steps\n\n::: {.steps}\n4. **Review** The board reads it.\n5. **Decide** Yes or no.\n:::\n')
QUOTES = ('## Quote\n\n::: {.quote role="accent"}\nThe purpose of computing\nis insight.\n— R. Hamming, 1962\n:::\n\n'
          '## Quote in a card\n\n:::: {.columns}\n::: {.column .card}\n### OPTION 1\nAbove.\n\n> Quoted in a card.\n\nBelow.\n:::\n'
          '::: {.column .card}\n### OPTION 2\nOther.\n:::\n::::\n\n'
          '## Quote in a column\n\n:::: {.columns}\n::: {.column}\n- one\n- two\n:::\n::: {.column}\nBefore.\n\n'
          '> In a plain column.\n\nAfter.\n:::\n::::\n')
AGENDA = ('# One\n\n## Where we are {.agenda}\n\n# Two\n\n## Two body\n\nText.\n\n'
          '# Backup {.unnumbered}\n\n## Backup body\n\nText.\n')   # slides 15 to 20
# a lead after a figure, as single-body.lua folds it into the figure's caption (pandoc: one centered
# paragraph, a:br between the lines, the markers inline); a lead before a table (pandoc's Content with
# Caption), the note folded into the table's caption, the notes line single-body.lua writes for it
FOLDED = ('## Lead under a figure\n\n![⟦engine:lead⟧\\\n**The headline** of it\\\nA plain line.\\\n⟦engine:end⟧](p.jpg)\n\n'
          '## Lead over a table\n\n::: {.lead}\n**The headline** of it\n:::\n\n| a | b | c | d |\n|---|---|---|---|\n| 1 | 2 | 3 | 4 |\n\n'
          'Table: Millions of core-hours.\n\n::: notes\nengine-layout: note\n:::\n')   # slides 21 and 22
MD = ('## Split\n\n:::: {.columns}\n::: {.column width="55%"}\n- One\n- Two\n:::\n::: {.column}\n![](p.jpg)\n:::\n::::\n\n'
      '## Wide\n\n:::: {.columns}\n::: {.column}\nLeft\n:::\n::: {.column}\nRight\n:::\n::::\n' + WIDE + '\n'
      '## Wide text\n\nA paragraph.\n' + WIDE + '\n'
      '## Three\n\n:::: {.columns}\n::: {.column}\n- AAA\n- aaa\n:::\n::: {.column}\nBBB\n:::\n::: {.column}\nCCC\n:::\n::::\n\n'
      '## Lead\n\n::: {.lead}\n**The headline**\nA plain line under it.\n:::\n\n'
      ':::: {.columns}\n::: {.column}\nLLL\n:::\n::: {.column}\nMMM\n:::\n::::\n\n'
      '## Lead alone\n\n::: {.lead}\nOnly a lead.\n:::\n\n' + CARDS + '\n' + DEFAULTS + '\n'
      '## Wide table\n\n| a | b | c |\n|---|---|---|\n| 1 | 2 | 3 |\n' + WIDE + '\n' + RAILS + '\n' + QUOTES + '\n' + AGENDA + '\n' + FOLDED)
STEP = """
import sys
from pathlib import Path
from pptx import Presentation
from pptx_steps import columns, grid, slide_layout
p = Path(sys.argv[1]); prs = Presentation(p)
slide_layout.run(prs, p); columns.run(prs, p); grid.run(prs, p); prs.save(p)
"""   # the driver's order


@pytest.fixture(scope="module", params=["ncar", "nasa-jsc-eg"])
def deck(request, tmp_path_factory):
    name = request.param
    pandoc = pandoc_exe()
    tmp = tmp_path_factory.mktemp(f"grid-{name}")
    shutil.copy(IMAGE, tmp / "p.jpg")
    d = ROOT / "liveries" / name
    out = tmp / "t.pptx"
    subprocess.run([pandoc, "--lua-filter", str(ENGINE / "grid-pptx.lua"), "--reference-doc", str(d / "template.pptx"),
                    "-o", str(out)], input=MD, text=True, cwd=tmp, check=True, capture_output=True)
    env = {**os.environ, "LIVERY_DIR": str(d), "PYTHONPATH": str(ENGINE)}
    subprocess.run([sys.executable, "-c", STEP, str(out)], env=env, check=True, capture_output=True)
    from pptx import Presentation
    return json.loads((d / "build" / "livery.json").read_text()), Presentation(str(out)), out


def boxes(livery, width, shares, wide=False):
    g = livery["geometry"]
    m, gap = g["wide-margin" if wide else "margin"], livery["boxes"]["columns"]["gap"]
    room, x, out = 1 - 2 * m - gap * (len(shares) - 1), m, []
    for s in shares:
        out.append((x * width, room * s * width))
        x += room * s + gap
    return out


def body_placeholders(slide):
    return {sh.placeholder_format.idx: sh for sh in slide.placeholders if sh.placeholder_format.idx in (1, 2)}


def test_shares(deck):
    livery, prs, _ = deck
    slide = prs.slides[0]
    (x1, w1), (x2, w2) = boxes(livery, prs.slide_width, [0.55, 0.45])
    ph = body_placeholders(slide)
    assert abs(ph[1].left - x1) <= 1 and abs(ph[1].width - w1) <= 1, f"left column {ph[1].left, ph[1].width}, want {x1, w1}"
    pic = next(sh for sh in slide.shapes if sh.shape_type == 13)
    assert abs(pic.width - w2) <= 1 and abs(pic.left - x2) <= 1, f"the picture {pic.left, pic.width}, want {x2, w2}"
    from PIL import Image
    iw, ih = Image.open(IMAGE).size
    assert abs(pic.height / pic.width - ih / iw) < 0.01, "the picture's aspect"
    assert not slide.notes_slide.notes_text_frame.text.strip().startswith("engine-columns"), "the line left in the notes"


def test_wide(deck):
    livery, prs, _ = deck
    (x1, w1), (x2, w2) = boxes(livery, prs.slide_width, [0.5, 0.5], wide=True)
    ph = body_placeholders(prs.slides[1])
    assert abs(ph[1].left - x1) <= 1 and abs(ph[2].left - x2) <= 1 and abs(ph[2].width - w2) <= 1, \
        f"wide columns {[(p.left, p.width) for p in ph.values()]}, want {x1, x2, w2}"
    m = livery["geometry"]["wide-margin"]
    body = body_placeholders(prs.slides[2])[1]
    assert abs(body.left - m * prs.slide_width) <= 1 and abs(body.width - (1 - 2 * m) * prs.slide_width) <= 1, \
        "a wide Title and Content body"
    assert body.top > 0 and body.height > 0, "the body's top and height kept (all four fields set)"
    table = next(sh for sh in prs.slides[8].shapes if getattr(sh, "has_table", False))   # a graphicFrame, not a body
    assert abs(table.left - m * prs.slide_width) <= 1 and abs(table.width - (1 - 2 * m) * prs.slide_width) <= 1, \
        f"a wide table {table.left, table.width}: pandoc's frame kept"
    cols = table.table._tbl.findall("{*}tblGrid/{*}gridCol")
    assert abs(sum(int(c.get("w")) for c in cols) - table.width) <= len(cols), "the grid columns not scaled to the frame"


def test_integers(deck):
    *_, out = deck
    with zipfile.ZipFile(out) as z:
        xml = "".join(z.read(n).decode() for n in z.namelist() if re.match(r"ppt/slides/slide\d+\.xml$", n))
    floats = re.findall(r'\b(x|y|cx|cy|w|h|sz|marL|indent|adj)="-?[0-9]+\.[0-9]+"', xml)
    assert not floats, f"non-integer EMUs: {floats[:5]}"


def test_three_columns(deck):
    """A three-column slide: three text boxes at the formula's frames, the placeholder gone, its
    list keeping the template's bullet (the box's list style is the master's body style)."""
    livery, prs, _ = deck
    slide = prs.slides[3]
    cols = sorted((sh for sh in slide.shapes if re.fullmatch(r"Column \d+", sh.name)), key=lambda sh: sh.name)
    assert [c.name for c in cols] == ["Column 1", "Column 2", "Column 3"], [sh.name for sh in slide.shapes]
    for c, (x, w) in zip(cols, boxes(livery, prs.slide_width, [0.3333, 0.3333, 0.3333])):
        assert abs(c.left - x) <= 2 * len(cols) and abs(c.width - w) <= 2 * len(cols), f"{c.name} {c.left, c.width}, want {x, w}"
    assert [c.text_frame.text.split("\n")[0] for c in cols] == ["AAA", "BBB", "CCC"], "the columns' text"
    assert not body_placeholders(slide), "the emptied placeholder left on the slide"
    xml = slide._element.xml
    assert "⟦" not in xml, "a marker left"
    assert cols[0]._element.find(".//{*}lstStyle/{*}lvl1pPr") is not None, "no list style in the box"


def test_column_rules(deck):
    """boxes.columns.rule: one rule in each gap where the livery sets one (nasa-jsc-eg), none where
    it is 0 (ncar); centered in the gap, its width the fraction of the slide's height."""
    livery, prs, _ = deck
    w = round(livery["boxes"]["columns"]["rule"] * prs.slide_height)
    counts = [sum(sh.name == "Column rule" for sh in s.shapes) for s in prs.slides]
    if w == 0:
        assert counts == [0] * 23, f"rules under a livery without one: {counts}"
        return
    assert counts == [1, 1, 0, 2, 1, 0, 2, 3, 0, 0, 0, 0, 0, 1, 1, 0, 0, 0, 0, 0, 0, 0, 0], f"rules per slide {counts}, want one per gap"
    slide = prs.slides[3]
    cols = sorted((sh for sh in slide.shapes if re.fullmatch(r"Column \d+", sh.name)), key=lambda sh: sh.name)
    for rule, (a, b) in zip(sorted((sh for sh in slide.shapes if sh.name == "Column rule"), key=lambda sh: sh.left), zip(cols, cols[1:])):
        mid = (a.left + a.width + b.left) / 2
        assert rule.width == w and abs(rule.left + rule.width / 2 - mid) <= 1, f"a rule at {rule.left}, {rule.width}; the gap's middle {mid}"
        assert rule.height > 0 and rule.top == a.top, "the rule runs down the row from its top"


def test_lead(deck):
    """A lead over two columns: its text box on the lead fill with a bar of bar-color at its left,
    cut to its text, its source lines two paragraphs; the columns below it, the slide a grid slide
    (two text boxes, no placeholder). A lead alone is drawn the same way."""
    livery, prs, _ = deck
    lead = livery["boxes"]["lead"]
    for n, text in ((4, ["The headline", "A plain line under it."]), (5, ["Only a lead."])):
        slide = prs.slides[n]
        by = {sh.name: sh for sh in flat(slide.shapes)}
        assert {"Lead", "Lead fill", "Lead bar", "Lead text"} <= set(by), f"slide {n + 1}: {sorted(by)}"
        assert by["Lead"].shape_type == 6 and len(by["Lead"].shapes) == 3, "the lead one group of its three shapes"
        assert str(by["Lead fill"].fill.fore_color.rgb) == lead["fill"][1:], "the lead's fill"
        assert str(by["Lead bar"].fill.fore_color.rgb) == lead["bar-color"][1:], "the bar's color"
        assert by["Lead bar"].width == round(lead["bar"] * prs.slide_height), "the bar's width"
        assert by["Lead fill"].height == by["Lead text"].height < prs.slide_height / 3, "the lead cut to its text"
        assert [p.text for p in by["Lead text"].text_frame.paragraphs] == text, "a source line a paragraph"
        assert "⟦" not in slide._element.xml, "a marker left"
    by = {sh.name: sh for sh in flat(prs.slides[4].shapes)}
    assert by["Column 1"].top >= by["Lead fill"].top + by["Lead fill"].height, "the columns under the lead"
    assert not body_placeholders(prs.slides[4]), "the emptied placeholder left"


def test_lead_under_a_figure(deck):
    """A lead after a figure reaches the pptx as the figure's caption (single-body.lua's fold): one
    centered paragraph, its lines parted by a:br, the markers inline. The step parts the lines, drops
    the fold's centering, and draws the lead under the picture: its group, its lines a paragraph each,
    left-aligned, the box starting below the picture; no marker left, no 'markers left' warning."""
    from pptx.enum.text import PP_ALIGN
    livery, prs, _ = deck
    slide = prs.slides[21]
    by = {sh.name: sh for sh in flat(slide.shapes)}
    assert {"Lead", "Lead fill", "Lead bar", "Lead text"} <= set(by), sorted(by)
    pic = next(sh for sh in slide.shapes if sh.shape_type == 13)
    assert by["Lead fill"].top >= pic.top + pic.height, "the lead under the picture"
    paras = by["Lead text"].text_frame.paragraphs
    assert [p.text for p in paras] == ["The headline of it", "A plain line."], "a caption line a paragraph"
    assert all(p.alignment in (None, PP_ALIGN.LEFT) for p in paras), "the fold's centering dropped"
    assert "⟦" not in slide._element.xml, "a marker left"


def test_lead_over_a_table(deck):
    """A lead before a table: pandoc lays the slide out on Content with Caption, the lead in the
    half-width side box beside the table. The step moves the slide onto Title and Content, draws
    the lead at the body's width, fits the table under it at that width, and anchors the note (the
    table's caption, pandoc's) at the body's foot under the table; no marker, no notes line left."""
    livery, prs, _ = deck
    slide = prs.slides[22]
    assert slide.slide_layout.name == "Title and Content", slide.slide_layout.name
    by = {sh.name: sh for sh in flat(slide.shapes)}
    assert {"Lead", "Lead fill", "Lead text"} <= set(by), sorted(by)
    table = next(sh for sh in slide.shapes if sh.has_table)
    fill = by["Lead fill"]
    assert abs(table.left - fill.left) <= 1 and abs(table.width - fill.width) <= 1, "the table as wide as the lead (the body)"
    assert table.top >= fill.top + fill.height, "the table under the lead"
    note = next(sh for sh in slide.shapes if sh.has_text_frame and "core-hours" in sh.text_frame.text)
    assert note.top >= table.top + table.height, "the note under the table"
    assert note.top + note.height <= round(livery["geometry"]["body-bottom"] * prs.slide_height) + 1, "the note inside the body"
    assert "⟦" not in slide._element.xml, "a marker left"
    assert "engine-layout" not in slide.notes_slide.notes_text_frame.text, "the notes line left"


def test_cards(deck):
    """Three cards: at the formula's frames, one height; each head band in its role, the title
    run bold in the subhead size and primary-text; the pick outlined in outline-color; the solid
    card's fill its role; the label in caps, the big number in the display size; no marker left."""
    import livery as livery_mod
    livery, prs, _ = deck
    slide = prs.slides[6]
    by = {sh.name: sh for sh in flat(slide.shapes)}
    cards = [by[f"Card {n} fill"] for n in (1, 2, 3)]
    for n in (1, 2, 3):   # each card one group: its fill, head or outline, text and strips inside
        g = by[f"Card {n}"]
        assert g.shape_type == 6 and f"Card {n} fill" in {sh.name for sh in flat(g.shapes)}, f"Card {n} is not a group of its shapes"
    for c, (x, w) in zip(cards, boxes(livery, prs.slide_width, [0.3333, 0.3333, 0.3333])):
        assert abs(c.left - x) <= 6 and abs(c.width - w) <= 6, f"{c.name} {c.left, c.width}, want {x, w}"
    assert len({c.height for c in cards}) == 1 and len({c.top for c in cards}) == 1, "the cards one height"
    role = lambda r: livery["roles"][r][1:]
    assert str(by["Card 1 head"].fill.fore_color.rgb) == role("primary"), "the head band in its role"
    assert str(by["Card 2 head"].fill.fore_color.rgb) == role("secondary")
    geom = lambda sh: sh._element.find(".//{*}prstGeom").get("prst")   # rounded cards: the head square at its foot
    rounded = livery["boxes"]["card"]["radius"] > 0
    assert geom(by["Card 1 fill"]) == ("roundRect" if rounded else "rect"), "the card's corners"
    assert geom(by["Card 1 head"]) == ("round2SameRect" if rounded else "rect"), "the head band's corners"
    assert "Card 3 head" not in by and str(by["Card 3 fill"].fill.fore_color.rgb) == role("accent"), "a solid card is its role"
    ol = by["Card 1 outline"]
    assert str(ol.line.color.rgb) == livery["boxes"]["card"]["outline-color"][1:], "the pick's outline color"
    assert ol.line.width == round(livery["boxes"]["card"]["outline"] * prs.slide_height), "the outline's width"
    assert "Card 2 outline" not in by, "an outline on a card not picked"
    title = by["Card 1 title"].text_frame.paragraphs[0].runs[0].font
    assert title.bold and title.size.pt == livery_mod.half_pt(livery["type"]["roles"]["subhead"]["pt"]), "the title's look"
    assert str(title.color.rgb) == role("primary-text"), "the title on the role in primary-text"
    body = by["Card 1 body"].text_frame.paragraphs
    label = next(p for p in body if p.text == "LOCAL INVESTMENT")
    big = next(p for p in body if p.text == "$20M")
    assert label.runs[0]._r.find("{*}rPr").get("cap") == "all", "the label in caps"
    assert big.runs[0].font.size.pt == livery_mod.half_pt(livery["type"]["roles"]["display"]["pt"]), "the big number's size"
    assert "⟦" not in slide._element.xml, "a marker left"


def test_rail(deck):
    """Steps and a timeline: one group each; the items at the formula's frames (equal shares), the
    text boxes one height; a step's badge a circle in its role with its number, bold, in the
    subhead size, knocked out in the background color; the rail boxes.rail.line of the height in line-color, through
    the marks' middle from the first mark to the row's edge; a done milestone's dot filled in the
    role, a pending one the background in a ring of it; the titles bold in the subhead size; no
    marker left."""
    import livery as livery_mod
    livery, prs, _ = deck
    rail, H = livery["boxes"]["rail"], prs.slide_height
    role = lambda r: livery["roles"][r][1:]
    subhead = livery_mod.half_pt(livery["type"]["roles"]["subhead"]["pt"])
    slide = prs.slides[9]
    by = {sh.name: sh for sh in flat(slide.shapes)}
    assert by["Steps"].shape_type == 6 and {"Rail", "Step 1 badge", "Step 3 text"} <= {sh.name for sh in flat(by["Steps"].shapes)}, \
        "the steps one group of their shapes"
    texts = [by[f"Step {k} text"] for k in (1, 2, 3)]
    for t, (x, w) in zip(texts, boxes(livery, prs.slide_width, [1 / 3] * 3)):
        assert abs(t.left - x) <= 2 and abs(t.width - w) <= 2, f"{t.name} {t.left, t.width}, want {x, w}"
    assert len({t.height for t in texts}) == 1 and len({t.top for t in texts}) == 1, "the steps one height"
    badge = by["Step 1 badge"]
    assert badge._element.find(".//{*}prstGeom").get("prst") == "ellipse", "the badge a circle"
    assert badge.width == badge.height == round(rail["badge"] * H), "the badge's diameter"
    assert str(badge.fill.fore_color.rgb) == role("secondary"), "the badge in the role"
    run = badge.text_frame.paragraphs[0].runs[0]
    assert run.text == "1" and run.font.bold and run.font.size.pt == subhead and str(run.font.color.rgb) == role("background"), \
        "the badge's number"
    assert by["Step 3 badge"].text_frame.text == "3"
    line = by["Rail"]
    assert line.height == round(rail["line"] * H) and str(line.fill.fore_color.rgb) == rail["line-color"][1:], "the rail"
    assert abs(line.top + line.height / 2 - (badge.top + badge.height / 2)) <= 1, "the rail through the badges' middle"
    assert line.left == badge.left + badge.width // 2 and line.left + line.width == texts[2].left + texts[2].width, \
        "the rail from the first badge to the row's edge"
    assert badge.top + badge.height == texts[0].top, "the text under the badge"
    title = texts[0].text_frame.paragraphs[0]
    assert title.text == "Propose" and title.runs[0].font.bold and title.runs[0].font.size.pt == subhead, "a step's title"
    assert [p.text for p in texts[1].text_frame.paragraphs] == ["Pilot", "Six weeks.", "A second line."], "a source line a line"
    assert "Step 1 cell" not in by, "no cell without boxes.rail.fill"
    assert "⟦" not in slide._element.xml, "a marker left"

    slide = prs.slides[10]
    by = {sh.name: sh for sh in flat(slide.shapes)}
    assert by["Timeline"].shape_type == 6, "the timeline one group"
    done, pending = by["Milestone 1 dot"], by["Milestone 2 dot"]
    assert str(done.fill.fore_color.rgb) == role("primary-contrast") and done.line.fill.type == 5, "a done milestone filled in the role, no line"   # MSO_FILL_TYPE.BACKGROUND
    assert str(pending.fill.fore_color.rgb) == role("background") and str(pending.line.color.rgb) == role("primary-contrast"), \
        "a pending one hollow, ringed in the role"
    date = by["Milestone 1 title"]
    assert date.text_frame.text == "Q1 2025" and date.text_frame.paragraphs[0].runs[0].font.bold, "the date"
    assert date.top + date.height == by["Rail"].top + by["Rail"].height // 2 - round(rail["dot"] * H) // 2, "the date above the dot"
    assert by["Milestone 2 text"].text_frame.text == "Full allocation", "the label under it"
    assert "⟦" not in slide._element.xml, "a marker left"

    slide = prs.slides[11]   # a process going on from 4: the badges follow the list's numbers
    by = {sh.name: sh for sh in flat(slide.shapes)}
    assert [by[f"Step {k} badge"].text_frame.text for k in (1, 2)] == ["4", "5"], "the badges from the list's start"
    assert str(by["Step 1 badge"].fill.fore_color.rgb) == role("primary-contrast"), "the default role"


def test_card_type_step(pandoc_deck, tmp_path):
    """A deck's type-body: 10 scales a card's big number as the rest (type_step runs after grid):
    ncar's display role (30 pt on its 12 pt body) to half_pt(30 x 10 / 12)."""
    import livery as livery_mod
    md = "---\ntype-body: 10\n---\n\n" + CARDS
    d = ROOT / "liveries" / "ncar"
    pandoc = pandoc_exe()
    out = tmp_path / "t.pptx"
    subprocess.run([pandoc, "--lua-filter", str(ENGINE / "grid-pptx.lua"), "--reference-doc", str(d / "template.pptx"),
                    "-o", str(out)], input=md, text=True, cwd=tmp_path, check=True, capture_output=True)
    step = STEP.replace("prs.save(p)", "from pptx_steps import type_step; type_step.run(prs, p); prs.save(p)")
    env = {**os.environ, "LIVERY_DIR": str(d), "PYTHONPATH": str(ENGINE)}
    subprocess.run([sys.executable, "-c", step, str(out)], env=env, check=True, capture_output=True)
    from pptx import Presentation
    j = json.loads((d / "build" / "livery.json").read_text())
    body = {sh.name: sh for sh in flat(Presentation(str(out)).slides[0].shapes)}["Card 1 body"].text_frame.paragraphs
    big = next(p for p in body if p.text == "$20M")
    want = livery_mod.half_pt(j["type"]["roles"]["display"]["pt"] * 10 / j["type"]["roles"]["body"]["pt"])
    assert big.runs[0].font.size.pt == want, f"the big number at {big.runs[0].font.size.pt}, want {want}"


def test_strip(deck):
    """A strip in a card: its rectangle the strip.tint of its role over the background, at the card's
    full width, between the body's text above and below it, its text in its own box."""
    import livery as livery_mod
    livery, prs, _ = deck
    by = {sh.name: sh for sh in flat(prs.slides[6].shapes)}
    strip, card = by["Card 1 strip 1"], by["Card 1 fill"]
    want = livery_mod.mix(livery["roles"]["accent"], livery["roles"]["background"], livery["boxes"]["strip"]["tint"])
    assert "#" + str(strip.fill.fore_color.rgb) == want, f"the strip's fill {strip.fill.fore_color.rgb}, want {want}"
    assert (strip.left, strip.width) == (card.left, card.width), "the strip at the card's full width"
    assert by["Card 1 strip 1 text"].text_frame.text == "To the consortium"
    above, below = by["Card 1 body"], by["Card 1 body 3"]
    assert above.top + above.height <= strip.top and strip.top + strip.height <= below.top, "the strip between the body's text"
    assert "one" in below.text_frame.text and "$20M" in above.text_frame.text, "the text on each side"


def test_quote(deck):
    """A quote: one group ("Quote") of a bar and a text box, no fill; the bar boxes.quote.bar wide in
    its role= (else the box's bar-color), the text box beside it padded by boxes.quote.pad, its lines
    in the quote size, the attribution "— " and its name at the caption size in muted; no marker.
    In a card it sits inside the body's pads between the text above and below it; in a plain column
    between that column's text boxes, at the column's left."""
    import livery as livery_mod
    livery, prs, _ = deck
    q, sh = livery["boxes"]["quote"], prs.slide_height
    size = lambda p: {r._r.get_or_add_rPr().get("sz") for r in p.runs}
    by = {s.name: s for s in flat(prs.slides[12].shapes)}
    g = by["Quote"]
    assert g.shape_type == 6 and {s.name for s in g.shapes} == {"Quote bar", "Quote text"}, "one group, a bar and its text, no fill"
    bar, text = by["Quote bar"], by["Quote text"]
    assert str(bar.fill.fore_color.rgb) == livery["roles"]["accent"][1:], "role= colors the bar"
    assert bar.width == round(q["bar"] * sh) and text.left == bar.left + bar.width and bar.height == text.height
    assert text.text_frame.margin_left == round(q["pad"][1] * livery_mod.type_pt(livery, "body") * 12700), "padded by boxes.quote.pad"
    paras = text.text_frame.paragraphs
    assert [p.text for p in paras] == ["The purpose of computing", "is insight.", "— R. Hamming, 1962"], "a source line a line"
    pt = lambda role: str(round(livery_mod.half_pt(livery_mod.type_pt(livery, role)) * 100))
    assert size(paras[0]) == {pt("quote")} and size(paras[2]) == {pt("caption")}, "the quote and caption sizes"
    assert {str(r.font.color.rgb) for r in paras[2].runs} == {livery["roles"]["muted"][1:]}, "the attribution muted"
    assert "⟦" not in prs.slides[12]._element.xml, "a marker left"
    pads = livery_mod.type_pt(livery, "body") * 12700
    by = {s.name: s for s in flat(prs.slides[13].shapes)}
    bar, card = by["Card 1 quote 1 bar"], by["Card 1 fill"]
    assert str(bar.fill.fore_color.rgb) == q["bar-color"][1:], "a bare > in the box's bar-color"
    assert abs(bar.left - (card.left + round(livery["boxes"]["card"]["pad"][1] * pads))) <= 1, "inside the card's pads"
    above, below = by["Card 1 body"], by["Card 1 body 3"]
    assert above.top + above.height <= bar.top and bar.top + bar.height <= below.top, "between the body's text"
    assert by["Card 1 quote 1 text"].text_frame.text == "Quoted in a card."
    by = {s.name: s for s in flat(prs.slides[14].shapes)}
    bar, (x, _w) = by["Column 2 quote 1 bar"], boxes(livery, prs.slide_width, [0.5, 0.5])[1]
    assert by["Column 2 quote 1"].shape_type == 6 and abs(bar.left - x) <= 6, "a group at the column's left"
    assert by["Column 2"].text_frame.text == "Before." and by["Column 2 3"].text_frame.text == "After."
    assert by["Column 2"].top + by["Column 2"].height <= bar.top <= bar.top + bar.height <= by["Column 2 3"].top
    assert "⟦" not in prs.slides[13]._element.xml + prs.slides[14]._element.xml, "a marker left"


def test_agenda(deck):
    """An agenda slide in section 1 of 2: one group ("Agenda"), a row per numbered section (the
    unnumbered Backup none), each a number box (its number, the eyebrow size, bold, in
    boxes.agenda.number-color) beside a text box (the title in the subhead size): section 1 current,
    bold in current-color, section 2 next, the text's color; rows top down; no marker."""
    import livery as livery_mod
    livery, prs, _ = deck
    slide = prs.slides[16]
    by = {s.name: s for s in flat(slide.shapes)}
    assert by["Agenda"].shape_type == 6, sorted(by)
    assert {n for n in by if n.startswith("Agenda ")} == {"Agenda 1 number", "Agenda 1 text", "Agenda 2 number", "Agenda 2 text"}
    pt = lambda role: str(round(livery_mod.half_pt(livery_mod.type_pt(livery, role)) * 100))
    for k in (1, 2):
        num, text = by[f"Agenda {k} number"], by[f"Agenda {k} text"]
        r = num.text_frame.paragraphs[0].runs[0]
        assert r.text == str(k) and r._r.get_or_add_rPr().get("sz") == pt("eyebrow") and r.font.bold
        assert str(r.font.color.rgb) == livery["boxes"]["agenda"]["number-color"][1:], "the number's color"
        assert num.top == text.top and num.left < text.left, "the number beside its title"
    one, two = (by[f"Agenda {k} text"].text_frame.paragraphs[0].runs[0] for k in (1, 2))
    assert one.text == "One" and one.font.bold and str(one.font.color.rgb) == livery["boxes"]["agenda"]["current-color"][1:]
    assert two.text == "Two" and not two.font.bold and two._r.get_or_add_rPr().get("sz") == pt("subhead")
    assert two._r.get_or_add_rPr().find("{*}solidFill") is None, "a later section in the text's own color"
    assert by["Agenda 1 text"].top < by["Agenda 2 text"].top
    assert "⟦" not in slide._element.xml, "a marker left"


def test_default_roles(deck):
    """Cards without role= take boxes.card.roles in turn by their place in the row, the fourth the
    first again: no two neighbors alike."""
    livery, prs, _ = deck
    by = {sh.name: sh for sh in flat(prs.slides[7].shapes)}
    roles = livery["boxes"]["card"]["roles"]
    want = [livery["roles"][roles[k % len(roles)]][1:] for k in range(4)]
    assert [str(by[f"Card {n} head"].fill.fore_color.rgb) for n in (1, 2, 3, 4)] == want, "the heads' roles by place"


def test_comparison_warns(tmp_path):
    """A two-column slide pandoc lays out on Comparison (text above a picture in a column) is not
    framed: columns.py warns, naming the slide and the layout, and the notes line is still gone."""
    md = '## Compare\n\n:::: {.columns}\n::: {.column width="60%"}\nText\n\n![](p.jpg)\n:::\n::: {.column}\nRight\n:::\n::::\n'
    shutil.copy(IMAGE, tmp_path / "p.jpg")
    d = ROOT / "liveries" / "ncar"
    pandoc = pandoc_exe()
    out = tmp_path / "t.pptx"
    subprocess.run([pandoc, "--lua-filter", str(ENGINE / "grid-pptx.lua"), "--reference-doc", str(d / "template.pptx"),
                    "-o", str(out)], input=md, text=True, cwd=tmp_path, check=True, capture_output=True)
    env = {**os.environ, "LIVERY_DIR": str(d), "PYTHONPATH": str(ENGINE)}
    r = subprocess.run([sys.executable, "-c", STEP, str(out)], env=env, check=True, capture_output=True, text=True)
    assert "columns: WARNING: slide 1: columns laid out on 'Comparison'" in r.stderr, r.stderr
    from pptx import Presentation
    slide = Presentation(str(out)).slides[0]
    assert slide.slide_layout.name == "Comparison", "the premise: pandoc's layout for text over a picture"
    assert "engine-columns" not in slide.notes_slide.notes_text_frame.text, "the line left in the notes"
