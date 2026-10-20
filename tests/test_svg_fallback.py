"""Check pptx_steps/svg_fallback.py on a deck pandoc writes with an SVG picture.

    python3 -m pytest tests/test_svg_fallback.py

pandoc puts a one-slide deck's SVG in as a plain .svg picture; after the step its blip's r:embed
is a real PNG at the step's dpi, the SVG hangs off the blip's svgBlip extension, [Content_Types]
has png and svg, python-pptx opens the deck and reads the picture, and there is no stroke="0"
(Google Slides refuses a deck that has one, #109). A second run changes nothing, and a deck with
no SVG is left as it is. Needs pandoc (Quarto's, $QUARTO_PANDOC) and rsvg-convert; the deck is
rendered once per module.
"""

import io
import re
import shutil
import zipfile

import pytest
from lxml import etree

from pptx_steps import svg_fallback as sf

SVG = (b'<svg xmlns="http://www.w3.org/2000/svg" width="72pt" height="36pt" viewBox="0 0 72 36">'
       b'<rect x="4" y="4" width="64" height="28" fill="#0057C2"/></svg>')


@pytest.fixture(scope="module")
def svg_deck(pandoc_deck, tmp_path_factory):
    """The one-slide deck of x.svg, as pandoc writes it."""
    assert shutil.which("rsvg-convert"), "rsvg-convert not found (conda-forge librsvg)"
    tmp = tmp_path_factory.mktemp("svg")
    (tmp / "x.svg").write_bytes(SVG)
    return pandoc_deck(tmp, "## s\n\n![](x.svg)\n", "svg.pptx")


@pytest.fixture(scope="module")
def fixed(svg_deck):
    """The same deck after the step."""
    parts = dict(svg_deck)
    sf.run(parts, "svg.pptx")
    return parts


def picture(parts):
    """The picture slide's blip and its slide's relationships, id to target."""
    slide = next(n for n in parts if re.match(r"ppt/slides/slide\d+\.xml$", n) and b"<p:pic>" in parts[n])
    rels = {r.get("Id"): r.get("Target") for r in etree.fromstring(parts[slide.replace("slides/", "slides/_rels/") + ".rels"])}
    return etree.fromstring(parts[slide]).find(".//{%s}blip" % sf.A), rels


def fallback_png(parts) -> tuple[str, bytes]:
    """The blip's r:embed target and its bytes."""
    blip, rels = picture(parts)
    png = rels.get(blip.get("{%s}embed" % sf.R), "")
    return png, parts.get("ppt/media/" + png.rsplit("/", 1)[-1], b"")


def test_pandoc_writes_svg_picture(svg_deck):
    assert any(n.endswith(".svg") for n in svg_deck), "pandoc wrote no .svg picture (the premise)"


def test_blip_embeds_png_at_dpi(fixed):
    png, data = fallback_png(fixed)
    assert png.endswith(".png"), f"the blip's r:embed is {png!r}, not a PNG"
    assert data[:8] == b"\x89PNG\r\n\x1a\n", "the fallback is not a PNG"
    width = int.from_bytes(data[16:20], "big") if len(data) > 20 else 0
    assert abs(width - sf.DPI) <= 1, f"the fallback is {width} px wide, not 1 in at {sf.DPI} dpi"


def test_svgblip_extension(fixed):
    blip, rels = picture(fixed)
    svg_blip = blip.find(".//{%s}svgBlip" % sf.ASVG)
    assert svg_blip is not None and svg_blip.getparent().get("uri") == sf.SVG_EXT, "no svgBlip in the blip's extension list"
    assert rels.get(svg_blip.get("{%s}embed" % sf.R), "").endswith(".svg"), "the svgBlip does not point at the SVG"


def test_content_types(fixed):
    ct = fixed["[Content_Types].xml"].decode()
    for ext, mime in (("png", "image/png"), ("svg", "image/svg+xml")):
        assert f'Extension="{ext}" ContentType="{mime}"' in ct, f"no {ext} Default content type"


def test_no_zero_stroke(fixed):
    assert not any(b'stroke="0"' in v for n, v in fixed.items() if n.endswith(".xml")), 'stroke="0" in the deck'


def test_python_pptx_reads_png(fixed):
    from pptx import Presentation
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        for n, v in fixed.items():
            z.writestr(n, v)
    pics = [sh for s in Presentation(buf).slides for sh in s.shapes if sh.shape_type == 13]
    assert len(pics) == 1 and pics[0].image.content_type == "image/png", "python-pptx does not read the picture as the PNG"


def test_second_run_changes_nothing(fixed):
    again = dict(fixed)
    sf.run(again, "svg.pptx")
    assert again == fixed, "a second run changed the deck"


def test_png_deck_untouched(fixed, pandoc_deck, tmp_path):
    (tmp_path / "y.png").write_bytes(fallback_png(fixed)[1])   # the fallback itself, as a PNG-only deck's picture
    plain = pandoc_deck(tmp_path, "## s\n\n![](y.png)\n", "png.pptx")
    before = dict(plain)
    sf.run(plain, "png.pptx")
    assert plain == before, "a deck with no SVG changed"
