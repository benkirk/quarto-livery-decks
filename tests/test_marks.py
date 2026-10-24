"""Check the content slides' marks by corner (geometry.marks) from the build to every drawer's numbers.

    python3 -m pytest tests/test_marks.py

One mark per corner, any corner, each named (`mark:`); the field roles wear an image per name
(marks:, a corner's light and dark the defaults). The title gives way to each top corner's mark on
its side (title_room(): livery.css's --engine-title-left|right and the template's title placeholder
agree), the body starts under the taller top mark, a flavor's null drops its parent's corner, and
the build refuses what the schema cannot see (a name worn twice, a top corner's bottom, a missing
aspect once merged, an image for a name no corner wears).
"""

import json
import re
import sys

import pytest
from pptx import Presentation

from conftest import ENGINE, ROOT, flavor, run, title_ph

sys.path.insert(0, str(ENGINE))
import livery_build as lb  # noqa: E402

# ncar with a second mark top left: the white lockup, half the logo's height, inset at the edge
TWO = """geometry:
  marks:
    top-left: {mark: emblem, light: ncar-white, dark: ncar-color, height: 0.05, top: 0.06, aspect: 2, inset: 0.01}
"""


def resolved(tmp, name, knobs, parent="ncar"):
    return lb.resolve(flavor(tmp, name, knobs, parent))[0]


@pytest.fixture(scope="module")
def two(tmp_path_factory):
    """The two-corner flavor built by the command: (dir, livery.json, sty, css, the template's master)."""
    tmp = tmp_path_factory.mktemp("two")
    d = flavor(tmp, "two", TWO)
    r = run("build", d)
    assert r.returncode == 0, r.stderr
    r = run("template", d, "-o", tmp / "t.pptx")
    assert r.returncode == 0, r.stderr
    b = d / "build"
    return (d, json.loads((b / "livery.json").read_text()), (b / "livery.sty").read_text(), (b / "livery.css").read_text(),
            Presentation(str(tmp / "t.pptx")))


def prop(css, name):
    return next(l for l in css.splitlines() if f"--{name}:" in l).split(":", 1)[1].strip().rstrip(";")


def test_two_top_corners_room(two):
    """Each top corner's room on its side: the css's calc and the template's title placeholder."""
    _, j, _, css, prs = two
    m, g = j["geometry"]["margin"], j["geometry"]["marks"]
    hw = j["type"]["canvas"]["h"] / j["type"]["canvas"]["w"]
    room = lambda e: e.get("inset", m) + e["height"] * e["aspect"] * hw + lb.TITLE_GAP - m
    left, right = room(g["top-left"]), room(g["top-right"])
    assert lb.title_room(j) == pytest.approx((left, right)), "title_room(): each side its own corner's"
    assert prop(css, "engine-title-left") == "calc(16px + 40px + 90.0px)", "the top-left mark's inset, the gap, its width"
    assert prop(css, "engine-title-right") == "calc(var(--engine-margin) + 40px + 228.2px)", "the logo's"
    ph = title_ph(prs.slide_master)
    W = prs.slide_width
    assert abs(ph.left - (m + left) * W) <= 1 and abs(ph.left + ph.width - (1 - m - right) * W) <= 1, "the placeholder between them"
    assert j["furniture"]["title"][0] == pytest.approx(m + left)


def test_two_top_corners_drawn(two):
    """The sty's corners and keys, the css's rules per field role, the master's pictures by corner."""
    _, j, sty, css, prs = two
    assert r"\def\engine@corners{top-left,top-right}" in sty, "the worn corners, in corner order"
    for k, v in (("mark", "emblem"), ("height", "0.05"), ("top", "0.06"), ("aspect", "2"), ("inset", "0.01")):
        assert rf"\expandafter\def\csname engine@corner@top-left@{k}\endcsname{{{v}}}" in sty, k
    assert r"\expandafter\def\csname engine@mark@background@emblem\endcsname" in sty, "the per-role image of the new name"
    assert ".reveal .slides section.slide .engine-mark-top-left { top: 54px; left: 16px; width: 90.0px; height: 45px; " \
           "background-position: left top; background-image: var(--engine-mark-background-emblem); }" in css
    assert ".reveal .slides section.engine-field .engine-mark-top-left { background-image: var(--engine-mark-primary-emblem); }" in css
    assert ".reveal .slides section.engine-dark .engine-mark-top-left { background-image: var(--engine-mark-dark-emblem); }" in css
    assert j["marks"]["background"]["emblem"].endswith("Color-White_RGB.png"), "light: the background role's"
    assert j["marks"]["primary"]["emblem"].endswith("FullColor_RGB.png") and j["marks"]["dark"]["emblem"] == j["marks"]["primary"]["emblem"]
    assert j["marks"]["primary"]["emblem-print"].endswith(".pdf"), "a name's print twin for LaTeX"
    pics = [s.name for s in prs.slide_master.shapes if s.shape_type == 13]
    assert pics == ["Mark emblem (light)", "Logo (light)"], "one picture per corner, in corner order"
    assert '--engine-mark-background-emblem: url("NSF-NCAR_Logo_Color-White_RGB.png");' in css


def test_symmetric_takes_the_wider(tmp_path):
    L = resolved(tmp_path, "sym", TWO + "frametitle: {align: center, symmetric: true}\n")
    left, right = lb.title_room(L)
    assert left == right == pytest.approx(lb.corner_room(L, "top-right")), "the logo's room is the wider, on both sides"
    css = lb.emit_css(L)
    assert prop(css, "engine-title-left") == prop(css, "engine-title-right")
    t = L["furniture"]["title"]
    assert t[0] + t[2] / 2 == pytest.approx(0.5), "centered on the slide"


def test_body_under_the_taller(tmp_path):
    L = resolved(tmp_path, "tall", "geometry:\n  marks:\n    top-left: {mark: emblem, light: ncar-white, dark: ncar-white, height: 0.2, top: 0.04, aspect: 1}\n")
    assert L["frametitle"]["body"] == round(0.24 + lb.BODY_GAP, 4)


def test_bottom_right_default_box(tmp_path):
    """A bottom corner's defaults (height 0.0644, bottom 0.03, the margin) in furniture, and a sign-off
    there is the fields' sign-off in that corner."""
    L = resolved(tmp_path, "br", "geometry:\n  marks:\n    bottom-right: {mark: signoff, light: ncar-color, dark: ncar-white, aspect: 4}\n")
    m, h = L["geometry"]["margin"], 0.0644
    w = h * 4 * L["type"]["canvas"]["h"] / L["type"]["canvas"]["w"]
    assert L["furniture"]["marks"]["bottom-right"] == {"mark": "signoff", "box": pytest.approx([1 - m - w, 1 - 0.03 - h, w, h])}
    so = L["compositions"]["section"]["signoff"]["box"]
    assert so[0] == pytest.approx(1 - m - w, abs=1e-4) and so[1] == pytest.approx(1 - 0.03 - h, abs=1e-4), "the divider's in its corner"
    assert L["marks"]["background"]["signoff"].endswith("FullColor_RGB.png")


def test_null_drops_a_parents_corner(tmp_path):
    L = resolved(tmp_path, "bare", "geometry:\n  marks:\n    bottom-left: null\n", parent="nasa")
    assert L["geometry"]["marks"]["bottom-left"] is None and L["furniture"]["marks"]["bottom-left"] is None
    assert "signoff" not in L["marks"]["background"], "no corner wears it: no image"
    assert L["compositions"]["section"]["signoff"] is None
    assert L["geometry"]["marks"]["top-right"]["mark"] == "logo", "the other corner inherited"


def test_another_mark_replaces_the_parents_entry(tmp_path):
    """A flavor's corner entry naming another mark replaces its parent's (nasa-ccp's insignia over
    nasa's logo corner), so none of the parent's keys (its *-height) leak into it; one naming the same
    mark, or none, merges over it (ncarucar's aspect)."""
    L = resolved(tmp_path, "ins", "geometry:\n  marks:\n    top-right: {mark: insignia, light: meatball, dark: meatball, aspect: 1.2}\n"
                 "    top-left: {mark: logo, aspect: 1, title-height: 0.2}\n", parent="nasa")
    tr = L["geometry"]["marks"]["top-right"]
    assert tr == {"mark": "insignia", "light": "meatball", "dark": "meatball", "aspect": 1.2, "height": lb.MARK["height"],
                  "top": lb.MARK["top"]}, "the parent's top-right entry not merged in"
    assert L["compositions"]["title"]["mark"]["box"][3] == 0.2, "the cover's mark: the logo entry's title-height"
    j = json.loads((ROOT / "liveries" / "nasa-ccp" / "build" / "livery.json").read_text())
    assert set(j["marks"]["background"]) == {"logo", "insignia"} and j["marks"]["background"]["insignia"].endswith("meatball.png")
    assert j["compositions"]["section"]["signoff"]["box"][0] == pytest.approx(j["furniture"]["marks"]["top-right"]["box"][0], abs=1e-3), \
        "the divider's insignia where the content slide's is"


@pytest.mark.parametrize("knobs, why", [
    ("geometry:\n  marks:\n    top-left: {mark: logo, aspect: 1}\n", "top-right.mark: 'logo' is already geometry.marks.top-left's"),
    ("geometry:\n  marks:\n    bottom-left: {mark: emblem, light: ncar-color, dark: ncar-color, aspect: 1, top: 0.1}\n",
     "a bottom corner's distance is bottom, not top"),
    ("geometry:\n  marks:\n    top-left: {mark: emblem, light: ncar-color, dark: ncar-color}\n", "geometry.marks.top-left.aspect is required"),
    ("marks:\n  background: {nope: ncar-color}\n", "not ['nope']"),
    ("geometry:\n  marks:\n    top-left: {mark: emblem, aspect: 1}\n", "marks.background.emblem None is not a logo.images name"),
    ("geometry:\n  marks:\n    bottom-left: {mark: signoff, light: ncar-color, dark: ncar-color, aspect: 1, title-height: 0.1}\n",
     "belong to the entry whose mark is logo"),
])
def test_refusals(tmp_path, knobs, why):
    with pytest.raises(SystemExit) as e:
        resolved(tmp_path, "bad", knobs)
    assert why in str(e.value), str(e.value)


def test_bundled_corners():
    """The ten liveries' corners as they were drawn before the move: ncar's family top right, cu's
    sign-off alone, nsi's mark top left over cu's foot, nasa's insignia and name, nasa-ccp's emblem."""
    want = {"ncar": {"top-right": "logo"}, "ucar": {"top-right": "logo"}, "ucp": {"top-right": "logo"},
            "ncarucar": {"top-right": "logo"}, "cu": {"bottom-left": "signoff"},
            "nsi": {"top-left": "logo", "bottom-left": "signoff"}, "nasa": {"top-right": "logo", "bottom-left": "signoff"},
            "nasa-dark": {"top-right": "logo", "bottom-left": "signoff"}, "nasa-jsc-eg": {"top-right": "logo", "bottom-left": "signoff"},
            "nasa-ccp": {"top-left": "logo", "top-right": "insignia"}}
    for name, corners in want.items():
        g = json.loads((ROOT / "liveries" / name / "build" / "livery.json").read_text())["geometry"]["marks"]
        assert {c: e["mark"] for c, e in g.items() if e} == corners, name
        assert not re.search(r"engine@geom@logo|engine@signoff(height|bottom|corner)",
                             (ROOT / "liveries" / name / "build" / "livery.sty").read_text()), f"{name}: the old names"
