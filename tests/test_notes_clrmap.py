"""Check pptx_steps/notes_clrmap.py on a deck pandoc writes with speaker notes.

    python3 -m pytest tests/test_notes_clrmap.py

pandoc writes a notes slide with no <p:clrMapOvr> (the premise; Keynote refuses such a deck).
After the step each notes slide has exactly one, right after its p:cSld, as PowerPoint writes it,
and python-pptx still reads the notes. A second run changes nothing, and a deck with no notes is
left as it is. Needs pandoc (Quarto's, $QUARTO_PANDOC); the deck is rendered once per module.
"""

import io
import zipfile

import pytest
from lxml import etree

from pptx_steps import notes_clrmap as nc

P = "http://schemas.openxmlformats.org/presentationml/2006/main"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
MD = "## a\n\nOne.\n\n::: notes\nA note.\n:::\n\n## b\n\nTwo.\n\n::: notes\nAnother.\n:::\n"


@pytest.fixture(scope="module")
def notes_deck(pandoc_deck, tmp_path_factory):
    """The two-slide deck with a note on each, as pandoc writes it."""
    return pandoc_deck(tmp_path_factory.mktemp("notes"), MD, "notes.pptx")


@pytest.fixture(scope="module")
def fixed(notes_deck):
    """The same deck after the step."""
    parts = dict(notes_deck)
    nc.run(parts, "notes.pptx")
    return parts


def notes_of(parts):
    return [n for n in parts if nc.NOTES.match(n)]


def test_pandoc_writes_no_clrmap(notes_deck):
    notes = notes_of(notes_deck)
    assert len(notes) == 2, f"pandoc wrote {len(notes)} notes slides, not 2"
    assert not any(b"clrMapOvr" in notes_deck[n] for n in notes), \
        "pandoc already writes clrMapOvr on notes slides (the premise; the step may go)"


def test_one_clrmap_after_csld(fixed):
    for n in notes_of(fixed):
        x = etree.fromstring(fixed[n])
        kids = [etree.QName(k).localname for k in x]
        assert kids[:2] == ["cSld", "clrMapOvr"] and kids.count("clrMapOvr") == 1, \
            f"{n}: children {kids}, not cSld then one clrMapOvr"
        assert x.find("{%s}clrMapOvr/{%s}masterClrMapping" % (P, A)) is not None, f"{n}: no masterClrMapping"


def test_python_pptx_reads_notes(fixed):
    from pptx import Presentation
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        for n, v in fixed.items():
            z.writestr(n, v)
    texts = [s.notes_slide.notes_text_frame.text for s in Presentation(buf).slides]
    assert texts == ["A note.", "Another."], f"python-pptx reads the notes as {texts}"


def test_second_run_changes_nothing(fixed):
    again = dict(fixed)
    nc.run(again, "notes.pptx")
    assert again == fixed, "a second run changed the deck"


def test_deck_without_notes_untouched(pandoc_deck, tmp_path):
    plain = pandoc_deck(tmp_path, "## a\n\nNo notes.\n", "plain.pptx")
    before = dict(plain)
    nc.run(plain, "plain.pptx")
    assert plain == before, "a deck with no notes changed"
