"""Check the boxes (livery.yml boxes:, livery_build.py BOXES): one set of numbers every format draws.

    python3 -m pytest tests/test_boxes.py

Every bundled livery's build/livery.json carries every box and key of BOXES, its fractions in
[0, 1], its pads two non-negative numbers, its colors hex; the css and the sty carry them too. A
livery's block merges over the defaults box by box; a color resolves from a role, a palette name or
a hex, and a name that is none of them, or an unknown key, fails with its path. The pptx caution
box takes its bar and pads from them (pptx_steps/slide_layout.py, in a subprocess under the
livery's LIVERY_DIR). Needs pandoc ($QUARTO_PANDOC) for that one.
"""

import json
import os
import re
import subprocess
import sys

import pytest

import livery_build as lb

from conftest import ENGINE, ROOT, run

HEX = re.compile(r"^#[0-9A-F]{6}$")
FRACTIONS = ("gap", "rule", "bar", "outline", "radius", "tint", "line", "badge", "dot", "ring")
EMS = {("references", "hang"), ("references", "gap"), ("agenda", "gap"), ("agenda", "number-width"),
       ("quote", "attribution-gap"), ("note", "gap")}   # em of a type role, not a fraction


def test_every_key(livery):
    b = json.loads((livery / "build" / "livery.json").read_text())["boxes"]
    assert set(b) == set(lb.BOXES), f"{livery.name}: boxes {sorted(b)} are not BOXES'"
    for box, keys in lb.BOXES.items():
        assert set(b[box]) == set(keys), f"{livery.name}: boxes.{box} keys {sorted(b[box])}"
        for k, v in b[box].items():
            where = f"{livery.name}: boxes.{box}.{k} = {v!r}"
            if k == "roles":
                assert v and all(r in json.loads((livery / "build" / "livery.json").read_text())["roles"] for r in v), f"{where}: not roles"
            elif k in lb.BOX_COLORS.get(box, ()):
                assert (v is None and lb.BOXES[box][k] is None) or HEX.match(v), f"{where} is not #RRGGBB"   # a nullable color
            elif (box, k) == ("quote", "family"):   # a bundled family, or null: the body's
                assert v is None or isinstance(v, str), f"{where} is not a family or null"
            elif (box, k) == ("quote", "style"):
                assert v in ("normal", "italic"), f"{where} is not normal or italic"
            elif k == "role":   # a role or blend name, kept as a name
                lv = json.loads((livery / "build" / "livery.json").read_text())
                assert v in lv["roles"] or v in lv["blends"], f"{where} is not a role or blend"
            elif (box, k) in EMS:
                assert isinstance(v, (int, float)) and v >= 0, f"{where} is not an em length"
            elif k in FRACTIONS:
                assert isinstance(v, (int, float)) and 0 <= v <= 1, f"{where} is not a fraction"
            else:
                assert k.endswith("pad") and len(v) == 2 and all(x >= 0 for x in v), f"{where} is not two pads"
    css = (livery / "build" / "livery.css").read_text()
    sty = (livery / "build" / "livery.sty").read_text()
    assert f"--engine-caution-pad: {b['caution']['pad'][0]:g}em {b['caution']['pad'][1]:g}em;" in css
    assert f"--engine-glyph-done: {b['status']['done']};" in css   # the glyphs'; --engine-status-* are the charts'
    assert f"\\definecolor{{EngineGlyphDone}}{{HTML}}{{{b['status']['done'][1:]}}}" in sty
    assert f"\\def\\engine@boxes@caution@padx{{{b['caution']['pad'][1]}}}" in sty
    assert f"\\definecolor{{EngineBoxesCardOutlineColor}}{{HTML}}{{{b['card']['outline-color'][1:]}}}" in sty
    assert f"--engine-rail-line-color: {b['rail']['line-color']};" in css
    assert f"--engine-rail-fill: {b['rail']['fill'] or 'transparent'};" in css
    assert ("EngineBoxesRailFill" in sty) == (b["rail"]["fill"] is not None), "a null fill defines no color"
    g = b["gallery"]   # the gallery: its gap, pads and fill in the css, its gap in the sty; the geometry it reads
    assert f"--engine-gallery-pad: {g['pad'][0]:g}em {g['pad'][1]:g}em;" in css and f"--engine-gallery-pad-y: {g['pad'][0]:g}em;" in css
    assert re.search(r"--engine-gallery-gap: [\d.]+px;", css) and f"--engine-gallery-fill: {g['fill'] or 'transparent'};" in css
    assert f"\\def\\engine@boxes@gallery@gap{{{g['gap']}}}" in sty
    assert ("EngineBoxesGalleryFill" in sty) == (g["fill"] is not None), "a null fill defines no color"
    geom = json.loads((livery / "build" / "livery.json").read_text())["geometry"]
    assert f"--engine-body-bottom: {round(geom['body-bottom'] * 900, 1):g}px;" in css
    assert f"\\def\\engine@geom@bodyheight{{{geom['body'][3]}}}" in sty
    assert "\\def\\engine@type@display{" in sty and "--engine-type-display:" in css
    q = b["quote"]   # the quote block: its bar, its color and pads in every format's file, its type role
    assert f"--engine-quote-bar-color: {q['bar-color']};" in css and f"--engine-quote-pad: {q['pad'][0]:g}em {q['pad'][1]:g}em;" in css
    assert f"\\definecolor{{EngineBoxesQuoteBarColor}}{{HTML}}{{{q['bar-color'][1:]}}}" in sty
    assert f"\\def\\engine@boxes@quote@bar{{{q['bar']}}}" in sty and f"\\def\\engine@boxes@quote@pady{{{q['pad'][0]}}}" in sty
    assert "\\def\\engine@type@quote{" in sty and "--engine-type-quote:" in css
    # the theme resolves a role= and the card roles by place through the role table, in the sty too
    j = json.loads((livery / "build" / "livery.json").read_text())
    assert f"\\def\\engine@boxes@card@roles{{{','.join(b['card']['roles'])}}}" in sty
    for role, tex in j["latex"]["roles"].items():
        assert f"\\expandafter\\def\\csname engine@role@{role}\\endcsname{{{tex}}}" in sty, \
            f"{livery.name}: {role} is not in the sty's role table"


def test_merge_box_by_box():
    hexof = lambda c, k: c if c.startswith("#") else {"surface": "#F1F0EE"}.get(c, "#000000")
    b = lb.boxes({"card": {"radius": 0.012}, "lead": {"bar-color": "#123456"}}, hexof)
    assert b["card"]["radius"] == 0.012 and b["card"]["pad"] == lb.BOXES["card"]["pad"], "a key merges into its box"
    assert b["lead"]["bar-color"] == "#123456" and b["lead"]["fill"] == "#F1F0EE", "the box's other keys keep their defaults"
    assert b["caution"] == lb.BOXES["caution"], "an untouched box is the default"


def test_colors_resolve(tmp_path, tmp_livery):
    rc, err, j = tmp_livery(tmp_path, "  lead: {fill: secondary, bar-color: orange}\n  card: {outline-color: \"#6840f0\"}\n")
    assert rc == 0, err
    assert j["boxes"]["lead"]["fill"] == j["roles"]["secondary"], "a role"
    assert j["boxes"]["lead"]["bar-color"] == j["palette"]["orange"], "a palette name"
    assert j["boxes"]["card"]["outline-color"] == "#6840F0", "a hex, upper-cased"
    assert j["boxes"]["status"]["done"] == j["roles"]["highlight"], "the default, a role"
    assert j["boxes"]["rail"]["line-color"] == j["blends"]["rule"], "the default, a blend"
    assert j["boxes"]["rail"]["fill"] is None, "the default, null: no cell"
    assert j["boxes"]["gallery"]["fill"] is None, "the default, null: no cell"
    rc, err, j = tmp_livery(tmp_path / "fill", "  rail: {fill: surface, line-color: muted}\n")
    assert rc == 0, err
    assert j["boxes"]["rail"]["fill"] == j["roles"]["surface"] and j["boxes"]["rail"]["line-color"] == j["roles"]["muted"], "a cell"
    rc, err, j = tmp_livery(tmp_path / "gallery", "  gallery: {fill: surface, gap: 0.02}\n")
    assert rc == 0, err
    assert j["boxes"]["gallery"]["fill"] == j["roles"]["surface"] and j["boxes"]["gallery"]["gap"] == 0.02, "a tile look"
    assert j["boxes"]["gallery"]["pad"] == lb.BOXES["gallery"]["pad"], "its other keys the defaults"


@pytest.mark.parametrize("boxes, expect", [
    ("  card: {radiuss: 0.01}\n", "boxes.card.radiuss: unknown key"),
    ("  grid: {gap: 0.01}\n", "boxes.grid: unknown key"),
    ("  lead: {fill: nope}\n", "boxes.lead.fill: 'nope' is neither a palette name nor a hex value"),
    ("  columns: {gap: 1.5}\n", "boxes.columns.gap: 1.5 is greater than the maximum of 1"),
    ("  caution: {pad: [0.5]}\n", "boxes.caution.pad: "),
    ("  card: {roles: [primary, nope]}\n", "boxes.card.roles[1]: 'nope' is neither"),
])
def test_refused_with_path(tmp_path, tmp_livery, boxes, expect):
    rc, err, _ = tmp_livery(tmp_path, boxes)
    assert rc != 0 and expect in err, f"rc {rc}: {err.strip()!r}"


STEP = """
import sys
from pathlib import Path
from pptx import Presentation
from pptx_steps import slide_layout
prs = Presentation(sys.argv[1]); slide_layout.run(prs, Path(sys.argv[1])); prs.save(sys.argv[1])
"""


@pytest.mark.parametrize("name", ["ncar", "nasa-jsc-eg"])
def test_pptx_caution(pandoc_deck, tmp_path, name):
    """The caution box's bar is boxes.caution.bar of the canvas height, its pads em of the body."""
    md = "## Careful\n\nA warning.\n\n::: notes\nengine-layout: caution\n:::\n"
    pandoc_deck(tmp_path, md, "t.pptx")
    env = {**os.environ, "LIVERY_DIR": str(ROOT / "liveries" / name), "PYTHONPATH": str(ENGINE)}
    subprocess.run([sys.executable, "-c", STEP, str(tmp_path / "t.pptx")], env=env, check=True, capture_output=True)
    from pptx import Presentation
    from pptx.util import Pt
    j = json.loads((ROOT / "liveries" / name / "build" / "livery.json").read_text())
    body, c = j["type"]["roles"]["body"]["pt"], j["boxes"]["caution"]
    slide = Presentation(str(tmp_path / "t.pptx")).slides[0]
    bars = [sh for sh in slide.shapes if not sh.is_placeholder and sh.shape_type == 1]   # autoshapes
    assert [b.width for b in bars] == [round(c["bar"] * Pt(j["type"]["canvas"]["h"]))], "one bar, its width the fraction"
    tf = next(sh for sh in slide.placeholders if sh.has_text_frame and "warning" in sh.text_frame.text).text_frame
    assert (tf.margin_top, tf.margin_left) == (round(c["pad"][0] * body * 12700), round(c["pad"][1] * body * 12700)), \
        "pads in em of the body"


def test_bullets(livery, tmp_path):
    """bullets: null where a livery sets none (each format keeps its shapes); where one does
    (nasa-jsc-eg), its glyphs and color in the json, the css, the sty and the template's master."""
    import zipfile
    j = json.loads((livery / "build" / "livery.json").read_text())
    knobs = (livery / "livery.yml").read_text()
    css, sty = (livery / "build" / "livery.css").read_text(), (livery / "build" / "livery.sty").read_text()
    if j["bullets"] is None:
        assert "\nbullets:" not in knobs and "EngineBullet" not in sty and "li::before { content:" not in css
        return
    glyphs = j["bullets"]["glyphs"]
    assert HEX.match(j["bullets"]["color"]), "the bullets' color as hex"
    assert f'content: "{glyphs[0]}"' in css and f"\\def\\engine@bullet@i{{{glyphs[0]}}}" in sty
    r = run("template", livery, "-o", tmp_path / "t.pptx")
    assert r.returncode == 0, r.stderr
    with zipfile.ZipFile(tmp_path / "t.pptx") as z:
        master = z.read("ppt/slideMasters/slideMaster1.xml").decode()
    chars = re.findall(r'<a:buChar char="([^"]*)"/>', master)
    want = [glyphs[min(i, len(glyphs) - 1)] for i in range(5)]
    assert chars[:5] == want, f"the master's bullets {chars[:5]}, want {want}"
    assert f'<a:buClr><a:srgbClr val="{j["bullets"]["color"][1:]}"/></a:buClr>' in master, "their color"


def test_geometry_derived(livery):
    """livery.json carries the content title band (frametitle.top/height/rule-y, every key, the defaults
    FRAMETITLE's), the body's top derived from it (BODY_GAP under the hairline or under a taller mark),
    the content body's box and the Two Content half (geometry.body, geometry.half), derived once by
    livery build; the template, the plots' figsize and columns.py read them, and the template's master
    draws the title rule at rule-y."""
    import zipfile
    j = json.loads((livery / "build" / "livery.json").read_text())
    g, m, ft = j["geometry"], j["geometry"]["margin"], j["frametitle"]
    assert set(ft) == set(lb.FRAMETITLE) | {"body"}, f"{livery.name}: frametitle keys"
    assert HEX.match(ft["rule-color"]) and 0 <= ft["rule-width"] <= 1, f"{livery.name}: frametitle.rule-color/rule-width"
    lg = g.get("logo") or {}
    below = lg.get("top", 0.0533) + lg.get("height", 0.0644) if lg.get("corner", "top-right") != "none" else 0
    assert ft["body"] == round(max(ft["rule-y"], below) + lb.BODY_GAP, 4), f"{livery.name}: frametitle.body"
    top = ft["body"]
    assert g["body"] == [m, top, round(1 - 2 * m, 4), round(g["body-bottom"] - top, 4)], f"{livery.name}: geometry.body"
    assert g["half"] == round((1 - 2 * m - j["boxes"]["columns"]["gap"]) / 2, 4), f"{livery.name}: geometry.half"
    with zipfile.ZipFile(livery / "template.pptx") as z:
        master = z.read("ppt/slideMasters/slideMaster1.xml").decode()
    rule = re.search(r'name="Title rule".*?<a:off x="\d+" y="(\d+)".*?<a:ln w="(\d+)"><a:solidFill><a:srgbClr val="([0-9A-F]{6})"',
                     master, re.S)
    h_emu = j["type"]["canvas"]["h"] * 12700   # the livery's pptx canvas (type.canvas, points)
    if ft["rule"]:
        assert rule, f"{livery.name}: the master draws no title rule"
        assert int(rule.group(1)) == round(ft["rule-y"] * h_emu), f"{livery.name}: the master's title rule is not at rule-y"
        assert int(rule.group(2)) == round(ft["rule-width"] * h_emu), f"{livery.name}: the master's title rule is not rule-width thick"
        assert rule.group(3) == ft["rule-color"][1:], f"{livery.name}: the master's title rule is not rule-color"
    else:
        assert not rule, f"{livery.name}: frametitle.rule is off, yet the master draws a title rule"


def _contrast(a, b):
    """WCAG contrast of two #RRGGBB colors."""
    def lum(h):
        c = [int(h[i:i + 2], 16) / 255 for i in (1, 3, 5)]
        c = [x / 12.92 if x <= 0.03928 else ((x + 0.055) / 1.055) ** 2.4 for x in c]
        return 0.2126 * c[0] + 0.7152 * c[1] + 0.0722 * c[2]
    hi, lo = sorted((lum(a), lum(b)), reverse=True)
    return (hi + 0.05) / (lo + 0.05)


def test_quote_bar_reads(livery):
    """The quote's default bar, a thin mark with no fill behind it, reads on the livery's background:
    3:1 at least (WCAG's non-text contrast). secondary did not (ncar 2.1, ucar 1.6)."""
    j = json.loads((livery / "build" / "livery.json").read_text())
    c = _contrast(j["boxes"]["quote"]["bar-color"], j["roles"]["background"])
    assert c >= 3, f"{livery.name}: the quote bar {j['boxes']['quote']['bar-color']} on {j['roles']['background']} is {c:.1f}:1"
