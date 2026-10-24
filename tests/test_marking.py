"""Check the marking strips (livery.yml marking:): their boxes from the build, the deck's word on
them (engine-common.lua), and pptx_steps/marking.py.

    python3 -m pytest tests/test_marking.py

The build places each named strip by its edge, offset, width and alignment, in the livery's order.
slide-layout.lua (under a `quarto` shim) writes every strip's text to docProps/custom.xml as
`engine-marking-<name>` from engine-common.lua's one reader of the deck's `marking:` (the deck's text,
else the strip's default, "none" for no strip), and the pptx step draws each strip on every slide.
The step runs in a subprocess, under the livery its LIVERY_DIR names. Needs pandoc ($QUARTO_PANDOC).
"""

import json
import os
import re
import subprocess
import sys

import pytest

import livery as livery_mod
import livery_build as lb

from conftest import ENGINE, ROOT, lua_common, pandoc_exe, slide_layout_shim

DEMO = "This presentation contains no sensitive information"
STEP = """
import sys
from pathlib import Path
from pptx import Presentation
from pptx_steps import marking, type_step
p = Path(sys.argv[1]); prs = Presentation(p); marking.run(prs, p); type_step.run(prs, p); prs.save(p)
"""   # the driver's order: the strips before the size step, which scales them as HTML and PDF do
TWO = """  banner: {text: TOP, fill: primary, bold: true, size: 14}
  legend: {text: FOOT, edge: bottom, width: 0.7, size: 8}
"""   # a two-strip block for ncar (tmp_livery)


def test_box():
    out = {"geometry": {"margin": 0.05}, "type": {"roles": {"marking": {"pt": 7, "leading": 1.2}}}}
    hexof = lambda c, k: {"primary": "#001122", "primary-text": "#FFFFFF", "muted": "#808080"}.get(c, c)
    one = lambda m: lb.markings("t", {"s": m}, out, hexof)[0]
    assert one({})["box"] == [0.1, 0.004, 0.8, 0.03], "the default: top, centered, 0.8 wide"
    m = one({"edge": "bottom", "align": "left", "width": 0.5, "height": 0.04, "offset": 0.01})
    assert m["box"] == [0.05, 0.95, 0.5, 0.04], "bottom-left at the margin"
    assert one({"align": "right", "width": 0.5})["box"][0] == 0.45
    assert one({})["color"] == "#808080" and one({})["fill"] is None, "muted text, no fill, as hex"
    assert one({"fill": "primary"})["color"] == "#FFFFFF", "text over a fill: primary-text, as hex"
    assert one({"fill": "#6840F0", "color": "#FFFFFF"})["color"] == "#FFFFFF", "a hex color"
    assert (one({})["size"], one({})["leading"]) == (7, 1.2), "the marking role's size by default"
    assert one({"size": 11})["size"] == 11


def test_names_and_order():
    out = {"geometry": {"margin": 0.05}, "type": {"roles": {"marking": {"pt": 7, "leading": 1.2}}}}
    hexof = lambda c, k: c
    names = lambda spec: [s["name"] for s in lb.markings("t", spec, out, hexof)]
    assert names({}) == ["banner"], "no marking: one text-less strip, so a deck's marking: works anywhere"
    assert lb.markings("t", {}, out, hexof)[0]["text"] is None
    assert names({"top": {}, "foot": {"edge": "bottom"}}) == ["top", "foot"], "the livery's order"
    assert names({"top": {}, "foot": None}) == ["top"], "null drops a parent's strip"
    with pytest.raises(SystemExit, match="marking.Bad"):
        lb.markings("t", {"Bad": {}}, out, hexof)


def strips(pandoc_deck, tmp, name, front, livery_dir=None):
    """The markings (name, text, size in pt) on each slide of a two-slide deck after the steps, under
    the livery name (or the directory livery_dir)."""
    d = livery_dir or ROOT / "liveries" / name
    md = f"---\nlivery: {d}\n{front}---\n\n## a\n\nOne.\n\n## b\n\nTwo.\n"   # the filter reads the livery's strips
    pandoc_deck(tmp, md, "t.pptx", "--lua-filter", str(slide_layout_shim(tmp)))
    env = {**os.environ, "LIVERY_DIR": str(d), "PYTHONPATH": str(ENGINE)}
    subprocess.run([sys.executable, "-c", STEP, str(tmp / "t.pptx")], env=env, check=True, capture_output=True)
    from pptx import Presentation
    prs = Presentation(str(tmp / "t.pptx"))
    return [[(sh.name, sh.text_frame.text, sh.text_frame.paragraphs[0].runs[0].font.size.pt)
             for sh in s.shapes if sh.name.startswith("Marking")] for s in prs.slides]


@pytest.mark.parametrize("name, front, want", [
    ("nasa-jsc-eg", "", DEMO),                       # the livery's default
    ("nasa-jsc-eg", "marking: CUI\n", "CUI"),        # the deck's own
    ("nasa-jsc-eg", "marking: none\n", None),        # turned off
    ("nasa-jsc-eg", "marking: {banner: none}\n", None),   # turned off by name
    ("ncar", "", None),                              # no default, no strip
    ("ncar", "marking: CUI//SP-EXPT\n", "CUI//SP-EXPT"),
])
def test_pptx_step(pandoc_deck, tmp_path, name, front, want):   # not `livery`: conftest's per-livery argument
    got = strips(pandoc_deck, tmp_path, name, front)
    if want is None:
        assert got == [[], []], f"a strip where none belongs: {got}"
    else:
        assert all(len(s) == 1 and s[0][0] == "Marking banner" and s[0][1].startswith(want) for s in got), \
            f"not {want!r} on every slide: {got}"


@pytest.fixture(scope="module")
def two(tmp_path_factory, tmp_livery):
    """ncar with two strips, built."""
    tmp = tmp_path_factory.mktemp("two")
    rc, err, j = tmp_livery(tmp, TWO, key="marking")
    assert rc == 0, err
    return tmp / "liveries" / "ncar", j


@pytest.mark.parametrize("front, want", [
    ("", [("Marking banner", "TOP", 14), ("Marking legend", "FOOT", 8)]),        # both defaults, each its size
    ("marking: CUI\n", [("Marking banner", "CUI", 14), ("Marking legend", "FOOT", 8)]),   # a string: the first
    ("marking: {legend: LEGEND}\n", [("Marking banner", "TOP", 14), ("Marking legend", "LEGEND", 8)]),
    ("marking: {legend: none}\n", [("Marking banner", "TOP", 14)]),
    ("marking: none\n", []),                                                     # none: every strip off
])
def test_two_strips(pandoc_deck, tmp_path, two, front, want):
    d, j = two
    assert [s["name"] for s in j["marking"]] == ["banner", "legend"]
    assert j["marking"][1]["box"][1] == round(1 - 0.004 - 0.03, 4), "the legend at the foot"
    got = strips(pandoc_deck, tmp_path, "ncar", front, livery_dir=d)
    assert got == [want, want], got


def test_type_step_scales_it(pandoc_deck, tmp_path):
    # a deck's type step scales the strip, as HTML's --engine-type-scale and beamer's typescale do:
    # ncar's 7 pt marking role under type-body: 10 (its body is 12), to the half point
    got = strips(pandoc_deck, tmp_path, "ncar", "marking: CUI\ntype-body: 10\n")
    want = livery_mod.half_pt(7 * 10 / 12)
    assert all(s == [("Marking banner", "CUI", want)] for s in got), f"not {want} pt on every slide: {got}"


DECK_MARKING = """
local doc = pandoc.read(io.read("a"), "markdown")
local strips = pandoc.json.decode(arg[2])
for _, t in ipairs(C.marking_texts(doc.meta, strips, "test")) do
  print(t.strip.name .. "=" .. (t.text and pandoc.utils.stringify(t.text) or "-"))
end
"""


@pytest.mark.parametrize("front, want, warned", [
    ("", ["banner=TOP", "legend=FOOT"], False),
    ("marking: CUI\n", ["banner=CUI", "legend=FOOT"], False),
    ("marking: none\n", ["banner=-", "legend=-"], False),
    ("marking: {legend: none}\n", ["banner=TOP", "legend=-"], False),
    ("marking: {banner: CUI, legend: FOOTER}\n", ["banner=CUI", "legend=FOOTER"], False),
    ("marking: {footer: X}\n", ["banner=TOP", "legend=FOOT"], True),   # an unknown name: warned, ignored
])
def test_deck_marking(tmp_path, front, want, warned):
    strips = json.dumps([{"name": "banner", "text": "TOP"}, {"name": "legend", "text": "FOOT"}])
    out, warnings = lua_common(tmp_path, DECK_MARKING, f"---\n{front}---\n\nx\n" if front else "x\n", strips)
    assert out.split() == want
    assert bool(warnings) == warned, warnings


# engine-revealjs.lua under plain pandoc: the `quarto` it calls (the dependency and header text are
# Quarto's to place, so they go nowhere here; its json gives nil for null, as Quarto's does)
REVEAL_SHIM = """
local function nonull(v)
  if v == pandoc.json.null then return nil end
  if type(v) == "table" then for k, x in pairs(v) do v[k] = nonull(x) end end
  return v
end
quarto = {{ utils = {{ resolve_path = function(p) return "{engine}/" .. p end }},
           doc = {{ is_format = function(f) return f == "revealjs" or f == "html" end,
                   add_html_dependency = function() end, include_text = function() end }},
           log = {{ warning = function(s) io.stderr:write("WARNING: " .. s .. "\\n") end }},
           json = {{ decode = function(s) return nonull(pandoc.json.decode(s, false)) end }} }}
PANDOC_SCRIPT_FILE = "{engine}/engine-revealjs.lua"
dofile(PANDOC_SCRIPT_FILE)
"""


def test_html_every_slide(tmp_path):
    # HTML's strips and watermark on every slide, as PDF and pptx draw them: one opened by a rule and
    # the blocks before the first heading too (nasa-ccp: two strips by default), each slide's last
    shim = tmp_path / "reveal-shim.lua"
    shim.write_text(REVEAL_SHIM.format(engine=ENGINE))
    md = (f"---\nlivery: {ROOT / 'liveries' / 'nasa-ccp'}\nwatermark: DRAFT\n---\n\nLead.\n\n## a\n\nOne.\n\n"
          "---\n\nUntitled.\n\n# Part\n\n## b\n\nTwo.\n")
    html = subprocess.run([pandoc_exe(), "-t", "revealjs", "--slide-level", "2", "--lua-filter", str(shim)],
                          input=md, text=True, capture_output=True, check=True, cwd=tmp_path).stdout
    slides = [s for s in re.split(r"<section\b", html)[1:] if not s.startswith(">")]   # not a divider's stack
    texts = ["Lead", "One", "Untitled", "Part", "Two"]
    assert [next(t for t in texts if t in s) for s in slides] == texts, "a slide per heading, rule and lead"
    for s, t in zip(slides, texts):
        got = re.findall(r'class="engine-(marking engine-marking-\w+|watermark)"', s)
        assert got == ["marking engine-marking-banner", "marking engine-marking-legend", "watermark"], (t, got)
        assert s.rstrip().endswith("</section>") and "engine-watermark" in s.rsplit("<div", 1)[-1], \
            f"{t}: the watermark is not the slide's last element"
