"""Check a content title's color (frametitle.color) and its mirrored room (frametitle.symmetric)
from the build to every drawer's numbers.

    python3 -m pytest tests/test_frametitle.py

The color resolves to hex (a role, a palette name or a hex) and reaches livery.sty, livery.css and
the template's title style, where the default stays the body text's slot (tx1), so the bundled
templates are unchanged. Symmetric, the title's room has the mark's side's gap on both sides: the
template's title placeholder is centered on the slide and the css's left and right agree.
"""

import json

from pptx import Presentation
from pptx.util import Emu

from conftest import ROOT, run

TITLE_STYLE = "{http://schemas.openxmlformats.org/presentationml/2006/main}titleStyle"


def title_style(template):
    """The master's title style XML of a template.pptx."""
    from lxml import etree
    m = Presentation(str(template)).slide_master
    return etree.tostring(m._element.find(f".//{TITLE_STYLE}")).decode()


def test_default_is_the_body_text():
    j = json.loads((ROOT / "liveries" / "ncar" / "build" / "livery.json").read_text())
    assert j["frametitle"]["color"] == j["roles"]["foreground"] and j["frametitle"]["symmetric"] is False
    for name in ("ncar", "cu", "nsi", "nasa-jsc-eg"):
        style = title_style(ROOT / "liveries" / name / "template.pptx")
        assert '<a:schemeClr val="tx1"/>' in style, f"{name}: the title in the body text's slot"


def test_color_and_symmetric(tmp_path, tmp_livery):
    rc, err, j = tmp_livery(tmp_path, "  align: center\n  symmetric: true\n  color: \"#FFFFFF\"\n", key="frametitle")
    assert rc == 0, err
    ft = j["frametitle"]
    assert ft["color"] == "#FFFFFF" and ft["symmetric"] is True
    d = tmp_path / "liveries" / "ncar"
    sty, css = (d / "build" / "livery.sty").read_text(), (d / "build" / "livery.css").read_text()
    assert r"\definecolor{EngineFrametitleColor}{HTML}{FFFFFF}" in sty and r"\def\engine@ft@symmetric{}" in sty
    assert "--engine-title-color: #FFFFFF;" in css
    left = next(l for l in css.splitlines() if "--engine-title-left:" in l).split(":", 1)[1]
    right = next(l for l in css.splitlines() if "--engine-title-right:" in l).split(":", 1)[1]
    assert left == right, (left, right)
    assert run("template", d).returncode == 0
    prs = Presentation(str(d / "template.pptx"))
    assert '<a:srgbClr val="FFFFFF"/>' in title_style(d / "template.pptx")
    ph = next(s for s in prs.slide_master.placeholders if s.placeholder_format.type == 1)   # the title
    center = ph.left + ph.width / 2
    assert abs(center - prs.slide_width / 2) < Emu(12700), "the title placeholder centered on the slide"


def test_role_name_resolves(tmp_path, tmp_livery):
    rc, err, j = tmp_livery(tmp_path, "  color: primary\n", key="frametitle")
    assert rc == 0, err
    assert j["frametitle"]["color"] == j["roles"]["primary"]
