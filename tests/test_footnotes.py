"""Check slide footnotes, `^[...]` notes lifted beside the dagger lines (engine-common.lua's
lift_footnotes() and take_footnotes(); the pptx filter inline-footnotes.lua; pptx_steps/style_footnotes.py),
and the references slide (ensure_refs(); slide_layout.py's references_style()).

    python3 -m pytest tests/test_footnotes.py

A slide's `^[...]` notes take the printer's markers after its own dagger lines (`† ‡ § ¶`): each a
superscript mark in the text (a Span `.engine-noteref`) and a paragraph led by the marker, which
the format filters then treat as any slide footnote. A note in the speaker notes stays pandoc's; a
fifth, or one holding more than paragraphs, warns and stays. engine-common.lua runs under
`pandoc lua` with a `quarto` shim (the harness below), the pptx filter under plain pandoc, the
step in a subprocess under ncar's LIVERY_DIR. A deck that cites without a `::: {#refs}` gets one
References slide last, with a warning, and none when it has the div or cites only cross-references;
on the pptx references slide every entry is in the footnote size, hanging by 1.5 em. Needs pandoc
($QUARTO_PANDOC).
"""

import os
import re
import shutil
import subprocess
import sys
import zipfile

import pytest

from conftest import ROOT

ENGINE = ROOT / "_extensions" / "engine"
NCAR = ROOT / "liveries" / "ncar"

# engine-common.lua's functions on markdown from stdin: `take` prints take_footnotes()'s kept
# blocks, a rule, the footnotes found; `refs` the blocks after ensure_refs() twice and whether they
# hold a #refs div; each warning on
# a line of its own after the markdown
HARNESS = r"""
local warnings = {}
quarto = { log = { warning = function(s) table.insert(warnings, s) end }, json = { decode = pandoc.json.decode } }
local C = dofile(arg[1])
local out = pandoc.Blocks({})
local doc = pandoc.read(io.read("a"), "markdown")
if arg[2] == "take" then
  local kept, found = C.take_footnotes(doc.blocks)
  out:extend(kept); out:insert(pandoc.HorizontalRule()); out:extend(found)
elseif arg[2] == "refs" then   -- twice: the second finds the slide the first added
  out = C.ensure_refs(C.ensure_refs(doc.blocks, doc.meta, 2), doc.meta, 2)
  out:insert(pandoc.Para(C.refs_div(out) and "(a refs div)" or "(no refs div)"))   -- the writer drops an empty div last in a document
end
io.write(pandoc.write(pandoc.Pandoc(out), "markdown", { wrap_text = "none" }))
for _, w in ipairs(warnings) do io.write("\nWARNING: " .. w) end
"""


def pandoc():
    exe = os.environ.get("QUARTO_PANDOC") or shutil.which("pandoc")
    if not exe:
        pytest.fail("pandoc not found: set QUARTO_PANDOC or source etc/config_env.sh")
    return exe


def common(tmp_path, verb, md):
    """The harness's output on md: (the markdown, the warnings)."""
    script = tmp_path / "harness.lua"
    script.write_text(HARNESS)
    r = subprocess.run([pandoc(), "lua", str(script), str(ENGINE / "engine-common.lua"), verb],
                       input=md, text=True, capture_output=True, check=True)
    text, *warnings = r.stdout.split("\nWARNING: ")
    return text, warnings


def test_lifted_after_daggers(tmp_path):
    """One dagger line and two notes: the notes take ‡ and §, three footnotes in order."""
    md = ("Text^[First note.] and more[^2].\n\n† A dagger line.\n\n"
          "[^2]: Second *note*\n\n    in two paragraphs.\n")
    text, warnings = common(tmp_path, "take", md)
    body, foot = re.split(r"\n-{20,}\n", text, maxsplit=1)
    assert "[^‡^]{.engine-noteref}" in body and "[^§^]{.engine-noteref}" in body, body
    assert not re.search(r"\[\^\d", text), "every note lifted"
    paras = [p for p in foot.strip().split("\n\n")]
    assert paras == ["† A dagger line.", "‡ First note.", "§ Second *note*\\\nin two paragraphs."], paras
    assert warnings == []


def test_speaker_notes_stay(tmp_path):
    """A note in the speaker notes is pandoc's, and takes no marker."""
    md = "Text^[Lifted.]\n\n::: notes\nSay this^[Kept.]\n:::\n"
    text, _ = common(tmp_path, "take", md)
    body, foot = re.split(r"\n-{20,}\n", text, maxsplit=1)
    assert "Say this[^1]" in body and "[^†^]{.engine-noteref}" in body, body
    assert foot.strip().split("\n\n") == ["† Lifted.", "[^1]: Kept."], foot


def test_fifth_and_blocks_stay(tmp_path):
    """Past ¶ a note warns and stays; so does one that holds a list (and takes no marker)."""
    md = "A^[a] B[^b] C^[c] D^[d] E^[e] F^[f]\n\n[^b]: Some text.\n\n    - an item\n"
    text, warnings = common(tmp_path, "take", md)
    body, foot = re.split(r"\n-{20,}\n", text, maxsplit=1)
    assert foot.strip().split("\n\n")[:4] == ["† a", "‡ c", "§ d", "¶ e"], foot
    assert "B[^1]" in body and "F[^2]" in body and "[^2]: f" in foot, text
    assert len(warnings) == 2 and any("more than paragraphs" in w for w in warnings), warnings


def slides_of(path):
    with zipfile.ZipFile(path) as z:
        return sorted(n for n in z.namelist() if n.startswith("ppt/slides/slide"))


def test_pptx_filter(tmp_path):
    """pptx: the notes lifted (pandoc's Notes slide gone); on a two-column slide, into the last column."""
    md = ("## One\n\nText^[A note.]\n\n## Two\n\n:::: {.columns}\n::: {.column}\nLeft^[In a column.]\n:::\n"
          "::: {.column}\nRight\n:::\n::::\n")
    out = tmp_path / "t.pptx"
    subprocess.run([pandoc(), "--lua-filter", str(ENGINE / "inline-footnotes.lua"),
                    "--reference-doc", str(NCAR / "template.pptx"), "-o", str(out)],
                   input=md, text=True, check=True, capture_output=True)
    assert len(slides_of(out)) == 2, "no Notes slide"
    from pptx import Presentation
    prs = Presentation(str(out))
    texts = [[sh.text_frame.text for sh in s.placeholders if sh.has_text_frame] for s in prs.slides]
    assert any("† A note." in t for t in texts[0]), texts[0]
    right = [t for t in texts[1] if t.startswith("Right")]
    assert right and "† In a column." in right[0], texts[1]
    # the step: the mark in the text (a superscript †) takes the accent, as HTML and PDF set it
    import json
    from pptx.oxml.ns import qn
    env = {**os.environ, "LIVERY_DIR": str(NCAR), "PYTHONPATH": str(ENGINE)}
    subprocess.run([sys.executable, "-c", STEP, str(out)], env=env, check=True, capture_output=True)
    accent = json.loads((NCAR / "build" / "livery.json").read_text())["pptx"]["scheme"]["accent4"].lstrip("#").upper()
    marks = [r for s in Presentation(str(out)).slides for sh in s.placeholders if sh.has_text_frame
             for p in sh.text_frame.paragraphs for r in p.runs
             if r.text == "†" and (r._r.find(qn("a:rPr")) is not None and r._r.find(qn("a:rPr")).get("baseline"))]
    assert len(marks) == 2 and all(str(r.font.color.rgb) == accent for r in marks), [(r.text, r.font.color.rgb) for r in marks]


STEP = """
import sys
from pathlib import Path
from pptx import Presentation
from pptx_steps import style_footnotes
prs = Presentation(sys.argv[1]); style_footnotes.run(prs, Path(sys.argv[1])); prs.save(sys.argv[1])
"""


def test_pptx_section_sign(pandoc_deck, tmp_path):
    """style_footnotes takes § and ¶ as it takes † and ‡: the footnote size, muted, the mark in the accent."""
    import json
    from pptx import Presentation
    from pptx.util import Pt
    pandoc_deck(tmp_path, "## S\n\n§ A third footnote.\n\n¶ A fourth.\n", "t.pptx")
    env = {**os.environ, "LIVERY_DIR": str(NCAR), "PYTHONPATH": str(ENGINE)}
    subprocess.run([sys.executable, "-c", STEP, str(tmp_path / "t.pptx")], env=env, check=True, capture_output=True)
    j = json.loads((NCAR / "build" / "livery.json").read_text())
    body = next(sh for sh in Presentation(str(tmp_path / "t.pptx")).slides[0].placeholders
                if sh.has_text_frame and "third" in sh.text_frame.text)
    for p, mark in zip(body.text_frame.paragraphs, "§¶"):
        assert p.runs[0].text.strip() == mark
        assert all(r.font.size == Pt(j["type"]["roles"]["footnote"]["pt"]) for r in p.runs), p.text
        assert str(p.runs[1].font.color.rgb) == j["roles"]["muted"][1:].upper()


BIB = "---\nbibliography: refs.bib\n---\n\n"


def test_refs_added_once(tmp_path):
    """A deck that cites without a #refs div gets one References slide last, and a warning."""
    text, warnings = common(tmp_path, "refs", BIB + "## Cites\n\nAs @knuth84 wrote.\n")
    assert text.count("## References") == 1 and text.rstrip().endswith("## References\n\n::: {#refs}\n:::\n\n(a refs div)"), text
    assert text.index("## Cites") < text.index("## References")
    assert len(warnings) == 1 and "cited but no ::: {#refs}" in warnings[0], warnings


@pytest.mark.parametrize("md", [
    BIB + "## Cites\n\nAs @knuth84 wrote.\n\n## Sources\n\n::: {#refs}\n:::\n",   # the author's own slide
    BIB + "## Plot\n\nSee @fig-plot.\n",                                             # a cross-reference only
    "## Cites\n\nAs @knuth84 wrote.\n",                                              # no bibliography
], ids=["has-refs", "crossref", "no-bibliography"])
def test_refs_left_alone(tmp_path, md):
    text, warnings = common(tmp_path, "refs", md)
    assert "## References" not in text and warnings == [], text


REFS_STEP = """
import sys
from pathlib import Path
from pptx import Presentation
from pptx_steps import slide_layout
prs = Presentation(sys.argv[1]); slide_layout.run(prs, Path(sys.argv[1])); prs.save(sys.argv[1])
"""


def test_pptx_references(pandoc_deck, tmp_path):
    """The references slide (slide-layout.lua's `references`): each entry in the footnote size, a
    1.5 em hanging indent of it, no bullet."""
    import json
    from pptx import Presentation
    from pptx.oxml.ns import qn
    from pptx.util import Pt
    md = ("## References\n\n::: {#refs .references .csl-bib-body .hanging-indent}\n"
          "::: {#ref-knuth84 .csl-entry}\nKnuth, Donald E. 1984. *The TeXbook*. Addison-Wesley.\n:::\n\n"
          "::: {#ref-sk .csl-entry}\nSkamarock, William C. 2012. A long title that runs on.\n:::\n:::\n\n"
          "::: notes\nengine-layout: references\n:::\n")
    pandoc_deck(tmp_path, md, "t.pptx")
    env = {**os.environ, "LIVERY_DIR": str(NCAR), "PYTHONPATH": str(ENGINE)}
    subprocess.run([sys.executable, "-c", REFS_STEP, str(tmp_path / "t.pptx")], env=env, check=True, capture_output=True)
    pt = json.loads((NCAR / "build" / "livery.json").read_text())["type"]["roles"]["footnote"]["pt"]
    slide = Presentation(str(tmp_path / "t.pptx")).slides[0]
    body = next(sh for sh in slide.placeholders if sh.has_text_frame and "Knuth" in sh.text_frame.text)
    paras = [p for p in body.text_frame.paragraphs if p.text.strip()]
    assert [p.text.split(",")[0] for p in paras] == ["Knuth", "Skamarock"]
    hang = round(1.5 * pt * 12700)
    for p in paras:
        ppr = p._p.find(qn("a:pPr"))
        assert (ppr.get("marL"), ppr.get("indent")) == (str(hang), str(-hang)), p.text
        assert ppr.find(qn("a:buNone")) is not None
        assert all(r.font.size == Pt(pt) for r in p.runs), p.text
    assert "engine-layout" not in slide.notes_slide.notes_text_frame.text
