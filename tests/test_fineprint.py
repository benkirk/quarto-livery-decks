"""Check the cover's fine print: a livery's default (compositions.<cover>.fineprint.text), the deck's
`fineprint:` over it, and its pptx drawer (slide_layout.py cover_fineprint()).

    python3 -m pytest tests/test_fineprint.py

A tmp flavor of ncar sets a two-paragraph default on `title`: the build gives every cover its
paragraphs and refuses the text on a closing. slide-layout.lua (under a `quarto` shim) flattens the
deck's fineprint: to `engine-fineprint`, a line per paragraph (a line break parts one too: Quarto
reads a metadata file's blank line as one); the pptx step draws the deck's, else the default, in
the box, size and color the livery gives, and none for `fineprint: none`. Needs pandoc.
"""

import json
import os
import shutil
import subprocess
import sys

import pytest

from conftest import ENGINE, ROOT, pandoc_exe, run, slide_layout_shim

NAME = "fp"
TEXT = "First notice, (a) and (b).\n\nSecond notice."
STEP = """
import sys
from pathlib import Path
from pptx import Presentation
from pptx_steps import slide_layout
p = Path(sys.argv[1]); prs = Presentation(p); slide_layout.run(prs, p); prs.save(p)
"""


def flavor(tmp, compositions):
    """ncar's flavor NAME with a compositions: block, built: (its directory, the run)."""
    liv = tmp / "liveries"
    shutil.copytree(ROOT / "liveries" / "ncar", liv / "ncar", ignore=shutil.ignore_patterns("reference"))
    d = liv / NAME; d.mkdir()
    (d / "_brand.yml").write_text("{}\n")
    (d / "livery.yml").write_text(f"schema: 1\nname: {NAME}\nextends: ncar\ncompositions:\n{compositions}")
    return d, run("build", d)


@pytest.fixture(scope="module")
def fp(tmp_path_factory):
    d, r = flavor(tmp_path_factory.mktemp("fp"), f"  title:\n    fineprint: {{text: {json.dumps(TEXT)}}}\n")
    assert r.returncode == 0, r.stderr
    return d, json.loads((d / "build" / "livery.json").read_text())


def test_default_reaches_every_cover(fp):
    _, j = fp
    c = j["compositions"]
    for k in ("title", "title-photo", "title-light", "title-photo-light", "title-band"):
        assert c[k]["fineprint"]["paragraphs"] == ["First notice, (a) and (b).", "Second notice."], k
        assert "text" not in c[k]["fineprint"], "the text becomes its paragraphs"
    assert all(not (c[k]["fineprint"] or {}).get("paragraphs") for k in ("closing", "section", "feature"))
    sty = (fp[0] / "build" / "livery.sty").read_text()
    assert r"\def\engine@c@title@fineprint@text{First notice, (a) and (b).\par Second notice.}" in sty


def test_text_on_a_closing_fails(tmp_path):
    _, r = flavor(tmp_path, "  closing:\n    fineprint: {box: [0.1, 0.8, 0.5, 0.1], anchor: top, text: No.}\n")
    assert r.returncode != 0 and "compositions.closing.fineprint.text: only a cover" in r.stderr, r.stderr


def flatten(tmp_path, front):
    """slide-layout.lua's engine-fineprint for the front matter (None: not set)."""
    r = subprocess.run([pandoc_exe(), "--lua-filter", str(slide_layout_shim(tmp_path)), "-t", "json"], input=f"---\n{front}---\n\n## a\n\nx\n",
                       text=True, capture_output=True, check=True)
    m = json.loads(r.stdout)["meta"].get("engine-fineprint")
    return m and m["c"]


@pytest.mark.parametrize("front, want", [
    ("fineprint: One line.\n", "One line."),
    ("fineprint: |\n  First para.\n\n  Second para.\n", "First para.\nSecond para."),
    ("fineprint: \"First line.\\\\\\nSecond line.\"\n", "First line.\nSecond line."),   # a line break
    ("fineprint: none\n", "none"),
    ("title: T\n", None),
])
def test_flattened_for_pptx(tmp_path, front, want):
    assert flatten(tmp_path, front) == want


def drawn(pandoc_deck, tmp, d, front):
    """The cover's Fineprint box after slide-layout.lua and the step: (its paragraphs, size pt, color
    hex, anchor), or None."""
    tmp.mkdir()
    pandoc_deck(tmp, f"---\ntitle: A deck\n{front}---\n\n## a\n\nOne.\n", "t.pptx", "--lua-filter", str(slide_layout_shim(tmp)))
    env = {**os.environ, "LIVERY_DIR": str(d), "PYTHONPATH": str(ENGINE)}
    r = subprocess.run([sys.executable, "-c", STEP, str(tmp / "t.pptx")], env=env, capture_output=True, text=True)
    assert r.returncode == 0, r.stderr
    from pptx import Presentation
    box = next((s for s in Presentation(str(tmp / "t.pptx")).slides[0].shapes if s.name == "Fineprint"), None)
    if box is None:
        return None
    run0 = box.text_frame.paragraphs[0].runs[0]
    return ([p.text for p in box.text_frame.paragraphs], run0.font.size.pt, str(run0.font.color.rgb),
            box.text_frame.vertical_anchor)


def test_pptx_default_and_deck(pandoc_deck, tmp_path, fp):
    d, j = fp
    fpc = j["compositions"]["title"]["fineprint"]
    got = drawn(pandoc_deck, tmp_path / "a", d, "")
    assert got and got[0] == ["First notice, (a) and (b).", "Second notice."], got
    assert got[1] == j["type"]["roles"]["fineprint"]["pt"], "the fineprint role's size"
    assert got[2] == j["roles"][fpc["role"]][1:].upper(), "its role's color"
    from pptx.enum.text import MSO_ANCHOR
    assert got[3] == (MSO_ANCHOR.BOTTOM if fpc["anchor"] == "bottom" else MSO_ANCHOR.TOP)
    got = drawn(pandoc_deck, tmp_path / "b", d, "fineprint: |\n  The deck's own.\n\n  Line two.\n")
    assert got and got[0] == ["The deck’s own.", "Line two."], got
    assert drawn(pandoc_deck, tmp_path / "c", d, "fineprint: none\n") is None
    assert drawn(pandoc_deck, tmp_path / "e", ROOT / "liveries" / "ncar", "") is None, "no default, no box"


COVER = r"""
\documentclass[aspectratio=169]{beamer}
\usetheme[livery=liveries/fp,fonts=bundled]{Engine}
\title{A deck}
%s
\begin{document}
\maketitle
\end{document}
"""


@pytest.mark.parametrize("preamble, shown", [("", True), (r"\titlefineprint{}", False), (r"\titlefineprint{Own words.}", False)])
def test_latex_default(fp, tmp_path, preamble, shown):
    """The theme draws the livery's default unless the deck's \\titlefineprint says otherwise."""
    if not (shutil.which("xelatex") and shutil.which("latexmk")):
        pytest.skip("xelatex and latexmk: the LaTeX layer compiles only with TeX")
    top = fp[0].parent.parent   # the tmp root holding liveries/ncar and liveries/fp
    (tmp_path / "c.tex").write_text(COVER % preamble)
    env = {**os.environ, "TEXINPUTS": f"{ENGINE}//:{top}/:{ROOT}/:", "TTFONTS": f"{top}/:{ROOT}/:{ENGINE}//:"}
    r = subprocess.run(["latexmk", "-xelatex", "-interaction=nonstopmode", "-halt-on-error", "c.tex"],
                       cwd=tmp_path, env=env, capture_output=True, text=True)
    assert r.returncode == 0, (tmp_path / "c.log").read_text(errors="replace")[-1500:]
    text = " ".join(subprocess.run(["pdftotext", str(tmp_path / "c.pdf"), "-"], capture_output=True, text=True).stdout.split())
    assert ("First notice, (a) and (b). Second notice." in text) == shown, text
    assert ("Own words." in text) == ("Own" in preamble), text
