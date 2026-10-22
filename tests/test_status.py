"""Check status items: - [x] ✓, - [!] ⚠, - [-] ✗ (engine-common.lua's status_lists(), grid-pptx.lua,
pptx_steps/status_glyphs.py).

    python3 -m pytest tests/test_status.py

In a `::: {.status}` list, pandoc's task markers become glyphs, and the pptx step makes each one the
paragraph's bullet (a:buChar) in the livery's boxes.status color, its text without the glyph, a
failed item's text muted; `- [ ]` stays a task list's unchecked box; outside a card or a .status
list `- [x]` is left alone, a literal glyph is a status item anywhere. A second run of the step
changes nothing. pandoc with the filter on ncar's template, the step in a subprocess under
LIVERY_DIR. Needs pandoc ($QUARTO_PANDOC).
"""

import json
import os
import subprocess
import sys
import zipfile

import pytest

from conftest import ENGINE, ROOT, pandoc_exe

NCAR = ROOT / "liveries" / "ncar"
MD = ("## Status\n\n::: {.status}\n- [x] Done thing\n- [!] Risky thing\n- [-] Failed thing\n- [ ] Open thing\n:::\n\n"
      "## Plain\n\n- [x] A task outside\n- ✓ A literal glyph\n\n"
      "## ⚠︎ Careful\n\n✓ A paragraph, not an item\n")
STEP = """
import sys
from pathlib import Path
from pptx import Presentation
from pptx_steps import status_glyphs
prs = Presentation(sys.argv[1]); status_glyphs.run(prs, Path(sys.argv[1])); prs.save(sys.argv[1])
"""


@pytest.fixture(scope="module")
def deck(tmp_path_factory):
    pandoc = pandoc_exe()
    out = tmp_path_factory.mktemp("status") / "t.pptx"
    subprocess.run([pandoc, "--lua-filter", str(ENGINE / "grid-pptx.lua"), "--reference-doc", str(NCAR / "template.pptx"),
                    "-o", str(out)], input=MD, text=True, check=True, capture_output=True)
    env = {**os.environ, "LIVERY_DIR": str(NCAR), "PYTHONPATH": str(ENGINE)}
    step = lambda: subprocess.run([sys.executable, "-c", STEP, str(out)], env=env, check=True, capture_output=True)
    slides = lambda: {n: z.read(n) for z in [zipfile.ZipFile(out)] for n in z.namelist() if n.startswith("ppt/slides/")}
    step()
    first = slides()
    step()
    from pptx import Presentation
    return Presentation(str(out)), first, slides()


def items(slide):
    body = next(sh for sh in slide.placeholders if sh.placeholder_format.idx == 1)
    return body.text_frame.paragraphs


def bullet(p):
    ppr = p._p.find("{*}pPr")
    bu = ppr.find("{*}buChar") if ppr is not None else None
    clr = ppr.find("{*}buClr/{*}srgbClr") if ppr is not None else None
    return (bu.get("char") if bu is not None else None, clr.get("val") if clr is not None else None)


def test_status_list(deck):
    prs, *_ = deck
    status = json.loads((NCAR / "build" / "livery.json").read_text())["boxes"]["status"]
    muted = json.loads((NCAR / "build" / "livery.json").read_text())["roles"]["muted"][1:]
    ps = items(prs.slides[0])
    assert [p.text for p in ps] == ["Done thing", "Risky thing", "Failed thing", "☐ Open thing"], [p.text for p in ps]
    assert [bullet(p) for p in ps[:3]] == [("✓", status["done"][1:]), ("⚠", status["warn"][1:]), ("✗", status["fail"][1:])]
    assert bullet(ps[3]) == (None, None), "an unchecked task is not a status item"
    assert all(str(r.font.color.rgb) == muted for r in ps[2].runs), "the failed item's text muted"
    assert all(r.font.color.type is None for r in ps[0].runs), "a done item's text colored"


def test_outside(deck):
    prs, *_ = deck
    ps = items(prs.slides[1])
    assert ps[0].text.startswith("☒") and bullet(ps[0]) == (None, None), "- [x] outside a card or .status list changed"
    assert ps[1].text == "A literal glyph" and bullet(ps[1])[0] == "✓", "a literal glyph is a status item anywhere"


def test_idempotent(deck):
    _, first, second = deck
    assert first == second, "a second run of status_glyphs changed the deck"


def test_not_list_items(deck):
    """A slide title and a paragraph led by a glyph are not status items: they keep the glyph."""
    prs, *_ = deck
    slide = prs.slides[2]
    assert slide.shapes.title.text_frame.text.startswith("⚠"), "the title lost its glyph"
    p = items(slide)[0]
    assert p.text.startswith("✓") and bullet(p) == (None, None), "a paragraph made a status item"
